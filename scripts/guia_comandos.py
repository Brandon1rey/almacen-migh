"""Genera la guía de instalación con comandos en Word y PDF a partir del Markdown."""
import sys

# La consola de Windows no siempre es UTF-8: sin esto, imprimir «→» o «·» lanza
# UnicodeEncodeError y el script parece fallar aunque ya haya hecho su trabajo.
for _flujo in (sys.stdout, sys.stderr):
    try:
        _flujo.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

import pathlib
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parents[1]
FUENTE = RAIZ / "docs" / "Guia_instalacion_comandos.md"
DOCX = RAIZ / "docs" / "Guia_instalacion_comandos.docx"
AZUL = RGBColor(0x1F, 0x3B, 0x73)
GRIS = RGBColor(0x4B, 0x55, 0x63)


def sombrear(parrafo, color="F2F4F7"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    parrafo._p.get_or_add_pPr().append(shd)


def borde(parrafo, color="D0D5DD"):
    pPr = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), color)
        bordes.append(b)
    pPr.append(bordes)


def bloque_codigo(doc, lineas):
    for i, linea in enumerate(lineas):
        p = doc.add_paragraph()
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.line_spacing = 1.0
        sombrear(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def tabla(doc, filas):
    encabezado, cuerpo = filas[0], filas[1:]
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    for i, texto in enumerate(encabezado):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(texto.replace("**", ""))
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = AZUL
    for fila in cuerpo:
        celdas = t.add_row().cells
        for i, texto in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(texto.replace("**", "").replace("`", ""))
            run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def escribir_runs(p, texto, tam=9.5, negritas=False, color=None):
    """Interpreta **negritas** y `codigo` del Markdown."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not trozo:
            continue
        run = p.add_run(trozo.strip("`").strip("*"))
        run.font.size = Pt(tam)
        if color is not None:
            run.font.color.rgb = color
        if negritas or (trozo.startswith("**") and trozo.endswith("**")):
            run.bold = True
        if trozo.startswith("`") and trozo.endswith("`"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(tam - 0.5)


def parrafo(doc, texto, negritas=False, color=None, tam=9.5, espacio=4):
    p = doc.add_paragraph()
    escribir_runs(p, texto, tam, negritas, color)
    p.paragraph_format.space_after = Pt(espacio)
    return p


def construir():
    if not FUENTE.exists():
        raise SystemExit(f"no existe {FUENTE}")
    lineas = FUENTE.read_text(encoding="utf-8").splitlines()

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.5)
        s.left_margin = s.right_margin = Cm(1.7)

    en_codigo = False
    codigo = []
    tabla_pendiente = []
    titulo_puesto = False

    for linea in lineas:
        limpia = linea.rstrip()

        if limpia.startswith("```"):
            if en_codigo:
                bloque_codigo(doc, codigo)
                codigo, en_codigo = [], False
            else:
                en_codigo = True
            continue
        if en_codigo:
            codigo.append(limpia)
            continue

        if limpia.startswith("|"):
            tabla_pendiente.append([c.strip() for c in limpia.strip("|").split("|")])
            continue
        if tabla_pendiente:
            filas = [f for f in tabla_pendiente if not all(set(c) <= set("-: ") for c in f)]
            if filas:
                tabla(doc, filas)
            tabla_pendiente = []

        if not limpia:
            continue

        if limpia.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(limpia[2:])
            run.bold = True
            run.font.size = Pt(17)
            run.font.color.rgb = None
            p.paragraph_format.space_after = Pt(2)
            parrafo(doc, "Comercializadora MIGH de México · Aseguramiento de Calidad", False, GRIS, 9, 12)
            titulo_puesto = True
        elif limpia.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(limpia[3:])
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = AZUL
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif limpia.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(limpia[4:])
            run.bold = True
            run.font.size = Pt(10.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
        elif limpia.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            escribir_runs(p, limpia[2:], 9.5)
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", limpia):
            numero, resto = re.match(r"^(\d+)\.\s+(.*)$", limpia).groups()
            p = doc.add_paragraph()
            run = p.add_run(f"{numero}. ")
            run.bold = True
            run.font.size = Pt(9.5)
            escribir_runs(p, resto, 9.5)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
        else:
            parrafo(doc, limpia)

    if codigo:
        bloque_codigo(doc, codigo)
    if tabla_pendiente:
        tabla(doc, [f for f in tabla_pendiente if not all(set(c) <= set("-: ") for c in f)])

    parrafo(doc, "Documento de trabajo. Verificar en sitio antes de entregar la PC.", False, GRIS, 7.5, 0)
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    return DOCX


def a_pdf(docx_path):
    import win32com.client as win32
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(docx_path))
    pdf = docx_path.with_suffix(".pdf")
    doc.ExportAsFixedFormat(str(pdf), 17)
    doc.Close(False)
    word.Quit()
    return pdf


if __name__ == "__main__":
    ruta = construir()
    print("docx:", ruta)
    print("pdf :", a_pdf(ruta))
