"""Siembra un juego de datos DEMO con ubicaciones pobladas y un historial de movimientos.

No toca el inventario real: escribe en data/demo/ y su propia bitácora.
Determinista (semilla fija): dos corridas producen el mismo demo.
"""
import datetime as dt
import difflib
import pathlib
import random
import re
import shutil
import sqlite3
import sys
import unicodedata

import docx
import openpyxl
from openpyxl.styles import Alignment, Font

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from app import ubicaciones  # noqa: E402

SEMILLA = 42
BASE = RAIZ / "data" / "Inventario_MP.xlsx"
DEMO_DIR = RAIZ / "data" / "demo"
DEMO_XLSX = DEMO_DIR / "Inventario_MP_demo.xlsx"
DEMO_DB = DEMO_DIR / "bitacora_demo.db"
MAESTRA = pathlib.Path(r"C:\Users\asegm\Desktop\USB BRANDON\trabajo2\auditoria-interna"
                       r"\04_PAQUETE_POR_CHECKLIST\CHK-A-12 PRP Almacenamiento y transporte"
                       r"\01_DOCUMENTOS\F02-ALM-PNO-001__Lista maestra de alérgenos.docx")
HOY = dt.date(2026, 10, 22)
USUARIOS = ["Esteban Soto", "Galilea Medina", "Maria Elena Becerra"]
NO_ALERGENO = {"", "no contiene", "n/a", "ninguno"}

# apoyo para el demo: la lista maestra identifica por SKU y el Excel del área no lo tiene.
# Se usa la palabra clave solo cuando el nombre no empata con la lista maestra.
PALABRAS_ALERGENO = {
    "Lácteos y derivados": ("LECHE", "LACTOSA", "SUERO", "WPC", "WPI", "WPH", "MPI ",
                            "CASEINATO", "CASEINA", "MANTEQUILLA", "LACTEO", "MILK"),
    "Soya": ("SOYA", "SOJA"),
    "Trigo": ("TRIGO", "GLUTEN"),
    "Pescado": ("PESCADO",),
    "Huevo": ("HUEVO", "ALBUMINA"),
}

# material de empaque y producto terminado: el Excel del área es solo de MP
EMPAQUE = [
    ("BOLSA LAMINADA 400G GROWVIT", "Empaque primario", "IMPREMEX", "pza", 48000, None, "BL-260901", "MEP-B"),
    ("BOLSA LAMINADA 900G GROWVIT", "Empaque primario", "IMPREMEX", "pza", 22500, None, "BL-260902", "MEP-B"),
    ("BOLSA LAMINADA 360G FORMULA 1", "Empaque primario", "IMPREMEX", "pza", 60000, None, "BL-260915", "MEP-B"),
    ("BOTE METALICO 400G BLANCO", "Empaque primario", "LATAS DEL NORTE", "pza", 36000, None, "BM-260910", "MEP-B"),
    ("BOTE METALICO 900G BLANCO", "Empaque primario", "LATAS DEL NORTE", "pza", 18000, None, "BM-260911", "MEP-B"),
    ("TAPA METALICA 400G", "Empaque primario", "LATAS DEL NORTE", "pza", 40000, None, "TM-260828", "MEP-C"),
    ("TAPA METALICA 900G", "Empaque primario", "LATAS DEL NORTE", "pza", 15000, None, "TM-260829", "MEP-C"),
    ("CUCHARA DOSIFICADORA 4.6 ML", "Empaque primario", "PLASTIMEX", "pza", 120000, None, "CD-260820", "MEP-C"),
    ("ETIQUETA GROWVIT VAINILLA 400G", "Empaque primario", "ETIQUETAS MX", "pza", 50000, None, "ET-260918", "MEP-C"),
    ("ETIQUETA FORMULA 1 ETAPA 1 900G", "Empaque primario", "ETIQUETAS MX", "pza", 30000, None, "ET-260919", "MEP-B"),
    ("CAJA CORRUGADA 12 PIEZAS", "Empaque secundario", "CARTONES DEL BAJIO", "pza", 4200, None, "CJ-260905", "MEP-D"),
    ("CAJA CORRUGADA 24 PIEZAS", "Empaque secundario", "CARTONES DEL BAJIO", "pza", 1800, None, "CJ-260906", "MEP-D"),
    ("TARIMA DE MADERA 1.2 X 1.0", "Embalaje", "TARIMAS SAN JOSE", "pza", 240, None, "TR-260701", "MEP-E"),
    ("PLAYO GRADO ALIMENTICIO", "Embalaje", "PLASTIMEX", "kg", 85, None, "PL-260812", "MEP-E"),
    ("SELLO DE SEGURIDAD NUMERADO", "Empaque secundario", "IMPREMEX", "pza", 12000, None, "SS-260914", "MEP-E"),
    ("BOLSA LAMINADA 400G FORMULA 1", "Empaque primario", "IMPREMEX", "pza", 8000, None, "BL-260622", "JAULA-PT"),
]
PT = [
    ("GROWVIT VAINILLA 400G", "Producto terminado", "MIGH", "pza", 960, "2027-08-31", "GV-261015", "MPT-B"),
    ("GROWVIT CHOCOLATE 400G", "Producto terminado", "MIGH", "pza", 480, "2027-09-15", "GV-261016", "MPT-B"),
    ("FORMULA 1 ETAPA 1 900G", "Producto terminado", "MIGH", "pza", 360, "2027-10-02", "F1-261017", "MPT-B"),
]


def normalizar(texto):
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9]+", " ", s)).strip().upper()


def catalogo_alergenos():
    if not MAESTRA.exists():
        return {}
    doc = docx.Document(str(MAESTRA))
    salida = {}
    for fila in doc.tables[0].rows[1:]:
        celdas = [c.text.strip() for c in fila.cells]
        if len(celdas) >= 3 and celdas[0].startswith("MIGH-"):
            salida.setdefault(normalizar(celdas[1]), celdas[2])
    return salida


def clasificar(descripcion, catalogo):
    clave = normalizar(descripcion)
    if clave in catalogo:
        return catalogo[clave], "lista maestra"
    parecido = difflib.get_close_matches(clave, list(catalogo), n=1, cutoff=0.88)
    if parecido and catalogo[parecido[0]].strip().lower() not in NO_ALERGENO:
        return catalogo[parecido[0]], "lista maestra (parecido)"
    for alergeno, palabras in PALABRAS_ALERGENO.items():
        if any(p in f" {clave} " for p in palabras):
            return alergeno, "palabra clave"
    return "No contiene", "sin registro"


def repartir(recursos, grupos, tope_por_rack=3):
    """Cada material se queda junto hasta llenar el tope del rack; luego pasa al siguiente."""
    asignaciones, puntero = {}, 0
    for filas in grupos:
        pendientes = list(filas)
        while pendientes:
            rack = recursos[puntero % len(recursos)]
            cupo = min(tope_por_rack, len(pendientes))
            for f in pendientes[:cupo]:
                asignaciones[f["fila"]] = rack
            pendientes = pendientes[cupo:]
            puntero += 1
    return asignaciones


def sembrar():
    random.seed(SEMILLA)
    DEMO_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BASE, DEMO_XLSX)

    wb = openpyxl.load_workbook(DEMO_XLSX)
    ws = wb.worksheets[0]
    col_ubicacion = next((c for c in range(1, ws.max_column + 1)
                          if normalizar(ws.cell(row=3, column=c).value) == "UBICACION"), None)
    if col_ubicacion is None:
        raise SystemExit("la base no tiene la columna UBICACIÓN; corre preparar_inventario.py")

    # ---------- filas nuevas de empaque y PT ----------
    fila = ws.max_row + 1
    nuevas = {}
    for desc, cat, prov, unidad, cantidad, caducidad, lote, zona in EMPAQUE + PT:
        ws.cell(row=fila, column=1, value=desc)
        ws.cell(row=fila, column=2, value=cat)
        ws.cell(row=fila, column=3, value=prov)
        ws.cell(row=fila, column=4, value=unidad)
        ws.cell(row=fila, column=5, value=cantidad)
        if caducidad:
            ws.cell(row=fila, column=6, value=dt.date.fromisoformat(caducidad))
        ws.cell(row=fila, column=7, value=lote)
        nuevas[fila] = zona
        fila += 1

    filas = []
    catalogo = catalogo_alergenos()
    for r in range(4, ws.max_row + 1):
        descripcion = ws.cell(row=r, column=1).value
        if descripcion in (None, ""):
            continue
        filas.append({
            "fila": r,
            "descripcion": str(descripcion).strip(),
            "lote": ws.cell(row=r, column=7).value,
            "observaciones": ws.cell(row=r, column=9).value,
            "es_nueva": r in nuevas,
        })

    # ---------- reparto de materia prima por zona ----------
    grupos_alergeno, grupos_no = {}, {}
    for f in filas:
        if f["es_nueva"]:
            continue
        alergeno, _ = clasificar(f["descripcion"], catalogo)
        destino = grupos_no if alergeno.strip().lower() in NO_ALERGENO else grupos_alergeno
        destino.setdefault(normalizar(f["descripcion"]), []).append(f)

    nombres_a = sorted(grupos_alergeno)
    nombres_na = sorted(grupos_no)
    random.shuffle(nombres_a)
    random.shuffle(nombres_na)

    racks_mpa = [c for c in ubicaciones.CODIGOS if c.startswith("MPA-A")]
    racks_mpna = [c for c in ubicaciones.CODIGOS if c.startswith("MPNA-")]
    asignaciones = repartir(racks_mpa, [grupos_alergeno[n] for n in nombres_a])
    asignaciones.update(repartir(racks_mpna, [grupos_no[n] for n in nombres_na]))

    # ---------- empaque y PT: un rack por zona, en orden ----------
    punteros = {}
    for f in filas:
        if not f["es_nueva"]:
            continue
        zona = nuevas[f["fila"]]
        recursos = [c for c in ubicaciones.CODIGOS if c.startswith(zona)]
        i = punteros.get(zona, 0)
        asignaciones[f["fila"]] = recursos[i % len(recursos)]
        punteros[zona] = i + 1

    # ---------- casos especiales ----------
    con_ubicacion = [f for f in filas if f["fila"] in asignaciones]

    def apartar(predicado):
        for f in con_ubicacion:
            if f["fila"] in asignaciones and predicado(f):
                del asignaciones[f["fila"]]
                return f
        return None

    lacteo = apartar(lambda f: "LECHE" in normalizar(f["descripcion"]))
    if lacteo:
        asignaciones[lacteo["fila"]] = "CUARENTENA-MP"
    premix = apartar(lambda f: "PREMIX" in normalizar(f["descripcion"]))
    if premix:
        asignaciones[premix["fila"]] = "LAB-RETENCION"
    rechazado = apartar(lambda f: "RESVERAT" in normalizar(f["descripcion"]))
    if rechazado:
        asignaciones[rechazado["fila"]] = "JAULA-MP"
    alergeno = apartar(lambda f: normalizar(f["descripcion"]) in grupos_alergeno)
    if alergeno:
        asignaciones[alergeno["fila"]] = "PASILLO-A"

    # ---------- cola de trabajo: 2 no encontrados + 5 recién llegados ----------
    candidatos = [f for f in filas if f["fila"] in asignaciones
                  and ubicaciones.CATALOGO[asignaciones[f["fila"]]]["tipo"] == "RACK"]
    random.shuffle(candidatos)
    no_encontrados = []
    for f in candidatos[:2]:
        no_encontrados.append({"fila": f["fila"], "antes": asignaciones[f["fila"]]})
        del asignaciones[f["fila"]]

    sin_ubicar = []
    excluidos = {n["fila"] for n in no_encontrados}

    def suelta(f):
        return (not f["es_nueva"] and f["fila"] in asignaciones
                and f["fila"] not in excluidos
                and ubicaciones.CATALOGO[asignaciones[f["fila"]]]["tipo"] == "RACK")

    candidatos_alta = [f for f in filas if suelta(f)
                       and "ALTA" in normalizar(f.get("observaciones"))]
    resto = [f for f in filas if suelta(f) and f not in candidatos_alta]
    random.shuffle(resto)
    for f in (candidatos_alta + resto)[:5]:
        del asignaciones[f["fila"]]
        sin_ubicar.append(f)

    # ---------- escribir ubicaciones ----------
    for f in filas:
        celda = ws.cell(row=f["fila"], column=col_ubicacion, value=asignaciones.get(f["fila"]))
        celda.alignment = Alignment(horizontal="left")
        celda.font = Font(size=9)

    aviso = wb.create_sheet("DEMO", 0)
    aviso["A1"] = "DATOS DE DEMOSTRACIÓN — NO ES UN REGISTRO DE CALIDAD"
    aviso["A1"].font = Font(bold=True, size=14, color="B01C1C")
    aviso["A2"] = ("Las ubicaciones y los movimientos de este archivo se generaron para mostrar "
                   "el sistema. No describen el almacén real.")
    aviso["A3"] = f"Semilla {SEMILLA} · generado el {HOY.isoformat()}"
    aviso.column_dimensions["A"].width = 100
    aviso.sheet_properties.tabColor = "B01C1C"
    wb.save(DEMO_XLSX)

    # ---------- bitácora ----------
    if DEMO_DB.exists():
        DEMO_DB.unlink()
    con = sqlite3.connect(DEMO_DB)
    con.executescript("""
    CREATE TABLE movimiento (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL,
        accion TEXT NOT NULL, fila INTEGER, descripcion TEXT NOT NULL, lote TEXT,
        ubicacion_origen TEXT, ubicacion_destino TEXT, usuario TEXT NOT NULL, nota TEXT,
        motivo TEXT, ref_doc TEXT);
    """)

    eventos = []
    for f in filas:
        final = asignaciones.get(f["fila"])
        if not final:
            continue
        inicio = dt.datetime.combine(HOY - dt.timedelta(days=random.randint(8, 14)),
                                     dt.time(random.randint(6, 11), random.randint(0, 59)))
        zona = ubicaciones.zona_de(final)
        mudanza = random.random() < 0.35 and zona in ("MP-ALERGENO", "MP-NO-ALERGENO", "EM")
        if mudanza:
            recursos = [c for c in ubicaciones.CODIGOS if ubicaciones.zona_de(c) == zona]
            previo = random.choice([r for r in recursos if r != final])
            eventos.append((inicio, "UBICAR", f, "ANDEN", previo, "recepción y ubicación inicial"))
            eventos.append((inicio + dt.timedelta(days=random.randint(1, 5),
                                                  hours=random.randint(1, 5)),
                            "MOVER", f, previo, final,
                            random.choice(["reacomodo por espacio en el pasillo",
                                           "se juntan los lotes del mismo material",
                                           "se recorre para liberar rack de surtido",
                                           "queda el resto del surtido y se reacomoda"])))
            if random.random() < 0.4:
                eventos.append((inicio + dt.timedelta(days=random.randint(6, 8),
                                                      hours=random.randint(1, 3)),
                                "CONFIRMAR", f, final, final, "conteo cíclico de zona"))
        else:
            eventos.append((inicio, "UBICAR", f, "ANDEN", final, "recepción y ubicación inicial"))
            if random.random() < 0.5:
                eventos.append((dt.datetime.combine(
                    HOY - dt.timedelta(days=random.randint(0, 3)),
                    dt.time(random.randint(6, 13), random.randint(0, 59))),
                    "CONFIRMAR", f, final, final, "verificación en piso"))

    for n in no_encontrados:
        f = next(x for x in filas if x["fila"] == n["fila"])
        eventos.append((dt.datetime.combine(HOY - dt.timedelta(days=1),
                                            dt.time(random.randint(8, 12), 20)),
                        "NO_ENCONTRADO", f, n["antes"], None,
                        "no estaba en el rack rotulado; se busca en el pasillo"))

    for f in sin_ubicar:
        eventos.append((dt.datetime.combine(HOY - dt.timedelta(days=random.randint(0, 2)),
                                            dt.time(random.randint(6, 12), random.randint(0, 59))),
                        "IMPORTAR", f, None, None,
                        "alta de recepción; pendiente de asignar rack"))

    eventos.sort(key=lambda e: e[0])

    def motivo_de(accion, nota):
        texto = (nota or "").lower()
        if accion == "UBICAR":
            return None, "F01-ALM-PNO-001"
        if accion == "MOVER":
            if "liberar" in texto:
                return "liberacion", None
            if "surtido" in texto:
                return "surtido", None
            return "reacomodo", None
        return None, None

    for ts, accion, f, origen, destino, nota in eventos:
        motivo, ref = motivo_de(accion, nota)
        con.execute(
            "INSERT INTO movimiento (ts, accion, fila, descripcion, lote, ubicacion_origen, "
            "ubicacion_destino, usuario, nota, motivo, ref_doc) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (ts.isoformat(sep=" ", timespec="seconds"), accion, f["fila"], f["descripcion"],
             f["lote"], origen, destino, random.choice(USUARIOS), nota, motivo, ref))
    con.commit()

    total_mov = con.execute("SELECT COUNT(*) FROM movimiento").fetchone()[0]
    acciones = dict(con.execute(
        "SELECT accion, COUNT(*) FROM movimiento GROUP BY accion ORDER BY 2 DESC").fetchall())
    con.close()

    por_zona = {}
    for f in filas:
        codigo = asignaciones.get(f["fila"])
        if codigo:
            zona = ubicaciones.zona_de(codigo)
            por_zona[zona] = por_zona.get(zona, 0) + 1
    ocupadas = sum(1 for f in filas if asignaciones.get(f["fila"]))

    print(f"demo         : {DEMO_XLSX}")
    print(f"bitácora     : {DEMO_DB}")
    print(f"lotes        : {len(filas)}")
    print(f"con ubicación: {ocupadas} ({100 * ocupadas / len(filas):.1f} %)")
    print(f"sin ubicar   : {len(filas) - ocupadas}")
    for zona, n in sorted(por_zona.items(), key=lambda kv: -kv[1]):
        print(f"  {zona:<16}{n:>4}")
    print(f"movimientos  : {total_mov} {acciones}")
    return 0


if __name__ == "__main__":
    raise SystemExit(sembrar())
