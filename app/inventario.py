"""El inventario vive en el Excel que ya usa el área; aquí solo se le agrega la columna UBICACIÓN."""
import datetime as dt
import ast
import difflib
import io
import math
import operator
import pathlib
import re
import unicodedata
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter

from app import fechas, inventarios, ubicaciones

RAIZ = pathlib.Path(__file__).resolve().parents[1]
RUTA = RAIZ / "data" / "Inventario_MP.xlsx"
FILA_ENCABEZADO = 3
FILA_DATOS = 4
ENCABEZADO_UBICACION = "UBICACIÓN"
ENCABEZADO_QR = "QR_ID"

# encabezado del Excel del área  ->  nombre interno
MAPA_COLUMNAS = {
    "descripcion": "descripcion",
    "producto": "descripcion",
    "material": "descripcion",
    "nombre del material": "descripcion",
    "presentacion": "presentacion",
    "piezas por caja": "piezas",
    "piezas": "piezas",
    "sku": "erp",
    "clave erp": "erp",
    "categoria": "categoria",
    "proveedor": "proveedor",
    "unidad de medida": "unidad",
    "unidad": "unidad",
    "conteo": "cantidad",
    "cantidad": "cantidad",
    "stock": "cantidad",
    "existencia": "cantidad",
    "um": "unidad",
    "caducidad vigente": "caducidad",
    "caducidad": "caducidad",
    "lote actual": "lote",
    "lote": "lote",
    "entrada erp": "erp",
    "observaciones": "observaciones",
    "ubicacion": "ubicacion",
    # etiqueta física opcional (QR o código de barras): no cambia cómo se identifica un lote
    "qr id": "qr_id",
    "qr": "qr_id",
    "barcode id": "qr_id",
    "barcode": "qr_id",
    "codigo qr": "qr_id",
    "codigo de barra": "qr_id",
    "codigo de barras": "qr_id",
}


UNIDADES = {"pz": "pza", "pzas": "pza", "pieza": "pza", "piezas": "pza", "pza": "pza",
            "kg": "kg", "kgs": "kg", "kilogramo": "kg", "kilogramos": "kg",
            "lt": "lt", "l": "lt", "litro": "lt", "litros": "lt",
            "m": "m", "ml": "ml", "g": "g", "caja": "caja", "millar": "millar"}


def normalizar_unidad(valor, defecto="kg"):
    """'PZ', 'Piezas', 'KG ' -> unidad canónica, para no sumar peras con manzanas."""
    clave = str(valor or "").strip().lower()
    return UNIDADES.get(clave, clave or defecto)


def normalizar_qr(texto):
    """Código de etiqueta tal como viene (sin quitarle guiones ni cambiarle el caso).

    No se usa `normalizar()` porque borra la puntuación: 7501-055 y 7501055 son códigos
    distintos y no pueden confundirse.
    """
    return str(texto or "").strip()


def clave_qr(texto):
    """Forma de comparación: sin espacios sobrantes y sin distinguir mayúsculas."""
    return normalizar_qr(texto).upper()


def normalizar(texto):
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9]+", " ", s)).strip().upper()


def _parecido(termino, texto):
    """Qué tan parecido es lo que escribió la persona a lo que hay en el inventario (0 a 1).

    Tolera letras cambiadas o faltantes: «protiena» encuentra «PROTEINA» y «lactsoa»
    encuentra «LACTOSA». No sustituye a la búsqueda exacta: solo se usa cuando esa
    no encontró nada, para ofrecer resultados aproximados y sugerencias.
    """
    if not termino or not texto:
        return 0.0
    if termino in texto:
        return 1.0
    mejor = 0.0
    palabras = texto.split()
    if any(palabra.startswith(termino) for palabra in palabras):
        mejor = 0.92
    if len(termino) < 3 or mejor >= 0.92:
        return mejor
    largo = len(termino)
    for palabra in palabras:
        razon = difflib.SequenceMatcher(None, termino, palabra).ratio()
        if razon > mejor:
            mejor = razon
        if len(palabra) > largo:
            # comparar también contra cada tramo del mismo largo (errores en medio)
            for inicio in range(0, len(palabra) - largo + 1):
                tramo = palabra[inicio:inicio + largo]
                razon = difflib.SequenceMatcher(None, termino, tramo).ratio()
                if razon > mejor:
                    mejor = razon
    return mejor


# Celdas que Excel guarda como TEXTO aunque sean números ("1458", "1,458.50").
# Solo se aceptan formas inequívocas: "1.458" en español puede ser mil cuatrocientos
# cincuenta y ocho o uno punto cuatro cinco ocho, así que NO se adivina: se reporta.
_NUMERO = re.compile(r"^-?\d+(?:\.\d+)?$")
_MILES = re.compile(r"^-?\d{1,3}(?:,\d{3})+(?:\.\d+)?$")


def numero(valor):
    """float si la celda es un número (o un texto numérico sin ambigüedad); si no, None."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(float(valor)) else None
    if not isinstance(valor, str):
        return None
    texto = valor.strip().replace(" ", "")
    if not texto:
        return None
    if _NUMERO.match(texto):
        # "1.458" es ambiguo en español: mil cuatrocientos cincuenta y ocho o uno punto
        # cuatro cinco ocho. Con tres decimales y sin coma de miles no se adivina.
        if "." in texto and len(texto.rsplit(".", 1)[1]) == 3:
            return None
        return float(texto)
    if _MILES.match(texto):
        return float(texto.replace(",", ""))
    return None


_UNIDADES_ESCRITAS = {
    "kg": "kg", "kgs": "kg", "kilo": "kg", "kilos": "kg",
    "kilogramo": "kg", "kilogramos": "kg",
    "g": "g", "gr": "g", "gramo": "g", "gramos": "g",
    "pza": "pza", "pzas": "pza", "pz": "pza", "pieza": "pza", "piezas": "pza",
    "lt": "lt", "lts": "lt", "l": "lt", "litro": "lt", "litros": "lt",
    "caja": "caja", "cajas": "caja", "millar": "millar",
}
_PARTES_CANTIDAD = re.compile(r"^(-?[\d.,]+)([a-záéíóúüñ]+)?$")


def _dividir_cantidad(texto):
    """'5 kg' -> ('5', 'kg'); '1458' -> ('1458', ''); 'basura' -> (None, None)."""
    crudo = re.sub(r"\s+", "", str(texto).strip().lower())
    coincidencia = _PARTES_CANTIDAD.match(crudo)
    if not coincidencia:
        return None, None
    return coincidencia.group(1), coincidencia.group(2) or ""


def unidad_escrita(texto):
    """Familia de la unidad que escribió la persona ('' si no escribió ninguna)."""
    _, escrita = _dividir_cantidad(texto)
    if not escrita:
        return ""
    return _UNIDADES_ESCRITAS.get(escrita, escrita)


def _candidatos_numero(texto):
    """Interpretaciones posibles de '1,458.50', '1.458,50', '1458,5', '1.458', '1,458'…"""
    if "," in texto and "." in texto:
        # el último separador que aparece es el decimal
        if texto.rfind(",") > texto.rfind("."):
            return [texto.replace(".", "").replace(",", ".")]
        return [texto.replace(",", "")]
    separador = "," if "," in texto else ("." if "." in texto else "")
    if not separador:
        return [texto]
    partes = texto.split(separador)
    decimal = "".join(partes[:-1]) + "." + partes[-1]
    candidatos = [decimal]
    # grupos de 3: interpretación de miles al estilo mexicano (1,458 = mil 458)
    if len(partes) > 1 and all(len(parte) == 3 for parte in partes[1:]):
        candidatos.append("".join(partes))
    return candidatos


def interpretar_cantidad(texto, total=None, unidad=None):
    """Lo que escribe la persona -> número. Para capturar cantidades a mano.

    Acepta «5», «5,5», «5.5», «5 kg», «1 458», «1,458.50», «1.458,50».
    Con separador y grupos de 3 dígitos se interpreta como miles (uso mexicano),
    salvo que se pase del total: ahí se usa la otra lectura si cabe.

    Si se escribe una unidad y no coincide con la del renglón (por ejemplo «5 g» en un
    lote de kg) devuelve None: es preferible pedir que la corrijan a mover 5 kg por 5 g.
    """
    if texto is None or isinstance(texto, bool):
        return None
    if isinstance(texto, (int, float)):
        valor = float(texto)
        return valor if math.isfinite(valor) else None
    crudo, escrita = _dividir_cantidad(texto)
    if crudo is None:
        return None
    if escrita:
        familia = _UNIDADES_ESCRITAS.get(escrita)
        if familia is None:
            return None
        if unidad and familia != str(unidad).strip().lower():
            return None
    # formas aceptadas: entero, con separador decimal, o con grupos de miles
    if not crudo or not re.fullmatch(
            r"-?(?:\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?|\d+(?:[.,]\d+)?)", crudo):
        return None

    candidatos = []
    for parte in _candidatos_numero(crudo):
        try:
            valor = float(parte)
        except ValueError:
            continue
        if math.isfinite(valor):
            candidatos.append(valor)
    if not candidatos:
        return None

    elegido = candidatos[0]
    if len(candidatos) > 1 and all(len(parte) == 3 for parte in re.split(r"[.,]", crudo)[1:]):
        # grupos de 3: se lee como miles (1,458 / 1.458 / 1,234,567 => miles)
        elegido = candidatos[-1]
    if total is not None and elegido > total + 1e-9:
        for alterno in candidatos:
            if alterno <= total + 1e-9:
                elegido = alterno
                break
    return elegido


_OPS = {ast.Add: operator.add, ast.Sub: operator.sub,
        ast.Mult: operator.mul, ast.Div: operator.truediv}


def evaluar_formula(texto):
    """Evalúa una suma/resta aritmética simple. Excel la muestra, pero el archivo
    puede no traer el valor calculado (openpyxl lo guarda sin caché)."""
    nodo = ast.parse(str(texto).lstrip("="), mode="eval").body

    def ev(n):
        if isinstance(n, ast.BinOp) and type(n.op) in _OPS:
            return _OPS[type(n.op)](ev(n.left), ev(n.right))
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, ast.USub):
            return -ev(n.operand)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return n.value
        raise ValueError("fórmula no soportada")

    return round(float(ev(nodo)), 3)


class Inventario:
    """Carga el libro del área, expone sus filas y escribe la columna UBICACIÓN."""

    def __init__(self, ruta=None, tipo="mp", con=None):
        self.tipo = inventarios.normalizar_tipo(tipo)
        self.ruta = pathlib.Path(ruta or inventarios.ruta(self.tipo))
        self.fila_encabezado = FILA_ENCABEZADO
        self.nombre = inventarios.TIPOS[self.tipo]["nombre"]
        self.con = con
        self.en_db = bool(con is not None and getattr(con, "es_postgres", False))
        if self.en_db:
            self.wb = None
            self.ws = None
            self.columnas = {}
            self.filas = self._leer_db()
            return
        if not self.ruta.exists():
            raise FileNotFoundError(f"no existe el inventario: {self.ruta}")
        self.wb = openpyxl.load_workbook(io.BytesIO(self.ruta.read_bytes()))
        self._fijar_imagenes()
        self.ws = self._hoja_inventario()
        self.columnas = self._detectar_columnas()
        self.filas = self._leer()

    # ---------- estructura ----------
    def _fijar_imagenes(self):
        """openpyxl reutiliza el buffer de las imágenes del libro y la segunda escritura
        falla con 'I/O operation on closed file'. Se congelan los bytes al cargar."""
        for hoja in self.wb.worksheets:
            for imagen in getattr(hoja, "_images", []):
                datos = imagen._data()

                def congelar(d):
                    return lambda: d

                imagen._data = congelar(datos)

    def _hoja_inventario(self):
        """La hoja del inventario es la que trae el encabezado del material, no siempre la primera.

        Acepta 'Descripción', 'Producto' o 'Material' y busca en las primeras filas, porque las
        tablas de empaque y producto terminado no siempre repiten el formato del archivo de MP.
        """
        validos = {"DESCRIPCION", "PRODUCTO", "MATERIAL", "NOMBRE DEL MATERIAL"}
        for hoja in self.wb.worksheets:
            for fila in range(1, min(hoja.max_row, 8) + 1):
                for col in range(1, min(hoja.max_column, 15) + 1):
                    if normalizar(hoja.cell(row=fila, column=col).value) in validos:
                        self.fila_encabezado = fila
                        return hoja
        return self.wb.worksheets[0]

    def _detectar_columnas(self):
        columnas = {}
        for col in range(1, self.ws.max_column + 1):
            valor = self.ws.cell(row=self.fila_encabezado, column=col).value
            if not valor:
                continue
            clave = MAPA_COLUMNAS.get(normalizar(valor).lower())
            if clave and clave not in columnas:
                columnas[clave] = col
        if "descripcion" not in columnas:
            raise ValueError(f"no se encontró la columna de descripción en la fila {self.fila_encabezado}")
        if "ubicacion" not in columnas:
            col = max(columnas.values()) + 1
            self._crear_columna_ubicacion(col)
            columnas["ubicacion"] = col
        return columnas

    def _crear_columna_ubicacion(self, col):
        origen = self.ws.cell(row=self.fila_encabezado, column=col - 1)
        celda = self.ws.cell(row=self.fila_encabezado, column=col, value=ENCABEZADO_UBICACION)
        if origen.has_style:
            celda._style = copy(origen._style)
        celda.value = ENCABEZADO_UBICACION
        self.ws.column_dimensions[get_column_letter(col)].width = 16
        for fila in range(self.fila_encabezado + 1, self.ws.max_row + 1):
            modelo = self.ws.cell(row=fila, column=col - 1)
            nueva = self.ws.cell(row=fila, column=col)
            if modelo.has_style:
                nueva._style = copy(modelo._style)

    def _asegurar_columna_qr(self):
        """Crea la columna QR_ID al final si el libro todavía no la tiene. Devuelve su número."""
        if "qr_id" in self.columnas:
            return self.columnas["qr_id"]
        col = self.ws.max_column + 1
        origen = self.ws.cell(row=self.fila_encabezado, column=col - 1)
        celda = self.ws.cell(row=self.fila_encabezado, column=col)
        celda.value = ENCABEZADO_QR
        if origen.has_style:
            celda._style = copy(origen._style)
        self.ws.column_dimensions[get_column_letter(col)].width = 22
        for fila in range(self.fila_encabezado + 1, self.ws.max_row + 1):
            modelo = self.ws.cell(row=fila, column=col - 1)
            nueva = self.ws.cell(row=fila, column=col)
            if modelo.has_style:
                nueva._style = copy(modelo._style)
        self.columnas["qr_id"] = col
        return col

    def _leer(self):
        filas, ocurrencias = [], {}
        for fila in range(self.fila_encabezado + 1, self.ws.max_row + 1):
            descripcion = self.ws.cell(row=fila, column=self.columnas["descripcion"]).value
            if descripcion in (None, ""):
                continue
            crudo = self._valor(fila, "caducidad")
            iso, motivo = fechas.parsear(crudo)
            lote = self._valor(fila, "lote")
            cantidad_cruda = self._valor(fila, "cantidad")
            cantidad = numero(cantidad_cruda)
            ubicacion_cruda = self._ubicacion_celda(fila)
            ubicacion = ubicaciones.canonico(ubicacion_cruda)
            # Una UBICACIÓN escrita a mano que no está en el catálogo NO es una ubicación:
            # se cuenta como sin ubicar y se marca para que la pantalla avise. El catálogo
            # es la única fuente de verdad y el Excel no puede inventar códigos.
            ubicacion_invalida = (str(ubicacion_cruda).strip()
                                  if ubicacion_cruda not in (None, "") and ubicacion is None
                                  else None)
            registro = {
                "fila": fila,
                "descripcion": str(descripcion).strip(),
                "categoria": self._valor(fila, "categoria"),
                "proveedor": self._valor(fila, "proveedor"),
                "unidad": normalizar_unidad(self._valor(fila, "unidad"),
                                            inventarios.TIPOS[self.tipo]["unidad"]),
                # si la cantidad no es un número, se conserva el texto para poder mostrarlo
                "cantidad": cantidad if cantidad is not None else cantidad_cruda,
                "cantidad_legible": cantidad is not None,
                "caducidad": iso,
                "caducidad_cruda": (crudo.date().isoformat() if isinstance(crudo, dt.datetime)
                                    else str(crudo) if crudo not in (None, "") else None),
                "motivo_caducidad": motivo,
                "lote": str(lote).strip() if lote not in (None, "") else None,
                "erp": self._valor(fila, "erp"),
                "qr_id": normalizar_qr(self._valor(fila, "qr_id")) or None,
                "observaciones": self._valor(fila, "observaciones"),
                "ubicacion": ubicacion,
                "ubicacion_cruda": (str(ubicacion_cruda).strip()
                                    if ubicacion_cruda not in (None, "") else None),
                "ubicacion_invalida": ubicacion_invalida,
                "cantidad_formula": self._es_formula(fila, "cantidad"),
                "presentacion": self._valor(fila, "presentacion"),
                "piezas": self._valor(fila, "piezas"),
                "tipo": self.tipo,
            }
            dias, sem = fechas.semaforo(iso)
            registro["dias_restantes"] = dias
            registro["semaforo"] = sem
            clave = (normalizar(registro["descripcion"]), normalizar(registro["lote"]))
            ocurrencias[clave] = ocurrencias.get(clave, 0) + 1
            registro["clave"] = f"{clave[0]}|{clave[1]}|{ocurrencias[clave]}"
            filas.append(registro)
        self._indice = None
        return filas

    def _leer_db(self):
        """Carga los renglones desde la tabla PostgreSQL `inventario_lote` (Supabase)."""
        filas, ocurrencias = [], {}
        registros_db = self.con.execute(
            "SELECT * FROM inventario_lote WHERE inventario=? ORDER BY fila ASC",
            (self.tipo,)).fetchall()
        for r in registros_db:
            crudo = r.get("caducidad_cruda") or r.get("caducidad")
            iso = r.get("caducidad")
            motivo = r.get("motivo_caducidad")
            if not iso and crudo:
                iso, motivo = fechas.parsear(crudo)
            ubicacion_cruda = r.get("ubicacion_cruda") or r.get("ubicacion")
            ubicacion = ubicaciones.canonico(ubicacion_cruda)
            ubicacion_invalida = r.get("ubicacion_invalida") or (
                str(ubicacion_cruda).strip()
                if ubicacion_cruda not in (None, "") and ubicacion is None else None
            )
            legible = bool(r.get("cantidad_legible", True)) and (r.get("cantidad") is not None)
            cantidad_val = float(r["cantidad"]) if legible else (
                r.get("cantidad_texto") if r.get("cantidad_texto") is not None else r.get("cantidad")
            )
            lote = r.get("lote")
            registro = {
                "fila": int(r["fila"]),
                "descripcion": str(r["descripcion"]).strip(),
                "categoria": r.get("categoria"),
                "proveedor": r.get("proveedor"),
                "unidad": normalizar_unidad(r.get("unidad"),
                                            inventarios.TIPOS[self.tipo]["unidad"]),
                "cantidad": cantidad_val,
                "cantidad_legible": legible,
                "caducidad": iso,
                "caducidad_cruda": str(crudo) if crudo not in (None, "") else None,
                "motivo_caducidad": motivo,
                "lote": str(lote).strip() if lote not in (None, "") else None,
                "erp": r.get("erp"),
                "qr_id": normalizar_qr(r.get("qr_id")) or None,
                "observaciones": r.get("observaciones"),
                "ubicacion": ubicacion,
                "ubicacion_cruda": (str(ubicacion_cruda).strip()
                                    if ubicacion_cruda not in (None, "") else None),
                "ubicacion_invalida": ubicacion_invalida,
                "cantidad_formula": bool(r.get("cantidad_formula", False)),
                "presentacion": r.get("presentacion"),
                "piezas": r.get("piezas"),
                "tipo": self.tipo,
            }
            dias, sem = fechas.semaforo(iso)
            registro["dias_restantes"] = dias
            registro["semaforo"] = sem
            clave = (normalizar(registro["descripcion"]), normalizar(registro["lote"]))
            ocurrencias[clave] = ocurrencias.get(clave, 0) + 1
            registro["clave"] = f"{clave[0]}|{clave[1]}|{ocurrencias[clave]}"
            filas.append(registro)
        self._indice = None
        return filas

    def _valor(self, fila, campo):
        col = self.columnas.get(campo)
        if not col:
            return None
        valor = self.ws.cell(row=fila, column=col).value
        if isinstance(valor, str):
            valor = valor.strip()
            if valor.startswith("=") and campo == "cantidad":
                try:
                    return evaluar_formula(valor)
                except (ValueError, SyntaxError, ZeroDivisionError, RecursionError):
                    return None
        return valor if valor != "" else None

    def _es_formula(self, fila, campo):
        col = self.columnas.get(campo)
        if not col:
            return False
        valor = self.ws.cell(row=fila, column=col).value
        return isinstance(valor, str) and valor.startswith("=")

    def _ubicacion_celda(self, fila):
        """Lo que dice la celda UBICACIÓN, tal cual (sin interpretarlo)."""
        col = self.columnas["ubicacion"]
        valor = self.ws.cell(row=fila, column=col).value
        if valor in (None, ""):
            return None
        return str(valor).strip()

    # ---------- consultas ----------
    def _indice_busqueda(self):
        """Precalcula el texto de búsqueda de cada renglón (con etiquetas de campo).

        Se arma una vez y se reutiliza: con 5,000 renglones, normalizar en cada tecla se siente
        lento. Incluye dos versiones: con espacios y sin espacios, para que «BL-2026-01»,
        «bl 2026 01» y «bl202601» encuentren el mismo lote.
        """
        if getattr(self, "_indice", None) is not None:
            return self._indice
        indice = {}
        for f in self.filas:
            campos = {
                "material": normalizar(f["descripcion"]),
                "lote": normalizar(f["lote"]),
                "erp": normalizar(f["erp"]),
                "proveedor": normalizar(f["proveedor"]),
                "categoria": normalizar(f["categoria"]),
                "ubicacion": normalizar(f["ubicacion"]),
                "qr": normalizar(f.get("qr_id")),
            }
            texto = " ".join([
                f"MATERIAL {campos['material']}", f"LOTE {campos['lote']}",
                f"PROVEEDOR {campos['proveedor']}", f"CATEGORIA {campos['categoria']}",
                f"ERP {campos['erp']}", f"UBICACION {campos['ubicacion']}",
                f"QR {campos['qr']}",
            ]).strip()
            indice[f["fila"]] = {"texto": texto, "compacto": texto.replace(" ", ""),
                                 **campos,
                                 **{f"{k}_c": v.replace(" ", "") for k, v in campos.items()}}
        self._indice = indice
        return indice

    @staticmethod
    def _puntaje(token, datos):
        """Cuánto se parece el renglón al término. 0 = no coincide.

        Prioriza el lote: si alguien teclea un lote, ese lote va primero, aunque otro material
        lo tenga escrito en la descripción.
        """
        campo, _, valor = token.partition(":")
        if campo in ("lote", "erp", "material", "proveedor", "ubicacion", "categoria") and valor:
            exacto = datos.get(campo) == valor or datos.get(f"{campo}_c") == valor.replace(" ", "")
            return 100 if exacto else (60 if valor in datos.get(campo, "") else 0)

        compacto = token.replace(" ", "")
        if not (token in datos["texto"] or (compacto and compacto in datos["compacto"])):
            return 0

        puntos = 1
        lote = datos.get("lote", "")
        if lote and token == lote:
            puntos += 100
        elif lote and (lote.startswith(token) or (compacto and datos.get("lote_c", "").startswith(compacto))):
            puntos += 60
        elif lote and (token in lote or (compacto and compacto in datos.get("lote_c", ""))):
            puntos += 45
        erp = datos.get("erp", "")
        if erp and (token in erp or (compacto and compacto in datos.get("erp_c", ""))):
            puntos += 25
        material = datos.get("material", "")
        if material and (token in material or (compacto and compacto in datos.get("material_c", ""))):
            puntos += 20
            if material.startswith(token):
                puntos += 10
        if datos.get("ubicacion") and (token in datos["ubicacion"]
                                      or (compacto and compacto in datos.get("ubicacion_c", ""))):
            puntos += 15
        if datos.get("proveedor") and token in datos["proveedor"]:
            puntos += 5
        return puntos

    def buscar(self, texto="", solo_sin_ubicar=False, solo_vencidos=False, incluir_cero=False,
               solo_ilegibles=False, difuso=False, limite=None):
        """Búsqueda por palabras: todas deben aparecer, en cualquier orden y en cualquier campo.

        Acepta «lote:», «erp:», «material:», «proveedor:», «ubicacion:» para apuntar a un campo,
        y encuentra el lote con o sin guiones («BL-2026-01» = «bl202601»).

        `difuso=True` se usa cuando la búsqueda exacta no encontró nada: compara por
        parecido (tolera letras cambiadas o faltantes) y devuelve lo más probable primero.
        `solo_ilegibles` muestra los renglones cuya CANTIDAD no es un número.
        """
        tokens = [t for t in (normalizar(p) for p in str(texto or "").split()) if t]
        indice = self._indice_busqueda() if (tokens or difuso) else {}
        filas = []
        for f in self.filas:
            if solo_ilegibles:
                if f["cantidad_legible"]:
                    continue
            elif not incluir_cero and not self._tiene_existencia(f):
                continue
            if solo_sin_ubicar and f["ubicacion"]:
                continue
            if solo_vencidos and f["semaforo"] != "VENCIDO":
                continue
            filas.append(f)

        if not tokens:
            return filas[:limite] if limite else filas

        if difuso:
            encontrados = []
            for f in filas:
                datos = indice[f["fila"]]
                puntajes = [_parecido(t, datos["texto"]) for t in tokens]
                if puntajes and min(puntajes) >= 0.68:
                    encontrados.append((sum(puntajes), f))
            encontrados.sort(key=lambda par: (-par[0], normalizar(par[1]["descripcion"]),
                                              str(par[1]["lote"] or "")))
            if limite:
                encontrados = encontrados[:limite]
            return [f for _, f in encontrados]

        salida = []
        for f in filas:
            datos = indice[f["fila"]]
            puntajes = [self._puntaje(t, datos) for t in tokens]
            if not all(p > 0 for p in puntajes):
                continue
            salida.append((sum(puntajes), f))
        salida.sort(key=lambda par: (-par[0], normalizar(par[1]["descripcion"]),
                                     str(par[1]["lote"] or "")))
        if limite:
            salida = salida[:limite]
        return [f for _, f in salida]

    def sugerencias(self, texto, limite=5):
        """Lotes y materiales parecidos, para cuando la búsqueda no encuentra nada.

        Se comparan los mismos campos que ve la búsqueda (material, lote, categoría,
        proveedor, código del ERP) con el criterio tolerante de la búsqueda aproximada,
        y se devuelve el texto **tal como está escrito** en el Excel. Así la sugerencia,
        al darle clic, siempre encuentra algo: nunca se inventa un nombre.
        """
        termino = normalizar(texto)
        if not termino:
            return []
        candidatos = {}
        for f in self.filas:
            for campo in ("lote", "descripcion", "categoria", "proveedor", "erp"):
                valor = str(f.get(campo) or "").strip()
                if valor:
                    candidatos.setdefault(normalizar(valor), valor)
        puntajes = [(_parecido(termino, clave), clave, valor)
                    for clave, valor in candidatos.items()]
        # 0.62 es un poco más permisivo que la búsqueda aproximada (0.68): aquí solo se
        # propone, la persona decide. Lo que no se parece en nada no se propone.
        buenos = sorted((p for p in puntajes if p[0] >= 0.62), key=lambda p: (-p[0], p[1]))
        return [valor for _, _, valor in buenos[:limite]]

    def por_fila(self, fila):
        for registro in self.filas:
            if registro["fila"] == fila:
                return registro
        raise KeyError(f"fila inexistente: {fila}")

    def buscar_por_qr(self, codigo):
        """Renglón ACTIVO de este libro con ese qr_id (o None)."""
        buscado = clave_qr(codigo)
        if not buscado:
            return None
        for registro in self.filas:
            if (registro.get("qr_id") and clave_qr(registro["qr_id"]) == buscado
                    and self._tiene_existencia(registro)):
                return registro
        return None

    def asignar_qr(self, fila, codigo):
        """Escribe (o borra, con None) el qr_id de un renglón. Devuelve (anterior, nuevo)."""
        registro = self.por_fila(fila)
        anterior = registro.get("qr_id")
        nuevo = normalizar_qr(codigo) or None
        if self.en_db:
            self.con.execute(
                "UPDATE inventario_lote SET qr_id=?, actualizado=NOW() "
                "WHERE inventario=? AND fila=?",
                (nuevo, self.tipo, int(fila)))
        else:
            col = self._asegurar_columna_qr()
            celda = self.ws.cell(row=fila, column=col)
            celda.value = nuevo            # asignar .value: con value=None en cell() NO borra
        registro["qr_id"] = nuevo
        self._indice = None
        return anterior, nuevo

    @staticmethod
    def _tiene_existencia(registro):
        cantidad = registro.get("cantidad")
        return isinstance(cantidad, (int, float)) and cantidad > 0

    def activas(self):
        """Renglones con existencia mayor que cero (los anulados quedan fuera)."""
        return [f for f in self.filas if self._tiene_existencia(f)]

    def resumen(self):
        activas = self.activas()
        total = len(activas)
        con = [f for f in activas if f["ubicacion"]]
        semaforo = {"VENCIDO": 0, "CRITICO": 0, "AVISO": 0, "OK": 0, "SIN DATO": 0}
        for f in activas:
            semaforo[f["semaforo"]] += 1
        por_unidad = {}
        for f in activas:
            if isinstance(f["cantidad"], (int, float)):
                unidad = (f["unidad"] or "kg").strip().lower() or "kg"
                por_unidad[unidad] = round(por_unidad.get(unidad, 0) + f["cantidad"], 3)
        return {
            "tipo": self.tipo,
            "nombre": inventarios.TIPOS[self.tipo]["nombre"],
            "total": total,
            "registros": len(self.filas),
            "anulados": len(self.filas) - total,
            "con_ubicacion": len(con),
            "sin_ubicacion": total - len(con),
            "cobertura": round(100 * len(con) / total, 1) if total else 0.0,
            "kg": round(por_unidad.get("kg", 0.0), 1),
            "por_unidad": {k: round(v, 3) for k, v in por_unidad.items()},
            "semaforo": semaforo,
            "sin_lote": sum(1 for f in activas if not f["lote"]),
            "sin_caducidad": sum(1 for f in activas if not f["caducidad"]),
            "con_formula": sum(1 for f in self.filas if f["cantidad_formula"]),
            "sin_cantidad_legible": sum(1 for f in self.filas if not f["cantidad_legible"]),
            "filas_sin_cantidad": [f["fila"] for f in self.filas if not f["cantidad_legible"]],
            # códigos escritos en el Excel que NO existen en el catálogo: no cuentan como
            # ubicación (esa es la regla de fuente única) y la pantalla los señala.
            "ubicaciones_invalidas": sum(1 for f in self.filas if f["ubicacion_invalida"]),
            "filas_ubicacion_invalida": [
                {"fila": f["fila"], "valor": f["ubicacion_invalida"]}
                for f in self.filas if f["ubicacion_invalida"]],
            "por_zona": self._por_zona(),
            "actualizado": dt.datetime.now().isoformat(sep=" ", timespec="seconds"),
        }

    def _por_zona(self):
        conteo = {}
        for f in self.filas:
            if not f["ubicacion"] or not self._tiene_existencia(f):
                continue
            zona = ubicaciones.zona_de(f["ubicacion"])
            conteo[zona] = conteo.get(zona, 0) + 1
        return conteo

    # ---------- escritura ----------
    def asignar(self, fila, codigo):
        """Escribe la ubicación en el libro. Devuelve (anterior, nueva).

        **Solo se escribe un código del catálogo** (fuente única). La captura tolerante
        se convierte al canónico; si no está en el catálogo se rechaza con ValueError y
        el Excel no se toca. Así no se pueden crear ubicaciones fantasma.

        Ojo: `ws.cell(..., value=None)` NO borra la celda (openpyxl solo asigna cuando el
        valor no es None). Para "quitar ubicación" hay que asignar `.value = None`; si no,
        la pantalla decía «sin ubicación» pero el Excel conservaba la anterior y al
        recargar el lote volvía a aparecer en su rack.
        """
        registro = self.por_fila(fila)
        if codigo in (None, "", "-"):
            nueva = None
        else:
            nueva = ubicaciones.canonico(codigo)
            if nueva is None:
                raise ValueError(f"la ubicación «{codigo}» no existe en el catálogo "
                                 f"ALM-PNO-002; solo se aceptan códigos del catálogo")
        anterior = registro["ubicacion"]
        if self.en_db:
            self.con.execute(
                "UPDATE inventario_lote SET ubicacion=?, ubicacion_cruda=?, "
                "ubicacion_invalida=NULL, actualizado=NOW() WHERE inventario=? AND fila=?",
                (nueva, nueva, self.tipo, int(fila)))
        else:
            celda = self.ws.cell(row=fila, column=self.columnas["ubicacion"])
            celda.value = nueva
            if self.ws.cell(row=fila, column=self.columnas["ubicacion"] - 1).has_style:
                celda._style = copy(
                    self.ws.cell(row=fila, column=self.columnas["ubicacion"] - 1)._style)
        registro["ubicacion"] = nueva
        registro["ubicacion_cruda"] = nueva
        registro["ubicacion_invalida"] = None
        self._indice = None
        return anterior, nueva

    def construir_workbook(self):
        """Genera un libro openpyxl en memoria a partir de `self.filas` (para exportar desde BD)."""
        if self.wb is not None and not self.en_db:
            return self.wb
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self.nombre[:31]
        ws["A1"] = f"INVENTARIO — {self.nombre.upper()}"
        encabezados = [
            ("DESCRIPCIÓN", "descripcion", 44),
            ("CATEGORÍA", "categoria", 18),
            ("PROVEEDOR", "proveedor", 24),
            ("PRESENTACIÓN", "presentacion", 18),
            ("PIEZAS POR CAJA", "piezas", 14),
            ("UNIDAD", "unidad", 10),
            ("CANTIDAD", "cantidad", 12),
            ("CADUCIDAD", "caducidad", 14),
            ("LOTE", "lote", 20),
            ("CLAVE ERP", "erp", 16),
            ("OBSERVACIONES", "observaciones", 28),
            (ENCABEZADO_UBICACION, "ubicacion", 16),
            (ENCABEZADO_QR, "qr_id", 20),
        ]
        for col_idx, (titulo, _, ancho) in enumerate(encabezados, start=1):
            ws.cell(row=self.fila_encabezado, column=col_idx, value=titulo)
            ws.column_dimensions[get_column_letter(col_idx)].width = ancho
        for f in sorted(self.filas, key=lambda x: x["fila"]):
            r_idx = max(int(f["fila"]), self.fila_encabezado + 1)
            for col_idx, (_, campo, _) in enumerate(encabezados, start=1):
                valor = f.get(campo)
                if campo == "ubicacion" and not valor and f.get("ubicacion_invalida"):
                    valor = f.get("ubicacion_invalida")
                if campo == "caducidad" and not valor and f.get("caducidad_cruda"):
                    valor = f.get("caducidad_cruda")
                ws.cell(row=r_idx, column=col_idx, value=valor)
        return wb

    def guardar(self, destino=None):
        """Guarda cambios en la BD (Supabase) o en el libro Excel en disco."""
        if self.en_db:
            self.con.commit()
            if destino is not None:
                destino = pathlib.Path(destino)
                destino.parent.mkdir(parents=True, exist_ok=True)
                self.construir_workbook().save(destino)
                return destino
            return self.ruta
        destino = pathlib.Path(destino or self.ruta)
        destino.parent.mkdir(parents=True, exist_ok=True)
        temporal = destino.with_suffix(".tmp.xlsx")
        try:
            self.wb.save(temporal)
            temporal.replace(destino)
        except Exception:
            try:
                temporal.unlink(missing_ok=True)
            except OSError:
                pass
            raise
        return destino

    # ---------- cantidades y parcialidades ----------
    def ajustar_cantidad(self, fila, nueva):
        registro = self.por_fila(fila)
        valor = float(nueva)
        if self.en_db:
            self.con.execute(
                "UPDATE inventario_lote SET cantidad=?, cantidad_texto=NULL, "
                "cantidad_legible=TRUE, actualizado=NOW() WHERE inventario=? AND fila=?",
                (valor, self.tipo, int(fila)))
        else:
            col = self.columnas.get("cantidad")
            if not col:
                raise ValueError("el libro no tiene columna de cantidad")
            modelo = self.ws.cell(row=fila, column=col)
            celda = self.ws.cell(row=fila, column=col, value=valor)
            if modelo.has_style:
                celda._style = copy(modelo._style)
        registro["cantidad"] = valor
        registro["cantidad_legible"] = True
        return registro

    def agregar_parcialidad(self, fila_origen, cantidad, destino, nota=None):
        """Crea un renglón nuevo con la parte que se mueve y deja el resto en el original."""
        if cantidad is None or cantidad <= 0:
            raise ValueError("la cantidad a mover debe ser mayor que cero")
        origen = self.por_fila(fila_origen)
        disponible = float(origen["cantidad"] or 0)
        if cantidad >= disponible:
            raise ValueError(f"la cantidad ({cantidad}) debe ser menor que la existencia "
                             f"({disponible}); para mover todo usa el movimiento completo")
        destino = ubicaciones.canonico(destino)
        if destino is None:
            raise ValueError("la ubicación de destino no existe en el catálogo ALM-PNO-002")

        if self.en_db:
            fila = max((f["fila"] for f in self.filas), default=self.fila_encabezado) + 1
            previo = str(origen.get("observaciones") or "").strip()
            marca = nota or f"parcialidad de la fila {fila_origen}"
            obs = f"{previo} · {marca}" if previo else marca
            self.con.execute(
                "INSERT INTO inventario_lote (inventario, fila, descripcion, categoria, "
                "proveedor, unidad, cantidad, cantidad_legible, caducidad, caducidad_cruda, "
                "motivo_caducidad, lote, erp, qr_id, observaciones, ubicacion, ubicacion_cruda, "
                "presentacion, piezas) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (self.tipo, fila, origen["descripcion"], origen.get("categoria"),
                 origen.get("proveedor"), origen.get("unidad") or "kg",
                 round(float(cantidad), 3), True, origen.get("caducidad"),
                 origen.get("caducidad_cruda"), origen.get("motivo_caducidad"),
                 origen.get("lote"), origen.get("erp"), None, obs, destino, destino,
                 origen.get("presentacion"), origen.get("piezas")))
            self.ajustar_cantidad(fila_origen, round(disponible - float(cantidad), 3))
            self.filas = self._leer_db()
            return self.por_fila(fila)

        fila = self.ws.max_row + 1
        col_desc = self.columnas["descripcion"]
        while fila > self.fila_encabezado + 1 and not self.ws.cell(row=fila - 1, column=col_desc).value:
            fila -= 1

        for campo, col in self.columnas.items():
            valor = self.ws.cell(row=fila_origen, column=col).value
            if campo == "cantidad":
                valor = round(float(cantidad), 3)
            elif campo == "ubicacion":
                valor = destino
            elif campo in ("presentacion", "piezas"):
                valor = self.ws.cell(row=fila_origen, column=col).value
            elif campo == "qr_id":
                # la etiqueta física se queda con el renglón que sigue en el rack; la
                # parcialidad es una tarima nueva y se etiqueta aparte si hace falta
                valor = None
            elif campo == "observaciones":
                previo = str(valor).strip() if valor not in (None, "") else ""
                marca = nota or f"parcialidad de la fila {fila_origen}"
                valor = f"{previo} · {marca}" if previo else marca
            celda = self.ws.cell(row=fila, column=col, value=valor)
            modelo = self.ws.cell(row=fila_origen, column=col)
            if modelo.has_style:
                celda._style = copy(modelo._style)

        self.ajustar_cantidad(fila_origen, round(disponible - float(cantidad), 3))
        self.filas = self._leer()
        return self.por_fila(fila)

    # ---------- alta de un lote nuevo ----------
    # la caducidad no se pide igual en las tres naturalezas: ver validar_alta()
    CAMPOS_OBLIGATORIOS = (("descripcion", "la descripción del material"),
                           ("lote", "el lote (sin lote no hay trazabilidad, ALM-PNO-002 §V.2.2.2)"),
                           ("proveedor", "el proveedor"),
                           ("cantidad", "la cantidad"))

    def validar_alta(self, datos, permitir_vencida=False):
        """Revisa todo antes de tocar el libro y devuelve los valores ya normalizados."""
        faltan = [texto for campo, texto in self.CAMPOS_OBLIGATORIOS
                  if str(datos.get(campo) or "").strip() in ("", "None")]
        if faltan:
            raise ValueError("falta " + ", ".join(faltan))

        descripcion = str(datos["descripcion"]).strip()
        if len(descripcion) < 3:
            raise ValueError("la descripción es demasiado corta")
        lote = str(datos["lote"]).strip()
        if len(lote) < 3:
            raise ValueError("el lote debe tener al menos 3 caracteres")
        proveedor = str(datos["proveedor"]).strip()

        reglas = inventarios.TIPOS[self.tipo]
        unidad = normalizar_unidad(datos.get("unidad"), reglas["unidad_preferida"])
        cantidad = interpretar_cantidad(datos.get("cantidad"), unidad=unidad)
        if cantidad is None:
            raise ValueError(f"la cantidad no es un número válido: {datos.get('cantidad')!r}")
        cantidad = round(cantidad, 3)
        if cantidad <= 0:
            raise ValueError("la cantidad debe ser mayor que cero")

        crudo_caducidad = str(datos.get("caducidad") or "").strip()
        iso, motivo = fechas.parsear(crudo_caducidad)
        if crudo_caducidad in ("", "None"):
            # la caducidad es opcional en material de empaque y obligatoria en MP y PT
            if reglas["exige_caducidad"]:
                raise ValueError(f"falta la caducidad: el inventario {reglas['corto']} la necesita "
                                 f"para el PEPS/FEFO (ALM-PNO-002)")
        elif not iso:
            raise ValueError(f"la caducidad no es una fecha válida ({motivo or 'ilegible'}); "
                             f"usa AAAA-MM-DD o DD/MM/AAAA")
        if iso:
            hoy = dt.date.today().isoformat()
            if not permitir_vencida and iso < hoy:
                raise ValueError(f"la caducidad ({iso}) ya venció; revisa la fecha antes de ingresar")

        ubicacion = ubicaciones.canonico(datos.get("ubicacion"))
        if str(datos.get("ubicacion") or "").strip() and ubicacion is None:
            raise ValueError(f"la ubicación «{datos.get('ubicacion')}» no existe en el catálogo "
                             f"de ALM-PNO-002")

        for f in self.filas:
            if (normalizar(f["descripcion"]) == normalizar(descripcion)
                    and normalizar(f["lote"]) == normalizar(lote)
                    and self._tiene_existencia(f)):
                raise ValueError(f"ya existe el lote {lote} de {descripcion} en la fila {f['fila']}; "
                                 f"si es una parcialidad, muévela desde ese renglón")

        qr_id = clave_qr(datos.get("qr_id")) or None
        if qr_id:
            for f in self.filas:
                if (f.get("qr_id") and clave_qr(f["qr_id"]) == qr_id
                        and self._tiene_existencia(f)):
                    raise ValueError(f"el código «{qr_id}» ya está en la fila {f['fila']} "
                                     f"({f['descripcion']}); una etiqueta no puede estar en dos "
                                     f"tarimas activas")

        return {"descripcion": descripcion, "lote": lote, "proveedor": proveedor,
                "cantidad": cantidad, "unidad": unidad, "caducidad": iso,
                "ubicacion": ubicacion or None,
                "qr_id": qr_id,
                "categoria": (str(datos.get("categoria") or "").strip() or None),
                "erp": (str(datos.get("erp") or "").strip() or None),
                "observaciones": (str(datos.get("observaciones") or "").strip() or None),
                "presentacion": (str(datos.get("presentacion") or "").strip() or None),
                "piezas": (str(datos.get("piezas") or "").strip() or None)}

    def agregar_alta(self, datos, permitir_vencida=False):
        """Agrega un lote al final del libro o BD. Todo se valida antes de escribir."""
        limpio = self.validar_alta(datos, permitir_vencida=permitir_vencida)
        if self.en_db:
            fila = max((f["fila"] for f in self.filas), default=self.fila_encabezado) + 1
            self.con.execute(
                "INSERT INTO inventario_lote (inventario, fila, descripcion, categoria, "
                "proveedor, unidad, cantidad, cantidad_legible, caducidad, caducidad_cruda, "
                "lote, erp, qr_id, observaciones, ubicacion, ubicacion_cruda, presentacion, "
                "piezas) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (self.tipo, fila, limpio["descripcion"], limpio["categoria"],
                 limpio["proveedor"], limpio["unidad"], limpio["cantidad"], True,
                 limpio["caducidad"], limpio["caducidad"], limpio["lote"], limpio["erp"],
                 limpio["qr_id"], limpio["observaciones"], limpio["ubicacion"],
                 limpio["ubicacion"], limpio["presentacion"], limpio["piezas"]))
            self.filas = self._leer_db()
            return self.por_fila(fila)
        self._asegurar_columna_qr()

        fila = max(self.ws.max_row + 1, self.fila_encabezado + 1)
        col_desc = self.columnas["descripcion"]
        while fila > self.fila_encabezado + 1 and not self.ws.cell(row=fila - 1,
                                                                   column=col_desc).value:
            fila -= 1

        valores = {
            "descripcion": limpio["descripcion"],
            "proveedor": limpio["proveedor"],
            "categoria": limpio["categoria"],
            "unidad": limpio["unidad"],
            "cantidad": limpio["cantidad"],
            "caducidad": limpio["caducidad"],
            "lote": limpio["lote"],
            "erp": limpio["erp"],
            "observaciones": limpio["observaciones"],
            "ubicacion": limpio["ubicacion"],
            "qr_id": limpio["qr_id"],
            "presentacion": limpio["presentacion"],
            "piezas": limpio["piezas"],
        }
        for campo, col in self.columnas.items():
            valor = valores.get(campo)
            celda = self.ws.cell(row=fila, column=col, value=valor)
            modelo = self.ws.cell(row=fila - 1, column=col)
            if modelo.has_style:
                celda._style = copy(modelo._style)

        self.filas = self._leer()
        return self.por_fila(fila)

    def sacar(self, fila, cantidad, nota=None):
        """Saca material del almacén hacia producción: baja la cantidad del renglón.

        Si se lleva todo, el renglón queda en cero y sin ubicación (ya no está en el almacén);
        si es parcial, el resto se queda donde estaba.
        """
        registro = self.por_fila(fila)
        disponible = float(registro["cantidad"] or 0)
        cantidad = float(cantidad)
        if cantidad <= 0:
            raise ValueError("la cantidad surtida debe ser mayor que cero")
        if cantidad > disponible + 0.001:
            raise ValueError(f"no puedes sacar {cantidad}; hay {disponible} {registro['unidad']} "
                             f"en ese renglón")
        nota = nota or "surtido a producción"
        if cantidad >= disponible - 0.001:
            self.anular(fila, nota)
        else:
            self.ajustar_cantidad(fila, round(disponible - cantidad, 3))
        return self.por_fila(fila)

    def anular(self, fila, nota):
        """Deja el renglón en cero y sin ubicación, con la marca en observaciones."""
        self.ajustar_cantidad(fila, 0)
        self.asignar(fila, None)
        if self.en_db:
            registro = self.por_fila(fila)
            previo = registro.get("observaciones")
            texto = f"{str(previo).strip() + ' · ' if previo else ''}{nota}"
            self.con.execute(
                "UPDATE inventario_lote SET observaciones=?, actualizado=NOW() "
                "WHERE inventario=? AND fila=?",
                (texto, self.tipo, int(fila)))
            registro["observaciones"] = texto
            return registro
        col = self.columnas.get("observaciones")
        if col:
            previo = self.ws.cell(row=fila, column=col).value
            texto = f"{str(previo).strip() + ' · ' if previo else ''}{nota}"
            celda = self.ws.cell(row=fila, column=col)
            celda.value = texto
            registro = self.por_fila(fila)
            if "observaciones" in registro:
                registro["observaciones"] = texto
        return self.por_fila(fila)

    def recargar(self):
        """Vuelve a leer los renglones del libro o BD en memoria."""
        self.filas = self._leer_db() if self.en_db else self._leer()
        return self.filas

    def borrar_renglones(self, filas):
        """Quita renglones completos del inventario (solo el administrador, con motivo).

        Devuelve {fila original: registro}, con lo que había en cada renglón, para dejarlo en la
        bitácora: el renglón desaparece, pero queda el rastro de qué se borró y quién lo hizo.
        """
        objetivo = sorted({int(f) for f in filas}, reverse=True)
        if not objetivo:
            raise ValueError("no se indicó ningún renglón")
        registros = {}
        max_fila = max((f["fila"] for f in self.filas), default=self.fila_encabezado) \
            if self.en_db else self.ws.max_row
        for fila in objetivo:
            if fila <= self.fila_encabezado or fila > max_fila:
                raise ValueError(f"la fila {fila} está fuera del inventario")
            try:
                registros[fila] = self.por_fila(fila)
            except KeyError as e:
                raise ValueError(f"la fila {fila} no tiene material") from e
        for fila in objetivo:          # de abajo hacia arriba, para no correr los índices
            if self.en_db:
                self.con.execute(
                    "DELETE FROM inventario_lote WHERE inventario=? AND fila=?",
                    (self.tipo, int(fila)))
                self.con.execute(
                    "UPDATE inventario_lote SET fila = fila - 1, actualizado=NOW() "
                    "WHERE inventario=? AND fila > ?",
                    (self.tipo, int(fila)))
            else:
                self.ws.delete_rows(fila, 1)
        self.recargar()
        return registros

    def copia(self, destino):
        destino = pathlib.Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb = self.construir_workbook() if self.en_db else self.wb
        wb.save(destino)
        return destino
