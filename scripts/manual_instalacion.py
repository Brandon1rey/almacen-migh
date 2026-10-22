"""Genera el manual de instalación y despliegue (Word + PDF)."""
import pathlib
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

DOCX = RAIZ / "docs" / "Manual_instalacion_almacen.docx"
AZUL = RGBColor(0x1F, 0x3B, 0x73)
ROJO = RGBColor(0xB0, 0x1C, 0x1C)
VERDE = RGBColor(0x1E, 0x6B, 0x3A)
GRIS = RGBColor(0x4B, 0x55, 0x63)


def parrafo(doc, texto, negritas=False, color=None, tam=10, espacio=4):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negritas
    run.font.size = Pt(tam)
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(espacio)
    return p


def vineta(doc, texto, tam=9, estilo="List Bullet"):
    p = doc.add_paragraph(style=estilo)
    run = p.add_run(texto)
    run.font.size = Pt(tam)
    p.paragraph_format.space_after = Pt(2)


def tabla(doc, encabezados, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    for i, h in enumerate(encabezados):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = AZUL
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8)
    if anchos:
        for fila in t.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)
    return t


def construir(destino=DOCX):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)

    parrafo(doc, "COMERCIALIZADORA MIGH DE MÉXICO S.A. DE C.V.", True, AZUL, 11, 0)
    parrafo(doc, "Aseguramiento de Calidad · Almacén de materias primas", False, GRIS, 8, 10)
    parrafo(doc, "Manual de instalación y despliegue", True, None, 16, 4)
    parrafo(doc, "Sistema de control de ubicación del inventario · versión 1.0 · octubre 2026",
            False, GRIS, 9, 12)

    parrafo(doc, "1. Requisitos", True, AZUL, 12)
    tabla(doc, ["Concepto", "Requisito", "Nota"], [
        ["Sistema operativo", "Windows 10 u 11", "Probado en Windows 11"],
        ["Python", "3.10 o superior con python.exe en PATH",
         "Si no hay internet, solo hace falta Python; las librerías van en vendor\\"],
        ["Dependencia externa", "openpyxl 3.1.5", "Único paquete; el servidor usa biblioteca estándar"],
        ["Espacio en disco", "≈ 40 MB", "Sin base de datos instalada"],
        ["Carpeta", "C:\\almacen-ubicaciones", "Nunca en Program Files (permisos de escritura)"],
        ["Red", "No requiere internet para operar", "Solo para instalar dependencias"],
    ], anchos=[3.6, 6.4, 7.4])

    parrafo(doc, "2. Instalación en una PC nueva", True, AZUL, 12)
    for paso in [
        "Copiar la carpeta del sistema a C:\\almacen-ubicaciones.",
        "Colocar el Excel del área en datos_iniciales\\Inventario_MP.xlsx "
        "(o indicar la ruta al instalar con --origen).",
        "Doble clic en instalar.cmd: crea .venv\\, instala openpyxl, agrega la columna UBICACIÓN "
        "al inventario y corre la verificación.",
        "Doble clic en arrancar.cmd y abrir http://127.0.0.1:8300.",
        "Escribir el nombre del usuario la primera vez (queda en la bitácora).",
    ]:
        vineta(doc, paso, estilo="List Number")

    parrafo(doc, "3. Verificación (obligatoria antes de entregar la PC)", True, AZUL, 12)
    parrafo(doc, "Ejecutar:  python scripts\\verificar_instalacion.py", False, None, 9, 3)
    tabla(doc, ["Revisión", "Debe reportar"], [
        ["Python", "3.10 o superior"],
        ["openpyxl", "openpyxl 3.1.5"],
        ["Estructura de la aplicación", "app completa"],
        ["Inventario (Excel del área)", "n lotes · tamaño en KB"],
        ["Permisos de escritura", "data/ escribible"],
        ["Bitácora SQLite", "n movimientos (0 en una instalación nueva)"],
        ["Puerto del servidor", "libre (o ya en uso si el sistema está corriendo)"],
    ], anchos=[6.0, 11.4])
    parrafo(doc, "Si alguna línea dice FALLA, no se entrega la PC: el detalle indica qué falta.",
            False, ROJO, 9, 8)

    parrafo(doc, "4. Varias PC: una sola verdad", True, AZUL, 12)
    parrafo(doc, "El sistema guarda los datos en la PC donde corre. Instalarlo en varias PC con copias "
                 "independientes produce varias verdades y bitácoras que no se pueden fusionar. "
                 "Dos formas correctas de operar:", False, None, 9, 4)
    vineta(doc, "PC servidor (recomendado): en la PC del almacén se corre servidor_lan.cmd; las demás "
                "PC solo abren http://IP:8300 en el navegador. No se instala nada en ellas.")
    vineta(doc, "Copia manual (sin red): cada PC con su copia y conciliación con el F01 impreso. "
                "Solo como excepción y con autorización de Calidad.")
    parrafo(doc, "Regla de firewall, una sola vez y como administrador:", False, None, 9, 2)
    parrafo(doc, "netsh advfirewall firewall add rule name=\"Almacen MIGH 8300\" dir=in action=allow "
                 "protocol=TCP localport=8300", False, AZUL, 8, 4)
    parrafo(doc, "El modo red no tiene contraseña: identifica al usuario por el nombre que escribe. "
                 "Es adecuado para la LAN de la planta; no debe exponerse a internet.",
            False, ROJO, 9, 8)

    parrafo(doc, "5. Respaldo y continuidad", True, AZUL, 12)
    tabla(doc, ["Qué", "Cómo", "Frecuencia"], [
        ["Inventario + bitácora", "respaldo.cmd", "Diario (o al cierre de turno)"],
        ["Copia externa", "respaldo.cmd --copia-a \"E:\\respaldo_almacen\"", "Semanal o al cerrar el mes"],
        ["Arranque automático", "arranque_automatico.cmd (tarea al iniciar sesión)", "Una vez por PC"],
        ["Rotación", "60 días por defecto (--dias 90 para ampliar)", "Automática"],
    ], anchos=[4.2, 8.6, 4.6])
    parrafo(doc, "Cada respaldo deja un manifiesto JSON con el SHA-256 de los archivos: sirve para probar "
                 "que la copia es idéntica a la original.", False, GRIS, 8, 6)

    parrafo(doc, "6. Actualización a una versión nueva", True, AZUL, 12)
    for paso in ["respaldo.cmd (obligatorio; sin respaldo no se actualiza).",
                 "Copiar la carpeta nueva encima, sin borrar data\\ ni datos_iniciales\\.",
                 "instalar.cmd otra vez.",
                 "arrancar.cmd y comparar el tablero: mismos lotes, mismas ubicaciones.",
                 "La bitácora se migra sola (columnas motivo y ref_doc)."]:
        vineta(doc, paso, estilo="List Number")

    parrafo(doc, "7. Diagnóstico de fallas frecuentes", True, AZUL, 12)
    tabla(doc, ["Síntoma", "Causa probable", "Acción"], [
        ["No se encontró Python", "Python sin PATH", "Reinstalar marcando “Add python.exe to PATH”"],
        ["No module named openpyxl", "Dependencia ausente", "instalar.cmd o copiar vendor\\"],
        ["El archivo está abierto en Excel", "Bloqueo del .xlsx", "Cerrar Excel y reintentar"],
        ["Puerto 8300 en uso", "Ya hay una instancia", "Usar la que corre o cambiar ALMACEN_PUERTO"],
        ["No abre desde otra PC", "Firewall / red de invitados", "Aplicar la regla; usar la red de planta"],
        ["Datos no cuadran con el piso", "Captura pendiente", "Contar por zona y ajustar con nota"],
    ], anchos=[5.2, 5.2, 7.0])

    parrafo(doc, "8. Qué queda fuera del alcance", True, AZUL, 12)
    for texto in ["Integración con el ERP: el código ERP es un campo de texto para conciliar a mano.",
                  "Usuarios y contraseñas: la identificación es el nombre capturado.",
                  "Firma electrónica: el reporte F01 se imprime y se firma a mano.",
                  "Acceso desde internet: el sistema es local o de red interna."]:
        vineta(doc, texto)

    parrafo(doc, "Anexo. Contenido del paquete", True, AZUL, 12)
    tabla(doc, ["Archivo / carpeta", "Para qué"], [
        ["instalar.cmd", "Instalación en una PC nueva"],
        ["arrancar.cmd", "Operación normal (localhost)"],
        ["servidor_lan.cmd", "Modo red: una PC con los datos, otras por navegador"],
        ["arranque_automatico.cmd / quitar_arranque.cmd", "Alta y baja de la tarea de inicio de sesión"],
        ["respaldo.cmd", "Respaldo con manifiesto SHA-256"],
        ["requirements.txt", "Dependencias de ejecución (openpyxl)"],
        ["requirements-docs.txt", "Dependencias para generar Word/PDF (no operativas)"],
        ["vendor\\", "openpyxl y et_xmlfile para PC sin internet ni pip"],
        ["data\\Inventario_MP.xlsx", "Inventario del área con la columna UBICACIÓN"],
        ["data\\bitacora.db", "Movimientos (rastro de auditoría)"],
        ["LEEME_INSTALACION.md", "Guía rápida de instalación"],
        ["MANIFIESTO.txt", "SHA-256 de cada archivo del paquete"],
    ], anchos=[7.4, 10.0])

    parrafo(doc, "Elaboró: Aseguramiento de Calidad · Revisó: Jefe de Almacén · Autorizó: "
                 "Jefe de Aseguramiento de Calidad", False, GRIS, 8, 4)
    parrafo(doc, "Documento de trabajo. No sustituye ALM-PNO-002 ni el formato F01-ALM-PNO-002.",
            False, GRIS, 7, 0)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


def a_pdf(docx_path):
    import win32com.client as win32
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(docx_path))
    pdf = docx_path.with_suffix(".pdf")
    doc.ExportAsFixedFormat(str(pdf), 17)
    doc.Close(False)
    word.Quit()
    return pdf


if __name__ == "__main__":
    ruta = construir()
    print("docx:", ruta)
    print("pdf :", a_pdf(ruta))
