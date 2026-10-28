"""Crea la plantilla vacía de un inventario (MP, MEP o PT) con el mismo formato del archivo del área.

    python scripts\\plantilla_inventario.py --tipo mep
    python scripts\\plantilla_inventario.py --tipo pt --destino "D:\\Inventario_PT.xlsx"

El archivo que produce se puede llenar a mano o pegar una tabla del área, siempre que conserve
los encabezados. El sistema agrega la columna UBICACIÓN si falta, pero es mejor traerla.
"""
import argparse
import pathlib
import sys

# La consola de Windows no siempre es UTF-8: sin esto, imprimir «→» o «·» lanza
# UnicodeEncodeError y el script parece fallar aunque ya haya hecho su trabajo.
for _flujo in (sys.stdout, sys.stderr):
    try:
        _flujo.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass


import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from app import inventarios  # noqa: E402

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]
ANCHOS = [46, 22, 16, 10, 11, 16, 20, 14, 34, 16]
AZUL = PatternFill("solid", fgColor="1F3B73")
GRIS = PatternFill("solid", fgColor="EEF2F7")
BORDE = Border(*[Side(style="thin", color="B7C0CC")] * 4)

EJEMPLOS = {
    "mp": [["WPC 80%", "VILHER", "Proteína", "kg", 1458, "2028-04-23", "22260425002",
            "MIGH-MP114", "", "MPA-A-005"]],
    "mep": [["BOLSA LAMINADA 2 KG", "IMPRESOS DEL NORTE", "Bolsa", "pza", 24000, "2029-01-31",
             "BL-2026-01", "MIGH-ME020", "caja de 2,000 pzas", "MEP-B-004"]],
    "pt": [["NUTRILIFE VAINILLA 900 G", "MIGH", "Polvo", "pza", 1200, "2028-02-15",
            "PT-2601-A", "MIGH-PT003", "caja colectiva 12 pzas", "MPT-B-002"]],
}

INSTRUCCIONES = [
    ("PLANTILLA DE INVENTARIO — {nombre} ({corto})", True, 13),
    ("Comercializadora MIGH de México · Aseguramiento de Calidad", False, 9),
    ("", False, 9),
    ("Base documental: ALM-PNO-002 §V.2.1.4 (zonas), §V.2.1.8 y §V.2.3.8 (reporte F01-ALM-PNO-002).",
     False, 9),
    ("", False, 9),
    ("Cómo llenarla", True, 11),
    ("1. Una fila por lote. Si un lote está en dos ubicaciones, van dos filas (parcialidad).", False, 9),
    ("2. No cambies los encabezados de la fila 3: el sistema los reconoce por nombre, no por posición.",
     False, 9),
    ("3. La columna UBICACIÓN se puede dejar vacía; se captura en el sistema o aquí con el código",
     False, 9),
    ("   del rack (por ejemplo MEP-B-004). Si el código no existe en el catálogo, el sistema lo rechaza.",
     False, 9),
    ("4. Caducidad: fecha de Excel o texto AAAA-MM-DD / DD/MM/AAAA / AAAA/mmm/DD. Si no es legible,",
     False, 9),
    ("   se muestra como «sin dato»; no se adivina.", False, 9),
    ("5. Cantidad: número. Las sumas de Excel se aceptan, pero es mejor dejar el valor calculado.",
     False, 9),
    ("", False, 9),
    ("Zonas que le corresponden según ALM-PNO-002", True, 11),
    ("{zonas}", False, 9),
    ("", False, 9),
    ("Restos y parcialidades", True, 11),
    ("El procedimiento pide incluirlos en el mismo reporte F01 (§V.2.2.6). En el sistema se registran",
     False, 9),
    ("moviendo solo la cantidad surtida: queda el resto en su rack y un renglón nuevo en el destino.",
     False, 9),
    ("", False, 9),
    ("Frecuencia del reporte: {reporte} ({base}).", False, 9),
]


def crear(tipo, destino, con_ejemplos=True):
    tipo = inventarios.normalizar_tipo(tipo)
    ficha = inventarios.TIPOS[tipo]
    destino = pathlib.Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws["A1"] = f"INVENTARIO DE {ficha['nombre'].upper()} — COMERCIALIZADORA MIGH DE MÉXICO"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Formato de captura · base ALM-PNO-002 · reporte {ficha['reporte']}"

    for i, texto in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=3, column=i, value=texto)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = AZUL
        c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ANCHOS[i - 1]
    ws.row_dimensions[3].height = 28

    for n, ejemplo in enumerate(EJEMPLOS[tipo] if con_ejemplos else [], start=4):
        for i, valor in enumerate(ejemplo, start=1):
            c = ws.cell(row=n, column=i, value=valor)
            c.font = Font(size=9, italic=True, color="6B7280")
            c.border = BORDE
            c.fill = GRIS

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:J{max(ws.max_row, 4)}"

    hoja = wb.create_sheet("INSTRUCCIONES")
    hoja.column_dimensions["A"].width = 108
    fila = 1
    for texto, negritas, tam in INSTRUCCIONES:
        limpio = (texto.replace("{nombre}", ficha["nombre"]).replace("{corto}", ficha["corto"])
                  .replace("{zonas}", ", ".join(ficha["zonas"]))
                  .replace("{reporte}", ficha["reporte"]).replace("{base}", ficha["base"]))
        c = hoja.cell(row=fila, column=1, value=limpio)
        c.font = Font(bold=negritas, size=tam)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        fila += 1

    wb.save(destino)
    return destino


def main(argv=None):
    ap = argparse.ArgumentParser(description="Plantilla de inventario por tipo")
    ap.add_argument("--tipo", required=True, choices=list(inventarios.TIPOS))
    ap.add_argument("--destino", default=None)
    ap.add_argument("--vacia", action="store_true", help="sin fila de ejemplo (para producción)")
    args = ap.parse_args(argv)
    destino = args.destino or inventarios.ruta(args.tipo)
    ruta = crear(args.tipo, destino, con_ejemplos=not args.vacia)
    print(f"Plantilla creada: {ruta}")
    print("Llénala con una fila por lote y colócala en data\\ (o usa --origen al preparar).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
