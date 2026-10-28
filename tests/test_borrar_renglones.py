"""Borrado de renglones de tajo (solo administrador): respaldo, auditoría y filas recorridas."""
import pathlib
import sqlite3
import tempfile
import unittest

from app import api, bitacora, inventario, usuarios

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]
FILAS = [("LECHE EN POLVO", "L-100", 500, "MPA-A-001"),
         ("AZUCAR", "L-200", 300, "MPA-A-002"),
         ("WPC 80%", "L-300", 200, None),
         ("MALTODEXTRINA", "L-400", 150, "MPA-A-004"),
         ("SOYA", "L-500", 90, "MPA-A-005")]


def libro(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=3, column=i, value=texto)
    for n, (desc, lote, cantidad, ubicacion) in enumerate(FILAS, start=4):
        ws.cell(row=n, column=1, value=desc)
        ws.cell(row=n, column=2, value="PROVEEDOR")
        ws.cell(row=n, column=4, value="kg")
        ws.cell(row=n, column=5, value=cantidad)
        ws.cell(row=n, column=6, value="2028-01-01")
        ws.cell(row=n, column=7, value=lote)
        if ubicacion:
            ws.cell(row=n, column=10, value=ubicacion)
    wb.save(ruta)
    return ruta


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.carpeta = pathlib.Path(self._tmp.name)
        self.ruta = libro(self.carpeta / "Inventario_MP.xlsx")
        self.inv = inventario.Inventario(self.ruta, tipo="mp")
        self.con = bitacora.conectar(self.carpeta / "b.db")
        usuarios.crear(self.con, "Maria Elena", "clave-admin-1", "admin")
        usuarios.crear(self.con, "Esteban", "clave-almacen-1", "almacen")
        self.sesion = self._sesion("Maria Elena")
        self.estado = type("Estado", (), {})()
        self.estado.carpeta_datos = self.carpeta
        self.estado.con = self.con
        self.estado.archivo = self.ruta.name
        self.estado.inventario = lambda tipo=None: self.inv

    def _sesion(self, nombre):
        fila = usuarios.por_nombre(self.con, nombre)
        return {"usuario_id": fila["id"], "nombre": nombre, "rol": fila["rol"],
                "debe_cambiar": 0}

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()

    def borrar(self, filas, motivo="renglón duplicado de prueba", sesion=None):
        return api.manejar(self.estado, "POST", "/api/borrar-renglones", {},
                           {"filas": filas, "motivo": motivo, "tipo": "mp"},
                           sesion or self.sesion)


class TestBorrarUnRenglon(Base):
    def test_desaparece_del_archivo(self):
        codigo, cuerpo, _ = self.borrar([6])
        self.assertEqual(codigo, 200, cuerpo)
        self.assertEqual(cuerpo["total"], 1)
        releido = openpyxl.load_workbook(self.ruta).worksheets[0]
        descripciones = [releido.cell(row=f, column=1).value for f in range(4, 9)]
        self.assertNotIn("WPC 80%", descripciones)

    def test_los_demas_suben_de_fila(self):
        self.borrar([5])
        registro = self.inv.por_fila(5)
        self.assertEqual(registro["descripcion"], "WPC 80%")   # antes fila 6
        self.assertEqual(len(self.inv.filas), 4)

    def test_queda_el_rastro_en_la_bitacora(self):
        self.borrar([6], motivo="captura equivocada de recepción")
        movimientos = bitacora.ultimos(self.con, 5, inventario="mp")
        purga = [m for m in movimientos if m["accion"] == "PURGAR"]
        self.assertEqual(len(purga), 1)
        self.assertEqual(purga[0]["lote"], "L-300")
        self.assertIn("renglón borrado", purga[0]["nota"])
        self.assertEqual(purga[0]["motivo"], "captura equivocada de recepción")

    def test_respalda_el_excel_antes_de_borrar(self):
        _, cuerpo, _ = self.borrar([4])
        respaldo = pathlib.Path(cuerpo["respaldo"])
        self.assertTrue(respaldo.exists())
        self.assertEqual(respaldo.parent, self.carpeta / "respaldos")
        original = openpyxl.load_workbook(respaldo).worksheets[0]
        self.assertEqual(original.cell(row=4, column=1).value, "LECHE EN POLVO")

    def test_el_resumen_cuadra(self):
        _, cuerpo, _ = self.borrar([4])
        self.assertEqual(cuerpo["resumen"]["total"], 4)


class TestBorrarVarios(Base):
    def test_borra_varios_de_tajo(self):
        codigo, cuerpo, _ = self.borrar([4, 7])
        self.assertEqual(codigo, 200, cuerpo)
        descripciones = [f["descripcion"] for f in self.inv.filas]
        self.assertEqual(descripciones, ["AZUCAR", "WPC 80%", "SOYA"])
        self.assertEqual(cuerpo["total"], 2)

    def test_suma_las_cantidades_borradas_en_el_detalle(self):
        _, cuerpo, _ = self.borrar([4, 8])
        cantidades = sorted(d["cantidad"] for d in cuerpo["borrados"])
        self.assertEqual(cantidades, [90, 500])

    def test_filas_repetidas_o_invertidas(self):
        _, cuerpo, _ = self.borrar([7, 4, 7])
        self.assertEqual(cuerpo["total"], 2)


class TestPermisosYValidacion(Base):
    def test_solo_admin(self):
        for rol in ("almacen",):
            codigo, cuerpo, _ = self.borrar([4], sesion=self._sesion("Esteban"))
            self.assertEqual(codigo, 403)
            self.assertIn("administrador", cuerpo["error"])
        self.assertEqual(len(self.inv.filas), 5)

    def test_pide_motivo(self):
        codigo, cuerpo, _ = self.borrar([4], motivo="x")
        self.assertEqual(codigo, 400)
        self.assertIn("motivo", cuerpo["error"])
        self.assertEqual(len(self.inv.filas), 5)

    def test_fila_inexistente(self):
        codigo, cuerpo, _ = self.borrar([99])
        self.assertEqual(codigo, 400)
        self.assertEqual(len(self.inv.filas), 5)

    def test_sin_filas(self):
        codigo, _, _ = api.manejar(self.estado, "POST", "/api/borrar-renglones", {},
                                   {"filas": [], "motivo": "motivo válido"}, self.sesion)
        self.assertEqual(codigo, 400)

    def test_no_permite_borrar_el_encabezado(self):
        codigo, cuerpo, _ = self.borrar([2])
        self.assertEqual(codigo, 400)
        self.assertIn("fuera del inventario", cuerpo["error"])
        self.assertEqual(len(self.inv.filas), 5)


class TestHistorialSigueCuadrando(Base):
    def test_las_filas_de_los_movimientos_se_recorren(self):
        # un movimiento en la fila 7 (MALTODEXTRINA) debe quedar en la 6 tras borrar la 4
        bitacora.registrar(self.con, "UBICAR",
                           {"fila": 7, "descripcion": "MALTODEXTRINA", "lote": "L-400"},
                           "Maria Elena", destino="MPA-A-004", usuario_id=1, inventario="mp")
        self.borrar([4])
        filas = [m["fila"] for m in bitacora.ultimos(self.con, 10, inventario="mp")]
        self.assertIn(6, filas)
        self.assertNotIn(7, filas)

    def test_el_historial_del_lote_sigue_apuntando_bien(self):
        bitacora.registrar(self.con, "UBICAR",
                           {"fila": 6, "descripcion": "WPC 80%", "lote": "L-300"},
                           "Maria Elena", destino="MPA-A-003", usuario_id=1, inventario="mp")
        self.borrar([5])       # borra AZUCAR (fila 5)
        historial = bitacora.por_fila(self.con, 5, 10, inventario="mp")
        self.assertTrue(any(m["accion"] == "UBICAR" and m["lote"] == "L-300" for m in historial))


if __name__ == "__main__":
    unittest.main()
