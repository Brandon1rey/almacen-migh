"""Tres inventarios (MP, MEP, PT) conviviendo: ETL, bitácora separada y F01 consolidado."""
import pathlib
import tempfile
import unittest

import openpyxl

from app import bitacora, inventarios, inventario, reportes

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones",
               "UBICACIÓN"]


def libro(ruta, filas, encabezados=None):
    """Escribe un xlsx con el mismo formato que el archivo del área (encabezado en fila 3)."""
    encabezados = encabezados or ENCABEZADOS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "INVENTARIO"
    for i, texto in enumerate(encabezados, start=1):
        ws.cell(row=3, column=i, value=texto)
    for n, fila in enumerate(filas, start=4):
        for i, valor in enumerate(fila, start=1):
            ws.cell(row=n, column=i, value=valor)
    wb.save(ruta)
    return ruta


def filas_de(tipo):
    if tipo == "mp":
        return [["WPC 80%", "VILHER", "Proteína", "kg", 1458, "2028-04-23", "22260425002",
                 "MIGH-MP114", "", "MPA-A-005"],
                ["LACA ROJO 40", "MEDINA", "Aditivo", "kg", 1.95, "2027-02-28", "260K1203",
                 "MIGH-MP001", "", "MPNA-A-072"]]
    if tipo == "mep":
        return [["BOLSA LAMINADA 2 KG", "IMPRESOS DEL NORTE", "Bolsa", "pza", 24000,
                 "2029-01-31", "BL-2026-01", "MIGH-ME020", "", "MEP-B-004"],
                ["TAPA ROSCA 38 MM", "PLASTICOS SUR", "Tapa", "pza", 8000, None, "TR-8891",
                 "MIGH-ME041", "", "MEP-C-030"]]
    return [["NUTRILIFE VAINILLA 900 G", "MIGH", "Polvo", "pza", 1200, "2028-02-15",
             "PT-2601-A", "MIGH-PT003", "caja colectiva 12 pzas", "MPT-B-002"],
            ["NUTRILIFE CHOCOLATE 900 G", "MIGH", "Polvo", "pza", 480, "2028-03-10",
             "PT-2601-B", "MIGH-PT004", "", "MPT-B-003"]]


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)
        self.rutas = {t: libro(self.dir / f"{t}.xlsx", filas_de(t)) for t in ("mp", "mep", "pt")}
        self.invs = {t: inventario.Inventario(self.rutas[t], tipo=t) for t in self.rutas}
        self.con = bitacora.conectar(self.dir / "b.db")

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()


class TestRegistro(Base):
    def test_tres_tipos_declarados(self):
        self.assertEqual(list(inventarios.TIPOS), ["mp", "mep", "pt"])
        self.assertEqual(inventarios.TIPOS["mep"]["nombre"], "Material de empaque")
        self.assertEqual(inventarios.TIPOS["pt"]["nombre"], "Producto terminado")

    def test_frecuencia_del_reporte(self):
        # ALM-PNO-002 §V.2.1.8 mensual para MP/MEP, §V.2.3.8 semanal para PT
        self.assertEqual(inventarios.TIPOS["mp"]["reporte"], "mensual")
        self.assertEqual(inventarios.TIPOS["mep"]["reporte"], "mensual")
        self.assertEqual(inventarios.TIPOS["pt"]["reporte"], "semanal")

    def test_ruta_por_tipo(self):
        self.assertEqual(inventarios.ruta("mp", self.dir).name, "Inventario_MP.xlsx")
        self.assertEqual(inventarios.ruta("mep", self.dir).name, "Inventario_MEP.xlsx")
        self.assertEqual(inventarios.ruta("pt", self.dir).name, "Inventario_PT.xlsx")

    def test_tipo_desconocido(self):
        with self.assertRaises(ValueError):
            inventarios.normalizar_tipo("xyz")


class TestLectura(Base):
    def test_cada_inventario_sabe_su_tipo(self):
        for tipo, inv in self.invs.items():
            self.assertEqual(inv.tipo, tipo)
            self.assertEqual(inv.resumen()["tipo"], tipo)
            self.assertTrue(all(f["tipo"] == tipo for f in inv.filas))

    def test_unidad_por_defecto_segun_tipo(self):
        ruta = libro(self.dir / "sin_unidad.xlsx",
                     [["BOLSA 1 KG", "PROV", "Bolsa", None, 100, None, "L-1", None, "", None]],
                     encabezados=[e for e in ENCABEZADOS if e != "Unidad de medida"])
        inv = inventario.Inventario(ruta, tipo="mep")
        self.assertEqual(inv.filas[0]["unidad"], "pza")

    def test_cantidad_y_lote_se_leen_igual_en_los_tres(self):
        self.assertEqual(self.invs["mp"].filas[0]["cantidad"], 1458)
        self.assertEqual(self.invs["mep"].filas[0]["lote"], "BL-2026-01")
        self.assertEqual(self.invs["pt"].filas[0]["unidad"], "pza")

    def test_resumen_no_mezcla_unidades(self):
        r = self.invs["mep"].resumen()
        self.assertEqual(r["por_unidad"], {"pza": 32000})
        self.assertEqual(r["kg"], 0.0)


class TestBitacora(Base):
    def test_movimientos_separados_por_inventario(self):
        bitacora.registrar(self.con, "MOVER", self.invs["mp"].filas[0], "Esteban",
                           origen="MPA-A-005", destino="MPA-A-006", inventario="mp")
        bitacora.registrar(self.con, "MOVER", self.invs["mep"].filas[0], "Esteban",
                           origen="MEP-B-004", destino="MEP-B-005", inventario="mep")
        de_mp = bitacora.por_fila(self.con, 4, inventario="mp")
        de_mep = bitacora.por_fila(self.con, 4, inventario="mep")
        self.assertEqual(len(de_mp), 1)
        self.assertEqual(de_mp[0]["ubicacion_destino"], "MPA-A-006")
        self.assertEqual(len(de_mep), 1)
        self.assertEqual(de_mep[0]["ubicacion_destino"], "MEP-B-005")

    def test_ultimo_vigente_no_se_cruza(self):
        bitacora.registrar(self.con, "MOVER", self.invs["mp"].filas[0], "Esteban",
                           origen="MPA-A-005", destino="MPA-A-006", inventario="mp")
        bitacora.registrar(self.con, "MOVER", self.invs["mep"].filas[0], "Esteban",
                           origen="MEP-B-004", destino="MEP-B-005", inventario="mep")
        self.assertEqual(bitacora.ultimo_vigente(self.con, 4, inventario="mep")["inventario"],
                         "mep")
        self.assertEqual(bitacora.ultimo_vigente(self.con, 4, inventario="mp")["inventario"],
                         "mp")

    def test_ultimos_filtra_por_tipo(self):
        bitacora.registrar(self.con, "MOVER", self.invs["mp"].filas[0], "Esteban",
                           origen="MPA-A-005", destino="MPA-A-006", inventario="mp")
        bitacora.registrar(self.con, "MOVER", self.invs["meP".lower()].filas[0], "Esteban",
                           origen="MEP-B-004", destino="MEP-B-005", inventario="mep")
        self.assertEqual(len(bitacora.ultimos(self.con, 50, inventario="mep")), 1)
        self.assertEqual(len(bitacora.ultimos(self.con, 50)), 2)


class TestReportes(Base):
    def test_f01_consolidado_trae_los_tres(self):
        destino = reportes.generar_f01([(t, self.invs[t]) for t in ("mp", "mep", "pt")],
                                       version="2", destino=self.dir / "f01.xlsx")
        wb = openpyxl.load_workbook(destino)
        self.assertIn("MP", wb.sheetnames)
        self.assertIn("MEP", wb.sheetnames)
        self.assertIn("PT", wb.sheetnames)
        self.assertIn("Portada", wb.sheetnames)
        texto = " ".join(str(c.value) for fila in wb["MEP"].iter_rows() for c in fila
                         if c.value is not None)
        self.assertIn("BOLSA LAMINADA 2 KG", texto)
        self.assertNotIn("WPC 80%", texto)

    def test_f01_portada_suma_los_tres(self):
        destino = reportes.generar_f01([(t, self.invs[t]) for t in ("mp", "mep", "pt")],
                                       version="2", destino=self.dir / "f01.xlsx")
        wb = openpyxl.load_workbook(destino)
        portada = wb["Portada"]
        valores = [c.value for fila in portada.iter_rows() for c in fila]
        self.assertIn(6, valores)  # 2 lotes por inventario

    def test_f01_de_un_solo_inventario_sigue_funcionando(self):
        destino = reportes.generar_f01(self.invs["mep"], version="2",
                                       destino=self.dir / "f01_mep.xlsx")
        wb = openpyxl.load_workbook(destino)
        self.assertEqual(wb.sheetnames, ["Hoja1"])

    def test_exportar_inventario_lleva_el_tipo_en_el_nombre(self):
        destino = reportes.exportar_inventario(self.invs["mep"],
                                               destino=self.dir / "MEP_con_ubicacion.xlsx")
        self.assertTrue(destino.exists())


class TestPlantilla(unittest.TestCase):
    def test_plantilla_trae_ubicacion_y_encabezados(self):
        from scripts import plantilla_inventario
        with tempfile.TemporaryDirectory() as tmp:
            destino = pathlib.Path(tmp) / "Inventario_MEP.xlsx"
            plantilla_inventario.crear("mep", destino)
            wb = openpyxl.load_workbook(destino)
            hoja = wb["Inventario"]
            encabezados = [hoja.cell(row=3, column=c).value for c in range(1, 11)]
            self.assertIn("UBICACIÓN", encabezados)
            self.assertIn("Descripción", encabezados)
            self.assertIn("INSTRUCCIONES", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
