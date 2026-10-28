"""El alta cambia según la naturaleza: MP, ME (empaque) y PT no piden lo mismo."""
import pathlib
import tempfile
import unittest

from app import api, bitacora, inventario, inventarios, usuarios  # primero app: habilita vendor\openpyxl

import openpyxl

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN",
               "Presentación", "Piezas por caja"]


def libro(ruta, filas=1):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=3, column=i, value=texto)
    for n in range(filas):
        ws.cell(row=4 + n, column=1, value=f"EXISTENTE {n + 1}")
        ws.cell(row=4 + n, column=2, value="PROVEEDOR")
        ws.cell(row=4 + n, column=4, value="kg")
        ws.cell(row=4 + n, column=5, value=10)
        ws.cell(row=4 + n, column=6, value="2028-01-01")
        ws.cell(row=4 + n, column=7, value=f"L-EXIST-{n + 1}")
    wb.save(ruta)
    return ruta


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.carpeta = pathlib.Path(self._tmp.name)
        self.invs = {}
        for tipo in ("mp", "mep", "pt"):
            ruta = libro(self.carpeta / inventarios.TIPOS[tipo]["archivo"])
            self.invs[tipo] = inventario.Inventario(ruta, tipo=tipo)
        self.con = bitacora.conectar(self.carpeta / "b.db")
        usuarios.crear(self.con, "Esteban", "clave-segura-1", "almacen")
        fila = usuarios.por_nombre(self.con, "Esteban")
        self.sesion = {"usuario_id": fila["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}
        self.estado = type("Estado", (), {})()
        self.estado.carpeta_datos = self.carpeta
        self.estado.con = self.con
        self.estado.archivo = "Inventario_MP.xlsx"
        self.estado.inventario = lambda tipo=None: self.invs[tipo or "mp"]

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()

    def alta(self, tipo, **datos):
        cuerpo = {"tipo": tipo, "descripcion": "MATERIAL DE PRUEBA", "lote": "LOTE-001",
                  "cantidad": 10, "proveedor": "PROVEEDOR X"}
        cuerpo.update(datos)
        return api.manejar(self.estado, "POST", "/api/alta", {}, cuerpo, self.sesion)


class TestMateriaPrima(Base):
    def test_exige_caducidad(self):
        codigo, cuerpo, _ = self.alta("mp", unidad="kg")
        self.assertEqual(codigo, 400)
        self.assertIn("falta la caducidad", cuerpo["error"])

    def test_ingresa_con_caducidad(self):
        codigo, cuerpo, _ = self.alta("mp", unidad="kg", caducidad="2028-05-01")
        self.assertEqual(codigo, 200, cuerpo)
        self.assertEqual(cuerpo["naturaleza"], "mp")
        self.assertEqual(cuerpo["archivo"], "Inventario_MP.xlsx")


class TestMaterialEmpaque(Base):
    def test_no_exige_caducidad(self):
        codigo, cuerpo, _ = self.alta("mep", unidad="pza")
        self.assertEqual(codigo, 200, cuerpo)
        self.assertIsNone(cuerpo["caducidad"])
        self.assertEqual(cuerpo["naturaleza"], "mep")
        self.assertEqual(cuerpo["archivo"], "Inventario_MEP.xlsx")

    def test_si_viene_caducidad_debe_ser_valida(self):
        codigo, cuerpo, _ = self.alta("mep", unidad="pza", caducidad="32/13/2028")
        self.assertEqual(codigo, 400)
        self.assertIn("caducidad", cuerpo["error"])

    def test_acepta_caducidad_cuando_la_hay(self):
        codigo, cuerpo, _ = self.alta("mep", unidad="pza", caducidad="2028-07-15")
        self.assertEqual(codigo, 200, cuerpo)
        self.assertEqual(cuerpo["caducidad"], "2028-07-15")

    def test_avisa_que_no_trae_caducidad(self):
        _, cuerpo, _ = self.alta("mep", unidad="pza")
        self.assertIn("sin caducidad", (cuerpo.get("aviso") or ""))


class TestProductoTerminado(Base):
    def test_exige_caducidad(self):
        codigo, cuerpo, _ = self.alta("pt", unidad="pza")
        self.assertEqual(codigo, 400)
        self.assertIn("falta la caducidad", cuerpo["error"])

    def test_avisa_si_falta_presentacion(self):
        _, cuerpo, _ = self.alta("pt", unidad="pza", caducidad="2027-12-01")
        self.assertIn("presentación", (cuerpo.get("aviso") or ""))

    def test_sin_aviso_cuando_trae_presentacion(self):
        _, cuerpo, _ = self.alta("pt", unidad="pza", caducidad="2027-12-01",
                                 presentacion="900 g caja 12", piezas="12")
        self.assertIsNone(cuerpo.get("aviso"))


class TestUnidades(Base):
    def test_avisa_una_unidad_inusual_sin_bloquear(self):
        # kg sí es válido en empaque (playo, tarima); «caja» no es unidad de inventario
        _, cuerpo, _ = self.alta("mep", unidad="caja")
        self.assertEqual(cuerpo["unidad"], "caja")
        self.assertIn("no es la habitual", (cuerpo.get("aviso") or ""))

    def test_acepta_litros_en_materia_prima(self):
        codigo, cuerpo, _ = self.alta("mp", unidad="lt", caducidad="2028-01-01")
        self.assertEqual(codigo, 200, cuerpo)
        self.assertEqual(cuerpo["unidad"], "lt")


class TestVaAlArchivoCorrecto(Base):
    def test_empaque_se_guarda_en_su_propio_excel(self):
        self.alta("mep", unidad="pza")
        ruta = self.carpeta / "Inventario_MEP.xlsx"
        hoja = openpyxl.load_workbook(ruta).worksheets[0]
        descripciones = [hoja.cell(row=f, column=1).value for f in range(4, hoja.max_row + 1)]
        self.assertIn("MATERIAL DE PRUEBA", descripciones)

    def test_el_mp_no_se_contamina(self):
        self.alta("mep", unidad="pza")
        descripciones = [f["descripcion"] for f in self.invs["mp"].filas]
        self.assertNotIn("MATERIAL DE PRUEBA", descripciones)


if __name__ == "__main__":
    unittest.main()
