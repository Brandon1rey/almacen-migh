import pathlib
import sys
import tempfile
import unittest

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from app import fechas, inventario, ubicaciones  # noqa: E402

ORIGEN = pathlib.Path(r"C:\Users\asegm\Desktop\Copia de 06 OCT 26 INVENTARIO MP.xlsx")


class TestUbicaciones(unittest.TestCase):
    def test_total(self):
        self.assertEqual(len(ubicaciones.CATALOGO), 383)   # 376 racks + 7 especiales
        self.assertEqual(len(ubicaciones.RACKS), 376)

    def test_rangos_del_pno(self):
        self.assertTrue(ubicaciones.existe("MPA-A-001"))
        self.assertTrue(ubicaciones.existe("MPA-A-040"))
        self.assertTrue(ubicaciones.existe("MPNA-A-041"))
        self.assertTrue(ubicaciones.existe("MPNA-B-136"))   # no está en el formato, sí en el PNO
        self.assertTrue(ubicaciones.existe("MPT-B-020"))
        self.assertTrue(ubicaciones.existe("MEP-E-220"))
        self.assertFalse(ubicaciones.existe("MPNA-B-137"))
        self.assertFalse(ubicaciones.existe("MEP-E-221"))

    def test_normaliza_y_zonas(self):
        self.assertEqual(ubicaciones.normalizar(" mpa-a-005 "), "MPA-A-005")
        self.assertEqual(ubicaciones.zona_de("mpna-a-042"), "MP-NO-ALERGENO")
        self.assertEqual(ubicaciones.zona_de("XXX"), "SIN UBICAR")


class TestFechas(unittest.TestCase):
    def test_formatos_reales_del_archivo(self):
        self.assertEqual(fechas.parsear("2027/agos/16"), ("2027-08-16", None))
        self.assertEqual(fechas.parsear("MAY/08/2027"), ("2027-05-08", None))
        self.assertEqual(fechas.parsear("12/26/2029"), ("2029-12-26", None))
        self.assertEqual(fechas.parsear("13/07/2027"), ("2027-07-13", None))

    def test_no_inventa(self):
        self.assertIsNone(fechas.parsear("26/17/07")[0])
        self.assertIsNone(fechas.parsear("2028/07/2028")[0])
        self.assertIsNone(fechas.parsear("05/07/2027")[0])   # ambigua
        self.assertIsNone(fechas.parsear("N/A")[0])

    def test_semaforo(self):
        self.assertEqual(fechas.semaforo("2026-10-01", hoy="2026-10-22")[1], "VENCIDO")
        self.assertEqual(fechas.semaforo("2026-11-10", hoy="2026-10-22")[1], "CRITICO")
        self.assertEqual(fechas.semaforo("2027-06-01", hoy="2026-10-22")[1], "OK")


@unittest.skipUnless(ORIGEN.exists(), "no está el Excel del área")
class TestInventarioReal(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ruta = pathlib.Path(self.tmp.name) / "Inventario_MP.xlsx"
        import shutil
        shutil.copy2(ORIGEN, self.ruta)
        self.inv = inventario.Inventario(self.ruta)

    def tearDown(self):
        self.tmp.cleanup()

    def test_lee_las_137_filas(self):
        self.assertEqual(len(self.inv.filas), 137)

    def test_agrega_la_columna_ubicacion(self):
        self.assertIn("ubicacion", self.inv.columnas)
        import openpyxl
        self.inv.guardar()
        ws = openpyxl.load_workbook(self.ruta).worksheets[0]
        self.assertEqual(ws.cell(row=3, column=self.inv.columnas["ubicacion"]).value, "UBICACIÓN")

    def test_no_pierde_datos_al_guardar(self):
        import openpyxl
        antes = [[c.value for c in fila] for fila in
                 openpyxl.load_workbook(self.ruta).worksheets[0].iter_rows(max_col=9)]
        self.inv.guardar()
        despues = [[c.value for c in fila] for fila in
                   openpyxl.load_workbook(self.ruta).worksheets[0].iter_rows(max_col=9)]
        self.assertEqual(antes, despues)

    def test_asignar_y_persistir(self):
        fila = self.inv.filas[0]["fila"]
        anterior, nueva = self.inv.asignar(fila, "mpa-a-005")
        self.assertIsNone(anterior)
        self.assertEqual(nueva, "MPA-A-005")
        self.inv.guardar()
        otra = inventario.Inventario(self.ruta)
        self.assertEqual(otra.por_fila(fila)["ubicacion"], "MPA-A-005")

    def test_rechaza_ubicacion_inexistente(self):
        with self.assertRaises(ValueError):
            self.inv.asignar(self.inv.filas[0]["fila"], "MPA-A-999")

    def test_limpiar_ubicacion(self):
        fila = self.inv.filas[0]["fila"]
        self.inv.asignar(fila, "MPA-A-005")
        anterior, nueva = self.inv.asignar(fila, None)
        self.assertEqual(anterior, "MPA-A-005")
        self.assertIsNone(nueva)

    def test_resumen_cuadra(self):
        r = self.inv.resumen()
        self.assertEqual(r["total"], 137)
        self.assertEqual(r["con_ubicacion"] + r["sin_ubicacion"], 137)
        self.assertEqual(sum(r["semaforo"].values()), 137)

    def test_kg_incluye_las_cantidades_en_formula(self):
        # 4 filas del archivo traen =suma sin valor calculado (Excel las muestra,
        # ningún lector de datos las ve). 71,576.4 + 639.0 = 72,215.4
        r = self.inv.resumen()
        self.assertAlmostEqual(r["kg"], 72215.4, places=1)
        self.assertEqual(r["con_formula"], 4)

    def test_evalua_formulas(self):
        self.assertEqual(inventario.evaluar_formula("=16.46+24.14+25"), 65.6)
        self.assertEqual(inventario.evaluar_formula("=150+22.3"), 172.3)
        with self.assertRaises(Exception):
            inventario.evaluar_formula("=SUMA(A1:A5)")

    def test_busqueda(self):
        self.assertTrue(self.inv.buscar("laca rojo"))
        self.assertTrue(self.inv.buscar("WPC"))
        self.assertEqual(self.inv.buscar("zzzzz"), [])


if __name__ == "__main__":
    unittest.main()
