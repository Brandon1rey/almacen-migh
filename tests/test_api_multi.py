"""API multi-inventario: cada petición opera sobre el tipo pedido y la bitácora no se cruza."""
import pathlib
import tempfile
import unittest

from app import api, bitacora, inventario, inventarios, usuarios

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)
from tests.test_multitipo import ENCABEZADOS, filas_de, libro


class EstadoDos:
    """Igual que app.server.Estado pero con dos inventarios de prueba."""

    def __init__(self, carpeta, con):
        self.carpeta_datos = pathlib.Path(carpeta)
        self.con = con
        self.archivo = "prueba"
        self._invs = {t: inventario.Inventario(self.carpeta_datos / f"{t}.xlsx", tipo=t)
                      for t in ("mp", "mep")}

    def inventario(self, tipo=None):
        tipo = inventarios.normalizar_tipo(tipo or "mp")
        if tipo not in self._invs:
            raise FileNotFoundError(tipo)
        return self._invs[tipo]

    def cargados(self):
        return list(self._invs.items())


class TestApiMulti(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)
        for tipo in ("mp", "mep"):
            libro(self.dir / f"{tipo}.xlsx", filas_de(tipo))
        self.con = bitacora.conectar(self.dir / "b.db")
        self.estado = EstadoDos(self.dir, self.con)
        usuarios.crear(self.con, "Esteban", "clave-segura-1", "almacen")
        fila = usuarios.por_nombre(self.con, "Esteban")
        self.sesion = {"usuario_id": fila["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()

    def api(self, metodo, ruta, params=None, cuerpo=None):
        return api.manejar(self.estado, metodo, ruta, params or {}, cuerpo or {}, self.sesion)

    def test_inventarios_lista_los_dos(self):
        codigo, res, _ = self.api("GET", "/api/inventarios")
        self.assertEqual(codigo, 200)
        claves = {i["clave"]: i for i in res["inventarios"]}
        self.assertTrue(claves["mp"]["cargado"])
        self.assertTrue(claves["mep"]["cargado"])
        self.assertEqual(claves["mep"]["lotes"], 2)

    def test_inventario_por_tipo(self):
        _, mp, _ = self.api("GET", "/api/inventario", {"tipo": "mp"})
        _, mep, _ = self.api("GET", "/api/inventario", {"tipo": "mep"})
        self.assertEqual(mp["tipo"], "mp")
        self.assertEqual(mep["tipo"], "mep")
        self.assertEqual(mp["filas"][0]["descripcion"], "WPC 80%")
        self.assertEqual(mep["filas"][0]["descripcion"], "BOLSA LAMINADA 2 KG")

    def test_mover_en_mep_no_toca_mp(self):
        fila_mep = self.estado.inventario("mep").filas[0]["fila"]
        codigo, res, _ = self.api("POST", "/api/mover", cuerpo={
            "fila": fila_mep, "destino": "MEP-B-010", "motivo": "reacomodo", "tipo": "mep",
            "cantidad": 500})
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["accion"], "PARCIAL")
        self.assertEqual(res["cantidad"], 500)

        mp = self.estado.inventario("mp")
        self.assertTrue(all(f["ubicacion"] in (None, "MPA-A-005", "MPNA-A-072") for f in mp.filas))
        movs_mep = bitacora.ultimos(self.con, 20, inventario="mep")
        self.assertEqual(len(movs_mep), 1)
        movs_mp = bitacora.ultimos(self.con, 20, inventario="mp")
        self.assertEqual(len(movs_mp), 0)

    def test_deshacer_usa_el_inventario_del_movimiento(self):
        fila_mep = self.estado.inventario("mep").filas[0]["fila"]
        _, res, _ = self.api("POST", "/api/mover", cuerpo={
            "fila": fila_mep, "destino": "MEP-B-010", "motivo": "reacomodo", "tipo": "mep"})
        mov = res["movimiento"]
        codigo, deshacer, _ = self.api("POST", "/api/deshacer", cuerpo={"movimiento": mov})
        self.assertEqual(codigo, 200, deshacer)
        self.assertEqual(deshacer["nueva"], "MEP-B-004")
        self.assertEqual(self.estado.inventario("mep").por_fila(fila_mep)["ubicacion"], "MEP-B-004")

    def test_descargar_f01_consolida(self):
        destino = api.descargar(self.estado, "f01")
        wb = openpyxl.load_workbook(destino)
        self.assertEqual(wb.sheetnames[0], "Portada")
        self.assertIn("MP", wb.sheetnames)
        self.assertIn("MEP", wb.sheetnames)

    def test_solo_admin_crea_usuarios(self):
        codigo, res, _ = self.api("POST", "/api/usuarios", cuerpo={
            "nombre": "No Debe", "password": "Migh.No.2026", "rol": "admin"})
        self.assertEqual(codigo, 403)
        self.assertIn("administrador", res["error"])

    def test_tipo_invalido_no_rompe(self):
        """Antes salía un ValueError sin manejar (500 en la cara del usuario): ahora es 400."""
        codigo, res, _ = self.api("GET", "/api/inventario", {"tipo": "xyz"})
        self.assertEqual(codigo, 400, res)
        self.assertIn("inventario desconocido", res["error"])


if __name__ == "__main__":
    unittest.main()
