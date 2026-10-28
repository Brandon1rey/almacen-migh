"""Códigos QR / Barcode: capa aditiva opcional sobre (tipo, fila) sin tocar el kardex ni SQLite."""
import datetime as dt
import io
import pathlib
import tempfile
import unittest
from contextlib import redirect_stdout
from unittest import mock

from app import api, bitacora, inventario, server, usuarios
from scripts import preparar_inventario

import openpyxl

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]


def crear_libro(ruta, filas=None, encabezados=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = encabezados or ENCABEZADOS
    for i, texto in enumerate(cols, start=1):
        ws.cell(row=3, column=i, value=texto)
    filas = filas if filas is not None else [
        ["WPC 80%", "VILHER", "Proteína", "kg", 100, "2028-04-23", "22260425002",
         "MIGH-MP114", "", "MPA-A-005"],
        ["LACA ROJO 40", "MEDINA", "Aditivo", "kg", 20, "2027-02-28", "260K1203",
         "MIGH-MP001", "", "MPNA-A-072"],
    ]
    for fila_idx, valores in enumerate(filas, start=4):
        for col_idx, valor in enumerate(valores, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=valor)
    wb.save(ruta)
    return ruta


def datos_alta(**cambios):
    base = {
        "descripcion": "LACTOSA 200 MESH",
        "proveedor": "LACTOSA MX",
        "categoria": "Carbohidrato",
        "unidad": "kg",
        "cantidad": 500,
        "caducidad": (dt.date.today() + dt.timedelta(days=400)).isoformat(),
        "lote": "LAC-2601",
        "erp": "MIGH-MP200",
        "ubicacion": "MPNA-A-050",
        "observaciones": "",
    }
    base.update(cambios)
    return base


class BaseQR(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)
        crear_libro(self.dir / "Inventario_MP.xlsx")
        crear_libro(self.dir / "Inventario_MEP.xlsx", [
            ["BOLSA LAMINADA 2 KG", "IMPRESOS", "Bolsa", "pza", 5000, "2029-01-31",
             "BL-2026-01", "MIGH-ME020", "", "MEP-B-004"],
        ])
        crear_libro(self.dir / "Inventario_PT.xlsx", [
            ["NUTRILIFE VAINILLA 900 G", "MIGH", "Polvo", "pza", 300, "2028-02-15",
             "PT-2601-A", "MIGH-PT003", "", "MPT-B-002"],
        ])
        self.estado = server.Estado(carpeta=self.dir)
        usuarios.crear(self.estado.con, "Esteban", "clave-segura-1", "almacen")
        usuarios.crear(self.estado.con, "Auditor", "clave-segura-1", "consulta")
        u_alm = usuarios.por_nombre(self.estado.con, "Esteban")
        u_con = usuarios.por_nombre(self.estado.con, "Auditor")
        self.sesion = {"usuario_id": u_alm["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}
        self.sesion_consulta = {"usuario_id": u_con["id"], "nombre": "Auditor", "rol": "consulta",
                                "debe_cambiar": 0}

    def tearDown(self):
        self.estado.con.close()
        self._tmp.cleanup()


class TestQRInventario(BaseQR):
    def test_01_mapa_columnas_reconoce_encabezados_qr_y_barcode(self):
        for encabezado in ("QR_ID", "BARCODE_ID", "Código QR", "Barcode",
                           "Código de barras", "Código de barra", "QR"):
            with self.subTest(encabezado=encabezado):
                ruta = self.dir / f"col_{inventario.normalizar(encabezado).replace(' ', '_')}.xlsx"
                crear_libro(ruta, [["WPC 80%", "VILHER", "Proteína", "kg", 50,
                                    "2028-04-23", "L-100", "SKU-1", "", "MPA-A-005", "QR-99"]],
                            encabezados=ENCABEZADOS + [encabezado])
                inv = inventario.Inventario(ruta, tipo="mp")
                self.assertIn("qr_id", inv.columnas)
                self.assertEqual(inv.filas[0]["qr_id"], "QR-99")

    def test_02_leer_expone_qr_id_y_convierte_numerico_a_texto_conservando_guiones(self):
        ruta = self.dir / "numerico.xlsx"
        crear_libro(ruta, [
            ["WPC 80%", "VILHER", "Proteína", "kg", 50, "2028-04-23", "L-1", "SKU-1",
             "", "MPA-A-005", 75010553108],
            ["LACTOSA", "VILHER", "Carbohidrato", "kg", 25, "2028-04-23", "L-2", "SKU-2",
             "", "MPA-A-006", None],
            ["COCOA", "VILHER", "Aditivo", "kg", 15, "2028-04-23", "L-3", "SKU-3",
             "", "MPA-A-007", "7501-055"],
        ], encabezados=ENCABEZADOS + ["QR_ID"])
        inv = inventario.Inventario(ruta, tipo="mp")
        self.assertEqual(inv.filas[0]["qr_id"], "75010553108")
        self.assertIsInstance(inv.filas[0]["qr_id"], str)
        self.assertIsNone(inv.filas[1]["qr_id"])
        self.assertEqual(inv.filas[2]["qr_id"], "7501-055")
        self.assertNotEqual(inventario.clave_qr("7501-055"), inventario.clave_qr("7501055"))

    def test_03_agregar_alta_crea_qr_id_lazy_y_persiste(self):
        inv = self.estado.inventario("mp")
        # Al solo abrir el libro no se crea la columna QR_ID (D6)
        self.assertNotIn("qr_id", inv.columnas)
        reg = inv.agregar_alta(datos_alta(qr_id="qr-0045"))
        self.assertIn("qr_id", inv.columnas)
        self.assertEqual(reg["qr_id"], "QR-0045")
        inv.guardar()
        releido = inventario.Inventario(inv.ruta, tipo="mp")
        self.assertEqual(releido.por_fila(reg["fila"])["qr_id"], "QR-0045")

    def test_04_alta_sin_qr_id_sigue_igual(self):
        inv = self.estado.inventario("mp")
        reg = inv.agregar_alta(datos_alta())
        self.assertIsNone(reg["qr_id"])
        self.assertEqual(reg["lote"], "LAC-2601")

    def test_05_alta_con_codigo_duplicado_en_mismo_libro_lanza_valueerror(self):
        inv = self.estado.inventario("mp")
        inv.agregar_alta(datos_alta(qr_id="QR-DUP-1"))
        with self.assertRaises(ValueError) as ctx:
            inv.agregar_alta(datos_alta(lote="LAC-2602", qr_id="qr-dup-1"))
        self.assertIn("QR-DUP-1", str(ctx.exception))

    def test_06_alta_con_codigo_duplicado_en_otro_libro_devuelve_409(self):
        inv_mep = self.estado.inventario("mep")
        inv_mep.asignar_qr(4, "QR-GLOBAL-1")
        inv_mep.guardar()
        codigo, res, _ = api.manejar(self.estado, "POST", "/api/alta", {},
                                     datos_alta(qr_id="qr-global-1"), self.sesion)
        self.assertEqual(codigo, 409, res)
        self.assertIn("MEP", res["error"])

    def test_07_codigo_en_renglon_anulado_se_puede_reutilizar(self):
        inv = self.estado.inventario("mp")
        inv.asignar_qr(4, "QR-REUSO")
        inv.anular(4, "consumido completo")
        inv.guardar()
        reg = inv.agregar_alta(datos_alta(qr_id="QR-REUSO"))
        self.assertEqual(reg["qr_id"], "QR-REUSO")

    def test_08_buscar_por_qr_ignora_renglones_sin_existencia(self):
        inv = self.estado.inventario("mp")
        inv.asignar_qr(4, "QR-ACTIVO")
        self.assertEqual(inv.buscar_por_qr("qr-activo")["fila"], 4)
        inv.ajustar_cantidad(4, 0)
        self.assertIsNone(inv.buscar_por_qr("QR-ACTIVO"))

    def test_09_api_qr_resuelve_activo_404_inexistente_y_400_vacio(self):
        inv_pt = self.estado.inventario("pt")
        inv_pt.asignar_qr(4, "PT-PALLET-01")
        inv_pt.guardar()

        c400, r400, _ = api.manejar(self.estado, "GET", "/api/qr", {}, {}, self.sesion)
        self.assertEqual(c400, 400, r400)

        c200, r200, _ = api.manejar(self.estado, "GET", "/api/qr",
                                    {"codigo": "pt-pallet-01"}, {}, self.sesion)
        self.assertEqual(c200, 200, r200)
        self.assertTrue(r200["encontrado"])
        self.assertEqual(r200["tipo"], "pt")
        self.assertEqual(r200["fila"], 4)
        self.assertEqual(r200["renglon"]["qr_id"], "PT-PALLET-01")

        c404, r404, _ = api.manejar(self.estado, "GET", "/api/qr",
                                    {"codigo": "NO-EXISTE"}, {}, self.sesion)
        self.assertEqual(c404, 404, r404)
        self.assertFalse(r404["encontrado"])

    def test_10_api_asignar_qr_bautiza_borra_y_valida_permisos(self):
        # Bautizar renglón existente
        c1, r1, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                {"tipo": "mp", "fila": 4, "qr_id": "qr-bautizo"}, self.sesion)
        self.assertEqual(c1, 200, r1)
        self.assertEqual(r1["qr_id"], "QR-BAUTIZO")
        releido = inventario.Inventario(self.dir / "Inventario_MP.xlsx", tipo="mp")
        self.assertEqual(releido.por_fila(4)["qr_id"], "QR-BAUTIZO")

        # Duplicado en otra fila -> 409
        c_dup, r_dup, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                      {"tipo": "mp", "fila": 5, "qr_id": "QR-BAUTIZO"}, self.sesion)
        self.assertEqual(c_dup, 409, r_dup)

        # Borrar etiqueta con cadena vacía
        c2, r2, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                {"tipo": "mp", "fila": 4, "qr_id": ""}, self.sesion)
        self.assertEqual(c2, 200, r2)
        self.assertIsNone(r2["qr_id"])
        releido2 = inventario.Inventario(self.dir / "Inventario_MP.xlsx", tipo="mp")
        self.assertIsNone(releido2.por_fila(4)["qr_id"])

        # Rol consulta -> 403
        c403, _, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                 {"tipo": "mp", "fila": 4, "qr_id": "QR-1"}, self.sesion_consulta)
        self.assertEqual(c403, 403)

        # Fila inexistente -> 400
        c400, _, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                 {"tipo": "mp", "fila": 999, "qr_id": "QR-1"}, self.sesion)
        self.assertEqual(c400, 400)

    def test_11_asignar_qr_con_excel_abierto_revierte_memoria(self):
        c_ini, _, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                  {"tipo": "mp", "fila": 4, "qr_id": "QR-INICIAL"}, self.sesion)
        self.assertEqual(c_ini, 200)
        inv = self.estado.inventario("mp")
        with mock.patch.object(inv, "guardar", side_effect=PermissionError("abierto")):
            codigo, res, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                         {"tipo": "mp", "fila": 4, "qr_id": "QR-NUEVO"},
                                         self.sesion)
        self.assertEqual(codigo, 409, res)
        self.assertEqual(inv.por_fila(4)["qr_id"], "QR-INICIAL")

    def test_12_parcialidad_conserva_qr_en_origen_y_no_duplica_en_hijo(self):
        c_ini, _, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                                  {"tipo": "mp", "fila": 4, "qr_id": "QR-ORIGEN-4"}, self.sesion)
        self.assertEqual(c_ini, 200)
        codigo, res, _ = api.manejar(self.estado, "POST", "/api/mover", {}, {
            "tipo": "mp", "fila": 4, "cantidad": 25, "destino": "MPA-A-006",
            "motivo": "reacomodo",
        }, self.sesion)
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["accion"], "PARCIAL")
        fila_hijo = res["fila_destino"]
        inv = self.estado.inventario("mp")
        self.assertEqual(inv.por_fila(4)["qr_id"], "QR-ORIGEN-4")
        self.assertIsNone(inv.por_fila(fila_hijo)["qr_id"])

        c_qr, r_qr, _ = api.manejar(self.estado, "GET", "/api/qr",
                                    {"codigo": "QR-ORIGEN-4"}, {}, self.sesion)
        self.assertEqual(c_qr, 200)
        self.assertEqual(r_qr["fila"], 4)

    def test_13_buscador_encuentra_por_qr_id_y_prefijo_qr(self):
        inv = self.estado.inventario("mp")
        inv.asignar_qr(5, "QR-0045")
        inv.guardar()
        por_texto = inv.buscar("QR-0045")
        self.assertEqual([f["fila"] for f in por_texto], [5])
        por_prefijo = inv.buscar("qr:QR-0045")
        self.assertEqual([f["fila"] for f in por_prefijo], [5])

        c_api, r_api, _ = api.manejar(self.estado, "GET", "/api/inventario",
                                      {"q": "qr:QR-0045", "tipo": "mp"}, {}, self.sesion)
        self.assertEqual(c_api, 200)
        self.assertEqual([f["fila"] for f in r_api["filas"]], [5])

    def test_14_preparar_inventario_reimportar_conserva_qr_y_avisa_claves_repetidas(self):
        destino = self.dir / "reimp.xlsx"
        origen = self.dir / "area_nuevo.xlsx"
        crear_libro(destino, [
            ["WPC 80%", "VILHER", "Proteína", "kg", 100, "2028-04-23", "LOTE-UNICO",
             "MIGH-MP114", "", "MPA-A-005"],
            ["LACTOSA", "VILHER", "Carbohidrato", "kg", 60, "2028-04-23", "LOTE-PARTIDO",
             "MIGH-MP200", "", "MPNA-A-050"],
            ["LACTOSA", "VILHER", "Carbohidrato", "kg", 40, "2028-04-23", "LOTE-PARTIDO",
             "MIGH-MP200", "parcialidad", "MPNA-A-051"],
        ])
        inv_ant = inventario.Inventario(destino, tipo="mp")
        inv_ant.asignar_qr(4, "QR-UNICO-1")
        inv_ant.asignar_qr(5, "QR-PARTIDO-1")
        inv_ant.guardar()

        # El archivo nuevo del área no trae columna QR_ID
        crear_libro(origen, [
            ["WPC 80%", "VILHER", "Proteína", "kg", 95, "2028-04-23", "LOTE-UNICO",
             "MIGH-MP114", "", ""],
            ["LACTOSA", "VILHER", "Carbohidrato", "kg", 100, "2028-04-23", "LOTE-PARTIDO",
             "MIGH-MP200", "", ""],
        ])
        salida = io.StringIO()
        with redirect_stdout(salida):
            preparar_inventario.main([
                "--tipo", "mp", "--origen", str(origen), "--destino", str(destino), "--reimportar"
            ])
        texto = salida.getvalue()
        self.assertIn("Códigos QR/Barcode conservados al reimportar: 1", texto)
        self.assertIn("AVISO: 1 clave(s) con varias filas tenían código", texto)

        inv_nuevo = inventario.Inventario(destino, tipo="mp")
        self.assertEqual(inv_nuevo.por_fila(4)["qr_id"], "QR-UNICO-1")
        self.assertIsNone(inv_nuevo.por_fila(5)["qr_id"])

    def test_15_flujo_manual_sin_codigos_y_esquema_sqlite_intactos(self):
        # Ubicar, mover y limpiar sin qr_id funcionan exactamente igual
        c_ub, _, _ = api.manejar(self.estado, "POST", "/api/ubicar", {},
                                 {"tipo": "mp", "fila": 4, "ubicacion": "MPA-A-007"}, self.sesion)
        self.assertEqual(c_ub, 200)
        c_mv, _, _ = api.manejar(self.estado, "POST", "/api/mover", {},
                                 {"tipo": "mp", "fila": 4, "destino": "MPA-A-008",
                                  "motivo": "reacomodo"}, self.sesion)
        self.assertEqual(c_mv, 200)
        c_lm, _, _ = api.manejar(self.estado, "POST", "/api/limpiar", {},
                                 {"tipo": "mp", "fila": 4}, self.sesion)
        self.assertEqual(c_lm, 200)

        # bitacora.db no tiene migraciones ni columnas de QR
        cols_mov = [r[1] for r in self.estado.con.execute("PRAGMA table_info(movimiento)")]
        self.assertNotIn("qr_id", cols_mov)

    def test_16_duplicado_imposible_entre_libros_cargados(self):
        c1, _, _ = api.manejar(self.estado, "POST", "/api/asignar_qr", {},
                               {"tipo": "mp", "fila": 4, "qr_id": "ETIQUETA-UNICA"}, self.sesion)
        self.assertEqual(c1, 200)
        for otro_tipo in ("mep", "pt"):
            c_asig, r_asig, _ = api.manejar(
                self.estado, "POST", "/api/asignar_qr", {},
                {"tipo": otro_tipo, "fila": 4, "qr_id": "etiqueta-unica"}, self.sesion
            )
            self.assertEqual(c_asig, 409, r_asig)


if __name__ == "__main__":
    unittest.main()
