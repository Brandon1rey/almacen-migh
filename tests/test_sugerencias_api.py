"""El endpoint de ubicaciones devuelve sugerencias ordenadas, con la ocupación real del inventario."""
import pathlib
import tempfile
import unittest

from app import api, bitacora, inventario, usuarios

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]


def libro(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=3, column=i, value=texto)
    filas = [("LECHE EN POLVO", "L-100", 500, "MPA-A-001"),
             ("AZUCAR", "L-200", 300, "MPA-A-002"),
             ("WPC 80%", "L-300", 200, None)]
    for n, (desc, lote, cantidad, ubicacion) in enumerate(filas, start=4):
        ws.cell(row=n, column=1, value=desc)
        ws.cell(row=n, column=2, value="PROVEEDOR")
        ws.cell(row=n, column=4, value="kg")
        ws.cell(row=n, column=5, value=cantidad)
        ws.cell(row=n, column=6, value="2028-01-01")
        ws.cell(row=n, column=7, value=lote)
        if ubicacion:
            ws.cell(row=n, column=10, value=ubicacion)
    wb.save(ruta)
    return ruta


class TestSugerenciasAPI(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        carpeta = pathlib.Path(self._tmp.name)
        self.inv = inventario.Inventario(libro(carpeta / "inv.xlsx"), tipo="mp")
        self.con = bitacora.conectar(carpeta / "b.db")
        usuarios.crear(self.con, "Esteban", "clave-segura-1", "almacen")
        fila = usuarios.por_nombre(self.con, "Esteban")
        self.sesion = {"usuario_id": fila["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}
        self.estado = type("Estado", (), {})()
        self.estado.carpeta_datos = carpeta
        self.estado.con = self.con
        self.estado.archivo = "inv.xlsx"
        self.estado.inventario = lambda tipo=None: self.inv

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()

    def pedir(self, consulta):
        codigo, cuerpo, _ = api.manejar(self.estado, "GET", "/api/ubicaciones", consulta, {},
                                        self.sesion)
        self.assertEqual(codigo, 200, cuerpo)
        return cuerpo

    def test_devuelve_sugerencias_ordenadas(self):
        cuerpo = self.pedir({"q": "MPA-A-00", "tipo": "mp"})
        codigos = [s["codigo"] for s in cuerpo["sugerencias"]]
        self.assertTrue(codigos[0].startswith("MPA-A-00"))
        self.assertNotIn("MEP-B-001", codigos)

    def test_marca_la_ocupacion_real(self):
        cuerpo = self.pedir({"q": "MPA-A-00", "tipo": "mp"})
        por_codigo = {s["codigo"]: s for s in cuerpo["sugerencias"]}
        self.assertEqual(por_codigo["MPA-A-001"]["lotes"], 1)
        self.assertFalse(por_codigo["MPA-A-001"]["libre"])
        self.assertTrue(por_codigo["MPA-A-005"]["libre"])
        self.assertEqual(cuerpo["ocupacion"]["MPA-A-001"], 1)

    def test_el_primer_lugar_es_un_rack_libre(self):
        cuerpo = self.pedir({"q": "MPA-A-00", "tipo": "mp"})
        self.assertTrue(cuerpo["sugerencias"][0]["libre"])

    def test_usa_el_material_del_lote_para_los_alergenos(self):
        cuerpo = self.pedir({"q": "", "tipo": "mp", "fila": 4, "limite": 400})   # LECHE EN POLVO
        self.assertEqual(cuerpo["material"], "LECHE EN POLVO")
        por_codigo = {s["codigo"]: s for s in cuerpo["sugerencias"]}
        # la zona de alérgenos va primero y no lleva aviso
        self.assertTrue(cuerpo["sugerencias"][0]["codigo"].startswith("MPA-A"))
        self.assertIsNone(cuerpo["sugerencias"][0]["aviso"])
        # el resto de zonas avisa que no admiten alérgenos
        self.assertIn("no admite alérgenos", por_codigo["MPNA-A-041"]["aviso"])

    def test_material_explicito(self):
        cuerpo = self.pedir({"q": "", "tipo": "mp", "material": "AZUCAR"})
        self.assertIsNone(cuerpo["sugerencias"][0]["aviso"])

    def test_no_ofrece_produccion_para_asignar(self):
        cuerpo = self.pedir({"q": "PRODUC", "tipo": "mp"})
        self.assertEqual([s for s in cuerpo["sugerencias"] if s["codigo"] == "PRODUCCION"], [])

    def test_respeta_el_limite(self):
        cuerpo = self.pedir({"q": "", "tipo": "mp", "limite": 3})
        self.assertLessEqual(len(cuerpo["sugerencias"]), 3)

    def test_sigue_entregando_el_catalogo_completo(self):
        cuerpo = self.pedir({"tipo": "mp"})
        self.assertIn("MPNA-A-041", cuerpo["codigos"])
        self.assertIn("PRODUCCION", cuerpo["codigos"])


if __name__ == "__main__":
    unittest.main()
