"""Catálogo de insumos (MP y ME) y salidas a producción: pruebas de las dos piezas."""
import datetime as dt
import os
import pathlib
import tempfile
import unittest

from app import api, bitacora, catalogo, inventario, usuarios

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)

ENCABEZADOS_INV = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
                   "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]
ENCABEZADOS_CAT = ["Clave de Insumo", "Descripción", "Tipo", "Unidad de Medida",
                   "Codigo de Barra", "Máximo", "Mínimo", "Fabricante"]


def libro_inventario(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS_INV, start=1):
        ws.cell(row=3, column=i, value=texto)
    ws.cell(row=4, column=1, value="WPC 80%")
    ws.cell(row=4, column=2, value="VILHER")
    ws.cell(row=4, column=4, value="kg")
    ws.cell(row=4, column=5, value=1000)
    ws.cell(row=4, column=6, value="2028-04-23")
    ws.cell(row=4, column=7, value="22260425002")
    ws.cell(row=4, column=8, value="MIGH-MP114")
    ws.cell(row=4, column=10, value="MPA-A-005")
    wb.save(ruta)
    return ruta


def libro_catalogo(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS_CAT, start=1):
        ws.cell(row=1, column=i, value=texto)
    filas = [
        ("MIGH-MP114", "WPC 80%", "Excipiente", "Kg", "", 5000, 100, "VILHER"),
        ("MIGH-MP001", "ACAIVERRY EXT", "Excipiente", "Kg", "", 1000, 50, "ACAIVERRY"),
        ("MIGH-MT010", "BOLSA LAMINADA 2 KG", "Empaque", "piezas", "", 90000, 5000, "IMPRESOS DEL NORTE"),
        ("MIGH-MT011", "TAPA ROSCA 43 MM", "Empaque", "piezas", "", 60000, 2000, "PLASTICOS SUR"),
        ("MIGH-LAB001", "REACTIVO X", "Principio Activo", "ml", "", 10, 1, "SIGMA"),
    ]
    for n, fila in enumerate(filas, start=2):
        for i, valor in enumerate(fila, start=1):
            ws.cell(row=n, column=i, value=valor)
    wb.save(ruta)
    return ruta


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)
        self.ruta = libro_inventario(self.dir / "inv.xlsx")
        self.ruta_cat = libro_catalogo(self.dir / "catalogo.xlsx")
        self.destino_cat = self.dir / "catalogo_insumos.json"
        catalogo.importar(self.ruta_cat, self.destino_cat)
        catalogo.RUTA = self.destino_cat
        catalogo._CACHE.update({"ruta": None, "mtime": None, "datos": None})
        self.inv = inventario.Inventario(self.ruta, tipo="mp")
        self.con = bitacora.conectar(self.dir / "b.db")
        usuarios.crear(self.con, "Esteban", "clave-segura-1", "almacen")
        fila = usuarios.por_nombre(self.con, "Esteban")
        self.sesion = {"usuario_id": fila["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}
        self.estado = type("Estado", (), {})()
        self.estado.carpeta_datos = self.dir
        self.estado.con = self.con
        self.estado.archivo = "inv.xlsx"
        self.estado.inventario = lambda tipo=None: self.inv

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()


class TestCatalogo(Base):
    def test_importa_el_catalogo(self):
        datos = catalogo.cargar(self.destino_cat)
        self.assertEqual(len(datos["insumos"]), 5)
        self.assertEqual(datos["por_tipo"]["Empaque"], 2)

    def test_el_catalogo_sigue_a_la_carpeta_de_datos(self):
        """El catálogo es el de ESTA instalación, no el del repositorio.

        demo.cmd y cualquier servidor arrancado con otra carpeta de datos tienen que leer
        su propio catálogo; si no, la validación del alta usaría datos de otra instalación.
        """
        self.assertEqual(catalogo.ruta(self.dir), self.destino_cat)
        self.assertTrue(catalogo.ruta(self.dir).exists())
        self.assertEqual(catalogo.ruta(), catalogo.RUTA)
        anterior = os.environ.get("ALMACEN_CATALOGO")
        os.environ["ALMACEN_CATALOGO"] = str(self.dir / "otro.json")
        try:      # la variable de entorno manda sobre todo
            self.assertEqual(catalogo.ruta(self.dir), self.dir / "otro.json")
        finally:
            if anterior is None:
                os.environ.pop("ALMACEN_CATALOGO", None)
            else:
                os.environ["ALMACEN_CATALOGO"] = anterior

    def test_busca_por_clave_y_descripcion(self):
        datos = catalogo.cargar(self.destino_cat)
        self.assertEqual(catalogo.buscar(datos, "MIGH-MP114")[0]["descripcion"], "WPC 80%")
        self.assertEqual(len(catalogo.buscar(datos, "bolsa")), 1)

    def test_clasifica_el_tipo_de_inventario(self):
        datos = catalogo.cargar(self.destino_cat)
        self.assertEqual(catalogo.para_inventario(datos, "mp")[0]["tipo"], "Excipiente")

    def test_guarda_todos_los_datos_del_insumo(self):
        """Clave, descripción, tipo, unidad, barra, máximo, mínimo y fabricante del ERP."""
        datos = catalogo.cargar(self.destino_cat)
        insumo = catalogo.por_sku(datos, "MIGH-MP114")
        self.assertEqual(insumo["descripcion"], "WPC 80%")
        self.assertEqual(insumo["tipo"], "Excipiente")
        self.assertEqual(insumo["unidad"], "kg")
        self.assertEqual(insumo["fabricante"], "VILHER")
        self.assertEqual(insumo["maximo"], 5000.0)
        self.assertEqual(insumo["minimo"], 100.0)
        self.assertIsNone(insumo["barra"])
        self.assertEqual(insumo["inventario"], "mp")

    def test_unidades_y_tipos_del_erp(self):
        """El ERP trae «piezas», «L» y producto terminado: se normalizan sin inventar nada."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta = pathlib.Path(tmp) / "catalogo.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(ENCABEZADOS_CAT)
            ws.append(["MIGH-MT999", "BOLSA X", "Empaque", "piezas", "750123", 100, 5, "IMPRESOS"])
            ws.append(["MIGH-PT999", "PRODUCTO Y", "Producto Terminado", "L", "", 10, 1, "PLANTA"])
            wb.save(ruta)
            datos = catalogo.importar(ruta, pathlib.Path(tmp) / "cat.json")
        bolsa = catalogo.por_sku(datos, "MIGH-MT999")
        self.assertEqual(bolsa["unidad"], "pza")
        self.assertEqual(bolsa["barra"], "750123")
        self.assertEqual(bolsa["inventario"], "mep")
        producto = catalogo.por_sku(datos, "MIGH-PT999")
        self.assertEqual(producto["unidad"], "lt")
        self.assertEqual(producto["inventario"], "pt")
        tipos = {i["tipo"] for i in catalogo.para_inventario(datos, "mep")}
        self.assertEqual(tipos, {"Empaque"})

    def test_por_sku(self):
        datos = catalogo.cargar(self.destino_cat)
        self.assertEqual(catalogo.por_sku(datos, "migh-mt010")["unidad"], "pza")


class TestSalidaProduccion(Base):
    def api(self, ruta, cuerpo):
        return api.manejar(self.estado, "POST", ruta, {}, cuerpo, self.sesion)

    def test_salida_parcial_baja_la_cantidad(self):
        codigo, res, _ = self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION",
                                                 "motivo": "surtido", "cantidad": 300,
                                                 "ref_doc": "OP-1234"})
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["accion"], "SALIDA")
        self.assertEqual(res["anterior"], "MPA-A-005")
        self.assertEqual(res["nueva"], "PRODUCCION")
        self.assertEqual(self.inv.por_fila(4)["cantidad"], 700)
        self.assertEqual(self.inv.por_fila(4)["ubicacion"], "MPA-A-005")

    def test_salida_total_deja_el_renglon_en_cero(self):
        codigo, res, _ = self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION",
                                                 "motivo": "surtido"})
        self.assertEqual(codigo, 200, res)
        registro = self.inv.por_fila(4)
        self.assertEqual(registro["cantidad"], 0)
        self.assertIsNone(registro["ubicacion"])
        self.assertIn("producción", (registro["observaciones"] or "").lower())

    def test_la_bitacora_registra_la_salida(self):
        self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION", "motivo": "surtido",
                                "cantidad": 300, "ref_doc": "OP-1234"})
        mov = bitacora.ultimos(self.con, 5, inventario="mp")[0]
        self.assertEqual(mov["accion"], "SALIDA")
        self.assertEqual(mov["ubicacion_destino"], "PRODUCCION")
        self.assertEqual(mov["ref_doc"], "OP-1234")
        self.assertAlmostEqual(mov["cantidad"], 300)

    def test_no_deja_sacar_mas_de_lo_que_hay(self):
        codigo, res, _ = self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION",
                                                 "motivo": "surtido", "cantidad": 5000})
        self.assertEqual(codigo, 400)
        self.assertIn("no puedes mover", res["error"])

    def test_deshacer_la_salida_devuelve_los_kilos(self):
        _, res, _ = self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION",
                                            "motivo": "surtido", "cantidad": 300})
        codigo, deshacer, _ = self.api("/api/deshacer", {"movimiento": res["movimiento"]})
        self.assertEqual(codigo, 200, deshacer)
        self.assertEqual(self.inv.por_fila(4)["cantidad"], 1000)

    def test_deshacer_una_salida_total_devuelve_la_ubicacion(self):
        _, res, _ = self.api("/api/mover", {"fila": 4, "destino": "PRODUCCION",
                                            "motivo": "surtido"})
        self.api("/api/deshacer", {"movimiento": res["movimiento"]})
        registro = self.inv.por_fila(4)
        self.assertEqual(registro["cantidad"], 1000)
        self.assertEqual(registro["ubicacion"], "MPA-A-005")

    def test_produccion_aparece_en_las_ubicaciones_de_la_interfaz(self):
        codigo, res, _ = api.manejar(self.estado, "GET", "/api/ubicaciones", {}, {}, self.sesion)
        self.assertEqual(codigo, 200)
        self.assertIn("PRODUCCION", res["codigos"])


class TestAltaConCatalogo(Base):
    def api(self, ruta, cuerpo):
        return api.manejar(self.estado, "POST", ruta, {}, cuerpo, self.sesion)

    def test_completa_descripcion_unidad_y_proveedor(self):
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-MP114", "lote": "L-2601",
                                                "cantidad": 100, "caducidad": "2028-01-01"})
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["descripcion"], "WPC 80%")
        self.assertEqual(res["unidad"], "kg")
        self.assertIsNone(res["aviso"])

    def test_acepta_clave_que_no_esta_en_el_catalogo(self):
        """El almacén no puede frenar una recepción: se guarda tal como se capturó, con aviso."""
        antes = len(self.inv.filas)
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-XX999", "descripcion": "MATERIAL RARO",
                                                "lote": "L-2602", "cantidad": 10, "unidad": "kg",
                                                "proveedor": "X", "caducidad": "2028-01-01"})
        self.assertEqual(codigo, 200, res)
        self.assertIn("no está en el catálogo", res["aviso"])
        self.assertEqual(len(self.inv.filas), antes + 1)
        fila = [f for f in self.inv.filas if f["lote"] == "L-2602"][0]
        self.assertEqual(fila["descripcion"], "MATERIAL RARO")
        self.assertEqual(fila["erp"], "MIGH-XX999")

    def test_acepta_material_de_otro_inventario_y_lo_pasa_a_su_libro(self):
        """La clave MIGH-MT010 es de Empaque: se acepta y el renglón cae en el libro correcto."""
        ruta_mep = self.dir / "Inventario_MEP.xlsx"
        ruta_mep.write_bytes(self.ruta.read_bytes())        # mismo formato, otro libro
        inv_mep = inventario.Inventario(ruta_mep, tipo="mep")
        self.addCleanup(inv_mep.wb.close)
        original = self.estado.inventario
        self.estado.inventario = (lambda tipo=None:
                                  inv_mep if (tipo or "mp") == "mep" else original(tipo))
        self.addCleanup(setattr, self.estado, "inventario", original)

        antes_mp = len(self.inv.filas)
        antes_mep = len(inv_mep.filas)
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-MT010", "lote": "L-2603",
                                                "cantidad": 100, "caducidad": "2028-01-01"})
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["naturaleza"], "mep", res)
        self.assertIn("Empaque", res["aviso"])
        self.assertEqual(len(self.inv.filas), antes_mp, "no debe quedar en el libro equivocado")
        self.assertEqual(len(inv_mep.filas), antes_mep + 1)
        self.assertTrue([f for f in inv_mep.filas if f["lote"] == "L-2603"])

    def test_la_ubicacion_sigue_siendo_lo_unico_obligatorio_del_catalogo(self):
        """Aunque el insumo no exista en el ERP, una ubicación inventada se rechaza."""
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-XX999", "descripcion": "RARO",
                                                "lote": "L-2610", "cantidad": 5,
                                                "proveedor": "X", "caducidad": "2028-01-01",
                                                "ubicacion": "RACK-INVENTADO"})
        self.assertEqual(codigo, 400, res)
        self.assertIn("no existe en el catálogo", res["error"])

    def test_avisa_si_el_insumo_no_trae_tipo(self):
        """El ERP real trae 7 insumos sin tipo: se aceptan, pero se avisa para revisar."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta = pathlib.Path(tmp) / "catalogo.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(ENCABEZADOS_CAT)
            ws.append(["MIGH-MH-001", "GUANTE DE NITRILO", "", "pieza", "", 10, 1, "PENDIENTE"])
            wb.save(ruta)
            catalogo.importar(ruta, self.destino_cat)
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-MH-001", "lote": "L-2609",
                                                "cantidad": 10, "caducidad": "2028-01-01",
                                                "proveedor": "PENDIENTE"})
        self.assertEqual(codigo, 200, res)
        self.assertIn("no trae tipo", res["aviso"])
        self.assertEqual(res["descripcion"], "GUANTE DE NITRILO")

    def test_avisa_si_la_descripcion_no_coincide(self):
        codigo, res, _ = self.api("/api/alta", {"erp": "MIGH-MP114", "descripcion": "OTRA COSA",
                                                "lote": "L-2604", "cantidad": 5, "unidad": "kg",
                                                "proveedor": "VILHER", "caducidad": "2028-01-01"})
        self.assertEqual(codigo, 200, res)
        self.assertIn("WPC 80%", res["aviso"])


if __name__ == "__main__":
    unittest.main()
