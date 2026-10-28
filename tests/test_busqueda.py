"""Búsqueda: por lote (con o sin guiones), por campo, por ubicación y por varias palabras."""
import pathlib
import tempfile
import unittest

from app import inventario

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]


def libro(ruta, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=3, column=i, value=texto)
    for n, fila in enumerate(filas, start=4):
        for i, valor in enumerate(fila, start=1):
            ws.cell(row=n, column=i, value=valor)
    wb.save(ruta)
    return ruta


FILAS = [
    ["BOLSA LAMINADA 2 KG", "IMPRESOS DEL NORTE", "Bolsa", "pza", 24000, "2029-01-31",
     "BL-2026-01", "MIGH-ME003", "", "MEP-B-004"],
    ["WPC 80%", "VILHER", "Proteína", "kg", 1458, "2028-04-23", "22260425002", "MIGH-MP114",
     "", "MPA-A-005"],
    ["TAPA ROSCA 43 MM", "PLASTICOS SUR", "Tapa", "pza", 50000, "2030-08-31", "TR43-B-2601",
     "MIGH-ME020", "", "MEP-C-030"],
    ["BOLSA LAMINADA 900 G", "IMPRESOS DEL NORTE", "Bolsa", "pza", 96000, "2029-05-30",
     "BL900-2601", "MIGH-ME002", "no confundir con BL-2026-01", None],
]


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.ruta = libro(pathlib.Path(self._tmp.name) / "inv.xlsx", FILAS)
        self.inv = inventario.Inventario(self.ruta, tipo="mep")

    def tearDown(self):
        self._tmp.cleanup()

    def lotes(self, texto):
        return [f["lote"] for f in self.inv.buscar(texto)]


class TestPorLote(Base):
    def test_lote_exacto(self):
        self.assertEqual(self.lotes("BL-2026-01"), ["BL-2026-01"])

    def test_lote_sin_guiones(self):
        self.assertEqual(self.lotes("bl202601"), ["BL-2026-01"])

    def test_lote_con_espacios(self):
        self.assertEqual(self.lotes("bl 2026 01"), ["BL-2026-01"])

    def test_lote_parcial(self):
        self.assertIn("BL-2026-01", self.lotes("2026"))

    def test_lote_numerico_largo(self):
        self.assertEqual(self.lotes("22260425002"), ["22260425002"])

    def test_con_la_palabra_lote(self):
        self.assertEqual(self.lotes("lote 22260425002"), ["22260425002"])

    def test_campo_explicito(self):
        self.assertIn("BL-2026-01", self.lotes("lote:bl202601"))

    def test_lote_exacto_va_primero(self):
        # «BL-2026-01» también aparece escrito en las observaciones del lote BL900-2601
        encontrados = self.lotes("BL-2026-01")
        self.assertEqual(encontrados[0], "BL-2026-01")

    def test_no_confunde_parciales_de_otro_lote(self):
        self.assertNotIn("BL900-2601", self.lotes("BL-2026-01"))


class TestPorOtrosCampos(Base):
    def test_ubicacion(self):
        self.assertEqual(self.lotes("MEP-B-004"), ["BL-2026-01"])

    def test_campo_ubicacion(self):
        self.assertEqual(self.lotes("ubicacion:mep-c-030"), ["TR43-B-2601"])

    def test_erp(self):
        self.assertEqual(self.lotes("MIGH-ME020"), ["TR43-B-2601"])

    def test_proveedor(self):
        self.assertEqual(len(self.lotes("plasticos")), 1)

    def test_varias_palabras_en_cualquier_orden(self):
        self.assertEqual(self.lotes("bolsa 2 kg"), self.lotes("2 kg bolsa"))
        self.assertEqual(len(self.lotes("bolsa 2 kg")), 1)

    def test_acentos_y_mayusculas(self):
        self.assertEqual(self.lotes("PROTEINA"), self.lotes("proteína"))

    def test_sin_resultados(self):
        self.assertEqual(self.lotes("material que no existe"), [])


class TestSugerencias(Base):
    def test_sugiere_por_lote_parecido(self):
        sugerencias = self.inv.sugerencias("BL2026")
        self.assertTrue(sugerencias)
        self.assertTrue(any("BL" in s for s in sugerencias))

    def test_sin_texto_no_sugiere(self):
        self.assertEqual(self.inv.sugerencias(""), [])


class TestIndice(Base):
    def test_la_busqueda_por_ubicacion_ve_los_cambios(self):
        fila = next(f for f in self.inv.filas if f["lote"] == "BL-2026-01")["fila"]
        self.inv.asignar(fila, "MEP-C-030")
        self.assertIn("BL-2026-01", self.lotes("MEP-C-030"))
        self.assertNotIn("BL-2026-01", self.lotes("MEP-B-004"))


if __name__ == "__main__":
    unittest.main()
