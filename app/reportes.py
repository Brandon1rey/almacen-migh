"""Salidas en Excel: inventario con ubicación, reporte F01, tabla macro y bitácora."""
import datetime as dt
import os
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

from app import fechas, ubicaciones

RAIZ = pathlib.Path(__file__).resolve().parents[1]
SALIDAS = pathlib.Path(os.environ.get("ALMACEN_SALIDAS") or (RAIZ / "data" / "salidas"))
AZUL = PatternFill("solid", fgColor="1F3B73")
GRIS = PatternFill("solid", fgColor="D9D9D9")
AMBAR = PatternFill("solid", fgColor="FFF2CC")
ROJO = PatternFill("solid", fgColor="F8D7DA")
BORDE = Border(*[Side(style="thin", color="808080")] * 4)
BLANCO_BOLD = Font(bold=True, size=9, color="FFFFFF")
NEGRITA = Font(bold=True, size=9)
NORMAL = Font(size=9)


def _marca():
    return dt.datetime.now().strftime("%Y%m%d-%H%M")


def _ruta(nombre):
    try:
        if os.environ.get("VERCEL") or os.environ.get("VERCEL_ENV"):
            raise OSError("entorno serverless de solo lectura")
        SALIDAS.mkdir(parents=True, exist_ok=True)
        return SALIDAS / nombre
    except OSError:
        import tempfile
        alterna = pathlib.Path(tempfile.gettempdir()) / "almacen-salidas"
        alterna.mkdir(parents=True, exist_ok=True)
        return alterna / nombre


def _encabezado_f01(ws, version, total_filas=0):
    ws["A1"] = "COMERCIALIZADORA MIGH DE MÉXICO S.A. DE C.V."
    ws["A1"].font = Font(bold=True, size=10)
    ws["A2"] = "TIPO DE DOCUMENTO: FORMATO"
    ws["A2"].font = NEGRITA
    ws["A3"] = "TÍTULO DEL DOCUMENTO: REPORTE DE INVENTARIOS"
    ws["A3"].font = Font(bold=True, size=11)
    ws["A4"] = "CÓDIGO DEL DOCUMENTO: F01-ALM-PNO-002"
    ws["A4"].font = NEGRITA
    ws["A5"] = f"VERSIÓN: {version}"
    ws["A5"].font = NEGRITA
    pagina = "PÁGINA: ____ de ____" if total_filas > 30 else "PÁGINA 1 de 1"
    ws["C5"] = f"FECHA: {dt.date.today().isoformat()} · {pagina}"
    ws["C5"].font = NEGRITA


def _tabla(ws, fila_inicio, encabezados):
    for i, texto in enumerate(encabezados, start=1):
        c = ws.cell(row=fila_inicio, column=i, value=texto)
        c.font = BLANCO_BOLD
        c.fill = AZUL
        c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center")
    return fila_inicio + 1


def _anchos(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _firmas(ws, fila):
    ws.cell(row=fila, column=1, value="Elaboró").font = NEGRITA
    ws.cell(row=fila, column=3, value="Revisó").font = NEGRITA
    ws.cell(row=fila, column=5, value="Autorizó").font = NEGRITA
    ws.cell(row=fila + 1, column=1, value="Esteban Soto — Líder de Almacén").font = NORMAL
    ws.cell(row=fila + 1, column=3, value="Galilea Medina — Jefe de Almacén").font = NORMAL
    ws.cell(row=fila + 1, column=5,
            value="Maria Elena Becerra — Jefe de Aseguramiento de Calidad").font = NORMAL
    ws.cell(row=fila + 3, column=1, value="Firma / Fecha: ____________________").font = NORMAL
    ws.cell(row=fila + 3, column=3, value="Firma / Fecha: ____________________").font = NORMAL
    ws.cell(row=fila + 3, column=5, value="Firma / Fecha: ____________________").font = NORMAL
    return fila + 4


def _hoja_inventario(ws, inv, version):
    """Escribe el reporte de un inventario en una hoja (formato F01 vigente)."""
    _encabezado_f01(ws, version, len(inv.filas))
    ws.cell(row=2, column=5, value=f"INVENTARIO: {inv.nombre}").font = NEGRITA

    if version == "3":
        encabezados = ["UBICACIÓN", "MATERIAL", "CANTIDAD", "UNIDAD", "LOTE",
                       "CADUCIDAD", "DÍAS"]
    else:
        encabezados = ["UBICACIÓN", "MATERIAL", "CANTIDAD", "LOTE"]
    fila = _tabla(ws, 6, encabezados)

    def valores_de(f, es_sin_ubicar=False):
        ubicacion = "(sin ubicación)" if es_sin_ubicar else f["ubicacion"]
        if version == "3":
            return [ubicacion, f["descripcion"], f["cantidad"], f["unidad"], f["lote"],
                    f["caducidad"], f["dias_restantes"]]
        return [ubicacion, f["descripcion"], f["cantidad"], f["lote"]]

    orden = {"MP-ALERGENO": 1, "MP-NO-ALERGENO": 2, "PT": 3, "EM": 4, "ESPECIAL": 5,
             "SIN UBICAR": 9}
    # Solo renglones con existencia: los anulados (cantidad 0) no se reportan y así el
    # total impreso cuadra con las líneas del formato.
    vigentes = inv.activas()
    sin_ubicar = [f for f in vigentes if not f["ubicacion"]]
    filas = [f for f in vigentes if f["ubicacion"]]
    filas.sort(key=lambda f: (orden.get(ubicaciones.zona_de(f["ubicacion"]), 8),
                              f["ubicacion"], str(f["lote"] or "")))

    zona_actual = None
    for f in filas:
        zona = ubicaciones.zona_de(f["ubicacion"])
        if zona != zona_actual:
            zona_actual = zona
            ws.cell(row=fila, column=1, value=f"ZONA: {zona}").font = NEGRITA
            for col in range(1, len(encabezados) + 1):
                ws.cell(row=fila, column=col).fill = GRIS
                ws.cell(row=fila, column=col).border = BORDE
            fila += 1
        for k, v in enumerate(valores_de(f), start=1):
            c = ws.cell(row=fila, column=k, value=v)
            c.font = NORMAL
            c.border = BORDE
        fila += 1

    if sin_ubicar:
        ws.cell(row=fila, column=1, value="ZONA: SIN UBICAR").font = NEGRITA
        for col in range(1, len(encabezados) + 1):
            ws.cell(row=fila, column=col).fill = AMBAR
            ws.cell(row=fila, column=col).border = BORDE
        fila += 1
        for f in sin_ubicar:
            for k, v in enumerate(valores_de(f, es_sin_ubicar=True), start=1):
                c = ws.cell(row=fila, column=k, value=v)
                c.font = NORMAL
                c.border = BORDE
            fila += 1

    fila += 1
    resumen = inv.resumen()
    ws.cell(row=fila, column=1,
            value=f"Total de lotes: {resumen['total']} · con ubicación: "
                  f"{resumen['con_ubicacion']} ({resumen['cobertura']} %) · sin ubicar: "
                  f"{resumen['sin_ubicacion']}").font = NEGRITA
    fila += 2
    _firmas(ws, fila)

    anchos = [16, 48, 12, 8, 22, 14, 8] if version == "3" else [16, 48, 12, 22, 8, 8, 8]
    _anchos(ws, anchos)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:6"
    ws.print_area = f"A1:{get_column_letter(len(encabezados))}{ws.max_row}"


def _portada(wb, pares, version):
    """Resumen de todos los materiales: es la primera página que firma jefatura."""
    ws = wb.create_sheet("Portada", 0)
    _encabezado_f01(ws, version, 99)
    ws["A3"] = "TÍTULO DEL DOCUMENTO: REPORTE DE INVENTARIOS — TODOS LOS MATERIALES"
    ws["A3"].font = Font(bold=True, size=11)
    ws["E2"] = "ALMACÉN: MP · MEP · PT"
    ws["E2"].font = NEGRITA

    encabezados = ["INVENTARIO", "LOTES", "CON UBICACIÓN", "SIN UBICAR", "COBERTURA",
                   "VENCIDOS", "POR VENCER 30 D", "CANTIDAD"]
    fila = _tabla(ws, 6, encabezados)
    totales = {"lotes": 0, "con": 0, "sin": 0, "vencidos": 0, "criticos": 0, "unidades": {}}
    for tipo, inv in pares:
        r = inv.resumen()
        totales["lotes"] += r["total"]
        totales["con"] += r["con_ubicacion"]
        totales["sin"] += r["sin_ubicacion"]
        totales["vencidos"] += r["semaforo"]["VENCIDO"]
        totales["criticos"] += r["semaforo"]["CRITICO"]
        for unidad, valor in r["por_unidad"].items():
            totales["unidades"][unidad] = totales["unidades"].get(unidad, 0) + valor
        cantidad = " · ".join(f"{v:,.1f} {u}" for u, v in sorted(r["por_unidad"].items()))
        valores = [f"{tipo.upper()} — {inv.nombre}", r["total"], r["con_ubicacion"],
                   r["sin_ubicacion"], f"{r['cobertura']} %", r["semaforo"]["VENCIDO"],
                   r["semaforo"]["CRITICO"], cantidad]
        for k, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=k, value=v)
            c.font = NORMAL
            c.border = BORDE
        fila += 1

    cobertura = round(100 * totales["con"] / totales["lotes"], 1) if totales["lotes"] else 0.0
    cantidad_total = " · ".join(f"{v:,.1f} {u}" for u, v in sorted(totales["unidades"].items()))
    valores = ["TOTAL", totales["lotes"], totales["con"], totales["sin"], f"{cobertura} %",
               totales["vencidos"], totales["criticos"], cantidad_total]
    for k, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=k, value=v)
        c.font = NEGRITA
        c.border = BORDE
        c.fill = GRIS
    fila += 2
    ws.cell(row=fila, column=1,
            value="Incluye restos y parcialidades de cada inventario (§V.2.2.6). "
                  "Reporte mensual para MP y MEP (§V.2.1.8); semanal para PT (§V.2.3.8).").font = NORMAL
    fila += 2
    _firmas(ws, fila)
    _anchos(ws, [40, 9, 13, 11, 11, 10, 15, 34])
    for fila_tabla in range(7, fila + 1):
        ws.cell(row=fila_tabla, column=8).font = Font(size=8)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = f"A1:H{ws.max_row}"


def generar_f01(inv, version="2", destino=None):
    """Reporte F01. Con una lista [(tipo, inventario)] genera el consolidado de todo el almacén."""
    pares = [(t, i) for t, i in inv] if isinstance(inv, (list, tuple)) else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"

    if pares is None:
        _hoja_inventario(ws, inv, version)
        nombre = f"F01-ALM-PNO-002_{inv.tipo.upper()}_v{version}_{_marca()}.xlsx"
    else:
        _portada(wb, pares, version)
        for tipo, inventario_tipo in pares:
            _hoja_inventario(wb.create_sheet(tipo.upper()), inventario_tipo, version)
        del wb["Hoja1"]
        wb.move_sheet("Portada", offset=-wb.sheetnames.index("Portada"))
        nombre = f"F01-ALM-PNO-002_TODOS_v{version}_{_marca()}.xlsx"

    destino = pathlib.Path(destino) if destino else _ruta(nombre)
    wb.save(destino)
    return destino


def exportar_inventario(inv, destino=None):
    """El Excel del área, tal cual, con la columna UBICACIÓN + hojas Resumen y Sin ubicar."""
    import openpyxl as ox
    inv.guardar()
    wb = inv.construir_workbook() if getattr(inv, "en_db", False) else ox.load_workbook(inv.ruta)
    resumen = inv.resumen()
    for previa in ("Resumen", "Sin ubicar"):
        if previa in wb.sheetnames:
            del wb[previa]

    hoja = wb.create_sheet("Resumen")
    hoja["A1"] = f"RESUMEN DE UBICACIÓN — {inv.nombre.upper()}"
    hoja["A1"].font = Font(bold=True, size=13)
    hoja["A2"] = f"Generado {resumen['actualizado']} · fuente: {inv.ruta.name}"
    datos = [
        ("Lotes en el inventario", resumen["total"]),
        ("Con ubicación asignada", resumen["con_ubicacion"]),
        ("Sin ubicación (pendientes)", resumen["sin_ubicacion"]),
        ("Cobertura de ubicación", f"{resumen['cobertura']} %"),
        ("Kilogramos declarados", resumen["kg"]),
        ("Lotes vencidos", resumen["semaforo"]["VENCIDO"]),
        ("Vencen en 30 días", resumen["semaforo"]["CRITICO"]),
        ("Vencen en 31-90 días", resumen["semaforo"]["AVISO"]),
        ("Sin caducidad legible", resumen["sin_caducidad"]),
        ("Sin número de lote", resumen["sin_lote"]),
        # un código que no está en el catálogo no es una ubicación: se cuenta como pendiente
        ("Ubicaciones no válidas (fuera del catálogo)", resumen["ubicaciones_invalidas"]),
    ]
    fila = 4
    for etiqueta, valor in datos:
        hoja.cell(row=fila, column=1, value=etiqueta).font = NORMAL
        c = hoja.cell(row=fila, column=2, value=valor)
        c.font = NEGRITA
        c.alignment = Alignment(horizontal="right")
        fila += 1
    fila += 1
    hoja.cell(row=fila, column=1, value="Lotes por zona").font = NEGRITA
    fila += 1
    for zona, n in sorted(resumen["por_zona"].items()):
        hoja.cell(row=fila, column=1, value=zona).font = NORMAL
        hoja.cell(row=fila, column=2, value=n).font = NORMAL
        fila += 1
    _anchos(hoja, [38, 14])

    pendientes = wb.create_sheet("Sin ubicar")
    encabezados = ["FILA", "DESCRIPCIÓN", "LOTE", "CANTIDAD", "UNIDAD", "CADUCIDAD", "SEMÁFORO",
                   "UBICACIÓN ESCRITA EN EL EXCEL"]
    fila = _tabla(pendientes, 1, encabezados)
    for f in inv.activas():
        if f["ubicacion"]:
            continue
        for i, v in enumerate([f["fila"], f["descripcion"], f["lote"], f["cantidad"],
                               f["unidad"], f["caducidad"], f["semaforo"],
                               f["ubicacion_invalida"]], start=1):
            c = pendientes.cell(row=fila, column=i, value=v)
            c.font = NORMAL
            c.border = BORDE
        fila += 1
    _anchos(pendientes, [7, 46, 20, 10, 8, 12, 10, 30])
    pendientes.freeze_panes = "A2"

    destino = pathlib.Path(destino) if destino else _ruta(
        f"Inventario_{inv.tipo.upper()}_con_ubicacion_{_marca()}.xlsx")
    wb.save(destino)
    return destino


def tabla_macro(inv, destino=None):
    """Matriz material x zona (kg) + hoja de caducidades ordenada para FEFO."""
    zonas = ["MP-ALERGENO", "MP-NO-ALERGENO", "PT", "EM", "ESPECIAL"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabla macro"
    ws["A1"] = f"TABLA MACRO DE OCUPACIÓN — {inv.nombre.upper()}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generada {dt.datetime.now():%Y-%m-%d %H:%M} · fuente: {inv.ruta.name}"
    ws["A2"].font = NORMAL

    encabezados = ["MATERIAL", "PROVEEDOR"] + zonas + ["TOTAL", "CADUCIDAD MÁS PRÓXIMA"]
    fila = _tabla(ws, 4, encabezados)

    materiales = {}
    for f in inv.activas():
        clave = f["descripcion"].strip().upper()
        d = materiales.setdefault(clave, {"proveedor": f["proveedor"], "zonas": {},
                                          "total": 0.0, "lotes": [],
                                          "caducidad": None})
        cantidad = f["cantidad"] if isinstance(f["cantidad"], (int, float)) else 0
        d["total"] += cantidad
        if f["ubicacion"]:
            zona = ubicaciones.zona_de(f["ubicacion"])
            d["zonas"][zona] = d["zonas"].get(zona, 0) + cantidad
        d["lotes"].append(f["lote"])
        if f["caducidad"] and (d["caducidad"] is None or f["caducidad"] < d["caducidad"]):
            d["caducidad"] = f["caducidad"]

    for nombre in sorted(materiales):
        d = materiales[nombre]
        ws.cell(row=fila, column=1, value=nombre).font = NORMAL
        ws.cell(row=fila, column=2, value=d["proveedor"]).font = NORMAL
        for i, zona in enumerate(zonas, start=3):
            if zona in d["zonas"]:
                ws.cell(row=fila, column=i, value=round(d["zonas"][zona], 3)).font = NORMAL
        ws.cell(row=fila, column=len(zonas) + 3, value=round(d["total"], 3)).font = NEGRITA
        ws.cell(row=fila, column=len(zonas) + 4, value=d["caducidad"]).font = NORMAL
        for col in range(1, len(encabezados) + 1):
            ws.cell(row=fila, column=col).border = BORDE
        fila += 1

    _anchos(ws, [42, 20] + [13] * len(zonas) + [12, 16])
    ws.freeze_panes = "C5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8  # A3
    ws.page_setup.fitToWidth = 1

    hoja = wb.create_sheet("Caducidades")
    hoja["A1"] = "CONTROL DE CADUCIDADES (FEFO) — ordenado por lo que vence primero"
    hoja["A1"].font = Font(bold=True, size=13)
    encabezados = ["UBICACIÓN", "MATERIAL", "LOTE", "CANTIDAD", "UNIDAD",
                   "CADUCIDAD", "DÍAS", "SEMÁFORO"]
    fila = _tabla(hoja, 3, encabezados)
    ordenados = sorted(inv.activas(),
                       key=lambda f: (f["caducidad"] is None, f["caducidad"] or "9999"))
    for f in ordenados:
        valores = [f["ubicacion"] or "(sin ubicar)", f["descripcion"], f["lote"],
                   f["cantidad"], f["unidad"],
                   f["caducidad"] or (f["caducidad_cruda"] or "sin dato"),
                   f["dias_restantes"], f["semaforo"]]
        for i, v in enumerate(valores, start=1):
            c = hoja.cell(row=fila, column=i, value=v)
            c.font = NORMAL
            c.border = BORDE
            if f["semaforo"] == "VENCIDO":
                c.fill = ROJO
            elif f["semaforo"] in ("CRITICO", "AVISO"):
                c.fill = AMBAR
        fila += 1
    _anchos(hoja, [15, 44, 20, 11, 8, 13, 7, 10])
    hoja.freeze_panes = "A4"
    hoja.auto_filter.ref = f"A3:H{fila - 1}"

    destino = pathlib.Path(destino) if destino else _ruta(
        f"Tabla_macro_{inv.tipo.upper()}_{_marca()}.xlsx")
    wb.save(destino)
    return destino


def exportar_bitacora(con, destino=None, inventario=None):
    from app import bitacora
    movimientos = bitacora.ultimos(con, limite=100000, inventario=inventario)
    movimientos.sort(key=lambda m: m["id"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws["A1"] = "BITÁCORA DE MOVIMIENTOS DE UBICACIÓN — ALMACÉN DE MATERIAS PRIMAS"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Exportada {dt.datetime.now():%Y-%m-%d %H:%M} · {len(movimientos)} movimientos"
    encabezados = ["ID", "FECHA-HORA", "INVENTARIO", "ACCIÓN", "DESCRIPCIÓN", "LOTE", "ORIGEN",
                   "DESTINO", "CANTIDAD", "UNIDAD", "MOTIVO", "REF. DOCUMENTO", "USUARIO", "NOTA"]
    fila = _tabla(ws, 4, encabezados)
    for m in movimientos:
        valores = [m["id"], m["ts"], m.get("inventario") or "mp", m["accion"], m["descripcion"],
                   m["lote"], m["ubicacion_origen"], m["ubicacion_destino"], m.get("cantidad"),
                   m.get("unidad"), m["motivo"], m["ref_doc"], m["usuario"], m["nota"]]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=i, value=v)
            c.font = NORMAL
            c.border = BORDE
        fila += 1
    _anchos(ws, [6, 20, 11, 15, 44, 20, 16, 16, 10, 8, 18, 16, 22, 40])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:N{max(fila - 1, 4)}"
    destino = pathlib.Path(destino) if destino else _ruta(f"Bitacora_movimientos_{_marca()}.xlsx")
    wb.save(destino)
    return destino
