"""Genera la guía operativa de una página (Word + PDF) para el área de almacén."""
import json
import pathlib
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from app import inventario  # noqa: E402

DOCX = RAIZ / "docs" / "Guia_operativa_ubicacion_almacen.docx"
AZUL = RGBColor(0x1F, 0x3B, 0x73)
ROJO = RGBColor(0xB0, 0x1C, 0x1C)
VERDE = RGBColor(0x1E, 0x6B, 0x3A)


def parrafo(doc, texto, negritas=False, color=None, tam=10, espacio=4, alinear=None):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negritas
    run.font.size = Pt(tam)
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(espacio)
    if alinear:
        p.alignment = alinear
    return p


def tabla(doc, encabezados, filas):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    for i, h in enumerate(encabezados):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = AZUL
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8)
    return t


def construir(resumen, destino=DOCX):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)

    parrafo(doc, "COMERCIALIZADORA MIGH DE MÉXICO S.A. DE C.V.", True, AZUL, 11, 0)
    parrafo(doc, "Almacén de materias primas · Aseguramiento de Calidad", False, None, 8, 10)
    parrafo(doc, "Guía operativa — control de ubicación del inventario", True, None, 15, 8)
    parrafo(doc, "Documento de trabajo. No sustituye ALM-PNO-002 ni el formato F01-ALM-PNO-002.",
            False, None, 8, 10).runs[0].italic = True

    parrafo(doc, "1. Punto de partida (medido el 22 de octubre de 2026)", True, AZUL, 11)
    tabla(doc, ["Indicador", "Valor"], [
        ["Lotes en el inventario del área", resumen["total"]],
        ["Kilogramos declarados", f"{resumen['kg']:,.1f}"],
        ["Lotes con ubicación registrada", f"0 de {resumen['total']}"],
        ["Lotes ya vencidos", resumen["semaforo"]["VENCIDO"]],
        ["Lotes que vencen en 30 días", resumen["semaforo"]["CRITICO"]],
        ["Lotes sin caducidad legible", resumen["sin_caducidad"]],
    ])
    parrafo(doc, "El archivo del área no puede contestar «¿qué se me vence primero?». Eso es PEPS/FEFO, "
                 "no papeleo: es merma y desabasto.", False, ROJO, 9, 10)

    parrafo(doc, "2. Cómo se usa (cuatro pasos, menos de un minuto)", True, AZUL, 11)
    for paso in [
        "Abrir el sistema: doble clic en arrancar.cmd. Se abre en el navegador y pide usuario y contraseña "
        "(cada movimiento queda firmado con quien lo hizo).",
        "Buscar el material o el lote y teclear el código del rack en la columna UBICACIÓN, o usar "
        "«Registrar movimiento» para cambiarlo de lugar eligiendo el motivo.",
        "Si se mueven solo unos kilos, escribir esa cantidad en el modal: el resto se queda en su rack y "
        "se crea un renglón nuevo con la parte movida (acción PARCIAL en la bitácora).",
        "Si hubo un error, pestaña Movimientos → Deshacer (el último del lote). El administrador puede "
        "borrar una captura de prueba dejando motivo y registro de auditoría.",
    ]:
        parrafo(doc, f"• {paso}", False, None, 9, 3)

    parrafo(doc, "3. Los cinco Excel, con un clic", True, AZUL, 11)
    tabla(doc, ["Botón", "Qué obtienes", "Para qué sirve"], [
        ["Inventario + ubicación", "El Excel del área con la columna UBICACIÓN, más hojas Resumen y Sin ubicar",
         "Trabajo diario y evidencia de avance"],
        ["Reporte F01", "F01-ALM-PNO-002 versión 2 (Ubicación, Material, Cantidad, Lote) por zona, con firmas",
         "El reporte del periodo, para imprimir y firmar"],
        ["F01 con caducidad (v3)", "La misma tabla más Caducidad y Días", "Propuesta de versión siguiente"],
        ["Tabla macro + FEFO", "Matriz material × zona y caducidades ordenadas", "Pegar en el almacén y planear surtido"],
        ["Bitácora", "Todos los movimientos con usuario, fecha, origen y destino", "Evidencia para auditoría"],
    ])

    parrafo(doc, "4. Reglas que no se negocian", True, AZUL, 11)
    for regla in [
        "No se inventa una ubicación: si no se sabe, se marca «no está».",
        "No se inventa una caducidad: si el dato no es legible, queda como «sin dato».",
        "No se llena el reporte a mano: se genera y se firma.",
        "El papel no cambia: R-ALM-001 y el checklist de recepción se siguen llenando igual.",
    ]:
        parrafo(doc, f"• {regla}", False, None, 9, 3)

    parrafo(doc, "5. Códigos de ubicación válidos", True, AZUL, 11)
    tabla(doc, ["Serie", "Rango", "Materiales"], [
        ["MPA-A", "001 – 040", "Materias primas, zona alérgenos"],
        ["MPNA-A / MPNA-B", "041 – 136", "Materias primas no alérgenas"],
        ["MPT-B", "001 – 020", "Material de producto terminado"],
        ["MEP-B a MEP-E", "001 – 220", "Material de empaque"],
        ["Especiales", "—", "PASILLO-A, JAULA-MP, JAULA-PT, LAB-RETENCION, CUARENTENA-MP, ADUANA-MP, ANDEN"],
    ])
    parrafo(doc, "Si el código no existe, el sistema lo rechaza. Las ubicaciones MPNA-B-131 a 136 y las "
                 "jaulas no están en el formato F01 vigente aunque sí en el procedimiento: hay que "
                 "agregarlas en la próxima revisión del formato.", False, VERDE, 9, 10)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


def a_pdf(docx_path):
    import win32com.client as win32
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(docx_path))
    pdf = docx_path.with_suffix(".pdf")
    doc.ExportAsFixedFormat(str(pdf), 17)
    doc.Close(False)
    word.Quit()
    return pdf


if __name__ == "__main__":
    resumen = inventario.Inventario().resumen()
    ruta = construir(resumen)
    print("docx:", ruta)
    print("pdf :", a_pdf(ruta))
