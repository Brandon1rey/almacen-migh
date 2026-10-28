"""Sincronización de la base real hacia la USB: copia consistente, manifiesto e historial."""
import json
import pathlib
import sqlite3
import tempfile
import unittest

from app import bitacora

import openpyxl  # noqa: E402  (después de app, que habilita vendor\)
from scripts import sincronizar_usb


def _inventario(ruta, filas=2):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(["Descripción", "Proveedor", "Categoría", "Unidad de medida",
                               "Cantidad", "Caducidad vigente", "Lote actual", "Entrada ERP",
                               "Observaciones", "UBICACIÓN"], start=1):
        ws.cell(row=3, column=i, value=texto)
    for n in range(filas):
        ws.cell(row=4 + n, column=1, value=f"MATERIAL {n}")
        ws.cell(row=4 + n, column=5, value=10 + n)
        ws.cell(row=4 + n, column=7, value=f"LOTE-{n}")
    wb.save(ruta)


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.raiz = pathlib.Path(self._tmp.name)
        self.origen = self.raiz / "data"
        self.destino = self.raiz / "usb"
        self.origen.mkdir()
        self.destino.mkdir()

        self.con = bitacora.conectar(self.origen / "bitacora.db")
        registro = {"fila": 4, "descripcion": "WPC 80%", "lote": "L-1", "tipo": "mp"}
        bitacora.registrar(self.con, "UBICAR", registro, "Esteban", destino="MPA-A-005",
                           inventario="mp")
        bitacora.registrar(self.con, "MOVER", registro, "Esteban", origen="MPA-A-005",
                           destino="MPA-A-006", inventario="mp")
        self.con.commit()

        _inventario(self.origen / "Inventario_MP.xlsx")
        _inventario(self.origen / "Inventario_MEP.xlsx")

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()


class TestCopia(Base):
    def test_copia_base_e_inventarios(self):
        resultado = sincronizar_usb.sincronizar(self.origen, self.destino)
        ultima = self.destino / "ultima"
        self.assertTrue((ultima / "bitacora.db").exists())
        self.assertTrue((ultima / "Inventario_MP.xlsx").exists())
        self.assertTrue((ultima / "Inventario_MEP.xlsx").exists())
        self.assertEqual(resultado["movimientos"], 2)

    def test_la_copia_es_una_base_valida_y_completa(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        copia = sqlite3.connect(self.destino / "ultima" / "bitacora.db")
        self.assertEqual(copia.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertEqual(copia.execute("SELECT COUNT(*) FROM movimiento").fetchone()[0], 2)
        self.assertEqual(copia.execute("SELECT COUNT(*) FROM usuario").fetchone()[0], 0)
        copia.close()

    def test_manifiesto_con_hashes_y_conteos(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        manifiesto = json.loads((self.destino / "ultima" / "sincronizacion.json").read_text("utf-8"))
        self.assertEqual(manifiesto["movimientos"], 2)
        self.assertIn("bitacora.db", manifiesto["archivos"])
        self.assertEqual(len(manifiesto["archivos"]["bitacora.db"]["sha256"]), 64)
        self.assertEqual(manifiesto["inventarios"]["mp"]["lotes"], 2)
        self.assertTrue(manifiesto["origen"].endswith("data"))

    def test_deja_leeme_en_la_raiz_del_destino(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        leeme = self.destino / "LEEME.txt"
        self.assertTrue(leeme.exists())
        self.assertIn("COPIA", leeme.read_text("utf-8"))

    def test_no_modifica_el_origen(self):
        antes = (self.origen / "bitacora.db").stat().st_mtime_ns
        sincronizar_usb.sincronizar(self.origen, self.destino)
        self.assertEqual((self.origen / "bitacora.db").stat().st_mtime_ns, antes)
        self.assertEqual(self.con.execute("SELECT COUNT(*) FROM movimiento").fetchone()[0], 2)

    def test_correr_dos_veces_no_falla(self):
        primera = sincronizar_usb.sincronizar(self.origen, self.destino)
        segunda = sincronizar_usb.sincronizar(self.origen, self.destino)
        self.assertEqual(primera["movimientos"], segunda["movimientos"])
        self.assertEqual(len(list((self.destino / "ultima").glob("*.xlsx"))), 2)

    def test_historial_guarda_una_copia_por_corrida(self):
        sincronizar_usb.sincronizar(self.origen, self.destino, historial=True)
        sincronizar_usb.sincronizar(self.origen, self.destino, historial=True)
        copias = list((self.destino / "historial").glob("bitacora_*.db"))
        self.assertEqual(len(copias), 2)

    def test_sin_inventarios(self):
        sincronizar_usb.sincronizar(self.origen, self.destino, con_inventarios=False)
        self.assertEqual(list((self.destino / "ultima").glob("*.xlsx")), [])

    def test_destino_inexistente_avisa(self):
        with self.assertRaises(FileNotFoundError):
            sincronizar_usb.sincronizar(self.origen, self.raiz / "no_existe", crear_destino=False)

    def test_destino_inexistente_se_crea(self):
        nuevo = self.raiz / "usb_nueva"
        sincronizar_usb.sincronizar(self.origen, nuevo)
        self.assertTrue((nuevo / "ultima" / "bitacora.db").exists())


class TestComparar(Base):
    def test_sin_copia_todo_pendiente(self):
        estado = sincronizar_usb.comparar(self.origen, self.destino)
        self.assertFalse(estado["al_dia"])
        self.assertEqual(estado["base"]["estado"], "falta en la copia")
        self.assertIn("bitacora.db", estado["pendientes"])

    def test_al_dia_despues_de_copiar(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        estado = sincronizar_usb.comparar(self.origen, self.destino)
        self.assertTrue(estado["al_dia"], estado)
        self.assertEqual(estado["base"]["estado"], "igual")
        self.assertIsNotNone(estado["ultima_copia"])

    def test_detecta_movimientos_nuevos(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        bitacora.registrar(self.con, "MOVER", {"fila": 4, "descripcion": "WPC 80%", "lote": "L-1"},
                           "Esteban", origen="MPA-A-005", destino="MPA-A-006", inventario="mp")
        self.con.commit()
        estado = sincronizar_usb.comparar(self.origen, self.destino)
        self.assertFalse(estado["al_dia"])
        self.assertEqual(estado["base"]["estado"], "hay cambios por copiar")
        self.assertEqual(estado["base"]["movimientos_origen"], 3)
        self.assertEqual(estado["base"]["movimientos_copia"], 2)

    def test_no_pisa_una_copia_mas_adelantada(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        # la USB trae movimientos que este servidor no tiene (viene de otra PC)
        copia = sqlite3.connect(self.destino / "ultima" / "bitacora.db")
        copia.execute("INSERT INTO movimiento (ts, accion, inventario, descripcion, usuario) "
                      "VALUES ('2026-10-20 10:00:00','UBICAR','mep','OTRO','Teresa')")
        copia.commit()
        copia.close()
        estado = sincronizar_usb.comparar(self.origen, self.destino)
        self.assertEqual(estado["base"]["estado"], "la copia está más al día")
        self.assertIn("bitacora.db", estado["adelantados"])

        manifiesto = sincronizar_usb.sincronizar(self.origen, self.destino)
        self.assertEqual(manifiesto["movimientos"], 3)      # sigue teniendo 3, no se pisó
        self.assertTrue(manifiesto["avisos"])
        comprobacion = sqlite3.connect(self.destino / "ultima" / "bitacora.db")
        self.assertEqual(comprobacion.execute("SELECT COUNT(*) FROM movimiento").fetchone()[0], 3)
        comprobacion.close()

    def test_forzar_si_sobrescribe(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        copia = sqlite3.connect(self.destino / "ultima" / "bitacora.db")
        copia.execute("INSERT INTO movimiento (ts, accion, inventario, descripcion, usuario) "
                      "VALUES ('2026-10-20 10:00:00','UBICAR','mep','OTRO','Teresa')")
        copia.commit()
        copia.close()
        manifiesto = sincronizar_usb.sincronizar(self.origen, self.destino, forzar=True)
        self.assertEqual(manifiesto["movimientos"], 2)

    def test_segunda_corrida_no_reescribe(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        manifiesto = sincronizar_usb.sincronizar(self.origen, self.destino)
        self.assertTrue(manifiesto["al_dia"])
        self.assertEqual(manifiesto["omitidos"].get("bitacora.db"), "igual")
        self.assertIn("Inventario_MP.xlsx", manifiesto["omitidos"])


class TestResumen(Base):
    def test_resumen_de_la_copia(self):
        sincronizar_usb.sincronizar(self.origen, self.destino)
        resumen = sincronizar_usb.resumen(self.destino / "ultima" / "bitacora.db")
        self.assertEqual(resumen["movimientos"], 2)
        self.assertEqual(resumen["por_inventario"], {"mp": 2})
        self.assertEqual(resumen["ultimo"]["accion"], "MOVER")
        self.assertEqual(resumen["integridad"], "ok")

    def test_resumen_sin_base(self):
        with self.assertRaises(FileNotFoundError):
            sincronizar_usb.resumen(self.raiz / "nada.db")


if __name__ == "__main__":
    unittest.main()
