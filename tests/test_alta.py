"""Alta de un lote nuevo: validaciones, escritura en el Excel y rastro en la bitácora."""
import datetime as dt
import pathlib
import tempfile
import unittest

from app import api, bitacora, inventario, usuarios  # primero app: habilita vendor\openpyxl

import openpyxl

ENCABEZADOS = ["Descripción", "Proveedor", "Categoría", "Unidad de medida", "Cantidad",
               "Caducidad vigente", "Lote actual", "Entrada ERP", "Observaciones", "UBICACIÓN"]


def libro(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, texto in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=3, column=i, value=texto)
    ws.cell(row=4, column=1, value="WPC 80%")
    ws.cell(row=4, column=2, value="VILHER")
    ws.cell(row=4, column=4, value="kg")
    ws.cell(row=4, column=5, value=100)
    ws.cell(row=4, column=6, value="2028-04-23")
    ws.cell(row=4, column=7, value="22260425002")
    ws.cell(row=4, column=8, value="MIGH-MP114")
    wb.save(ruta)
    return ruta


def datos(**cambios):
    base = {
        "descripcion": "LACTOSA 200 MESH",
        "proveedor": "LACTOSA MX",
        "categoria": "Carbohidrato",
        "unidad": "kg",
        "cantidad": 500,
        "caducidad": (dt.date.today() + dt.timedelta(days=400)).isoformat(),
        "lote": "LAC-2601",
        "erp": "MIGH-MP200",
        "ubicacion": "MPNA-A-050",
        "observaciones": "",
    }
    base.update(cambios)
    return base


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)
        self.ruta = libro(self.dir / "inv.xlsx")
        self.inv = inventario.Inventario(self.ruta, tipo="mp")
        self.con = bitacora.conectar(self.dir / "b.db")
        usuarios.crear(self.con, "Esteban", "clave-segura-1", "almacen")
        fila = usuarios.por_nombre(self.con, "Esteban")
        self.sesion = {"usuario_id": fila["id"], "nombre": "Esteban", "rol": "almacen",
                       "debe_cambiar": 0}
        self.estado = type("Estado", (), {})()
        self.estado.carpeta_datos = self.dir
        self.estado.con = self.con
        self.estado.archivo = "inv.xlsx"
        self.estado.inventario = lambda tipo=None: self.inv

    def tearDown(self):
        self.con.close()
        self._tmp.cleanup()


class TestAltaDirecta(Base):
    def test_agrega_el_renglon_al_excel(self):
        registro = self.inv.agregar_alta(datos())
        self.assertEqual(registro["descripcion"], "LACTOSA 200 MESH")
        self.assertEqual(registro["cantidad"], 500)
        self.assertEqual(registro["ubicacion"], "MPNA-A-050")
        self.inv.guardar()                      # el API guarda dentro de _persistir
        wb = openpyxl.load_workbook(self.ruta)
        ws = wb.active
        self.assertEqual(ws.cell(row=5, column=1).value, "LACTOSA 200 MESH")
        self.assertEqual(ws.cell(row=5, column=7).value, "LAC-2601")

    def test_queda_en_la_lista_y_en_el_resumen(self):
        antes = self.inv.resumen()["total"]
        self.inv.agregar_alta(datos())
        self.assertEqual(self.inv.resumen()["total"], antes + 1)
        self.assertEqual(len(self.inv.buscar("LAC-2601")), 1)

    def test_normaliza_unidad_y_ubicacion(self):
        registro = self.inv.agregar_alta(datos(unidad="KG", ubicacion="mpna-a-050"))
        self.assertEqual(registro["unidad"], "kg")
        self.assertEqual(registro["ubicacion"], "MPNA-A-050")

    def test_rechaza_campos_faltantes(self):
        for campo in ("descripcion", "lote", "proveedor", "cantidad", "caducidad"):
            with self.subTest(campo=campo):
                with self.assertRaises(ValueError):
                    self.inv.agregar_alta(datos(**{campo: ""}))

    def test_rechaza_cantidad_no_positiva(self):
        for cantidad in (0, -5):
            with self.subTest(cantidad=cantidad):
                with self.assertRaises(ValueError):
                    self.inv.agregar_alta(datos(cantidad=cantidad))

    def test_rechaza_lote_duplicado_del_mismo_material(self):
        self.inv.agregar_alta(datos())
        with self.assertRaises(ValueError):
            self.inv.agregar_alta(datos(cantidad=10))

    def test_permite_el_mismo_lote_en_otro_material(self):
        self.inv.agregar_alta(datos())
        registro = self.inv.agregar_alta(datos(descripcion="LACTOSA 100 MESH", erp="MIGH-MP201"))
        self.assertEqual(registro["lote"], "LAC-2601")

    def test_rechaza_caducidad_vencida(self):
        ayer = (dt.date.today() - dt.timedelta(days=1)).isoformat()
        with self.assertRaises(ValueError):
            self.inv.agregar_alta(datos(caducidad=ayer))

    def test_rechaza_caducidad_ilegible(self):
        with self.assertRaises(ValueError):
            self.inv.agregar_alta(datos(caducidad="26/17/07"))

    def test_rechaza_ubicacion_inexistente(self):
        with self.assertRaises(ValueError):
            self.inv.agregar_alta(datos(ubicacion="ZZZ-999"))

    def test_sin_ubicacion_queda_pendiente(self):
        registro = self.inv.agregar_alta(datos(ubicacion=""))
        self.assertIsNone(registro["ubicacion"])

    def test_quitar_la_ubicacion_se_guarda_en_el_excel(self):
        registro = self.inv.agregar_alta(datos())
        self.inv.asignar(registro["fila"], None)
        self.inv.guardar()
        releido = inventario.Inventario(self.ruta, tipo="mp")
        self.assertIsNone(releido.por_fila(registro["fila"])["ubicacion"])


class TestAltaApi(Base):
    def api(self, cuerpo):
        return api.manejar(self.estado, "POST", "/api/alta", {}, cuerpo, self.sesion)

    def test_alta_por_api(self):
        codigo, res, _ = self.api(datos())
        self.assertEqual(codigo, 200, res)
        self.assertEqual(res["lote"], "LAC-2601")
        self.assertEqual(res["accion"], "ALTA")
        self.assertIn("resumen", res)

    def test_la_bitacora_registra_el_alta(self):
        self.api(datos())
        movs = bitacora.ultimos(self.con, 10, inventario="mp")
        self.assertEqual(movs[0]["accion"], "ALTA")
        self.assertEqual(movs[0]["lote"], "LAC-2601")
        self.assertEqual(movs[0]["usuario"], "Esteban")
        self.assertEqual(movs[0]["cantidad"], 500)

    def test_error_de_validacion_devuelve_400_con_mensaje(self):
        codigo, res, _ = self.api(datos(lote=""))
        self.assertEqual(codigo, 400)
        self.assertIn("lote", res["error"].lower())

    def test_deshacer_un_alta_la_anula(self):
        _, res, _ = self.api(datos())
        movimiento = res["movimiento"]
        codigo, deshacer, _ = api.manejar(self.estado, "POST", "/api/deshacer", {},
                                          {"movimiento": movimiento}, self.sesion)
        self.assertEqual(codigo, 200, deshacer)
        registro = self.inv.por_fila(res["fila"])
        self.assertEqual(registro["cantidad"], 0)
        self.assertIsNone(registro["ubicacion"])


if __name__ == "__main__":
    unittest.main()
